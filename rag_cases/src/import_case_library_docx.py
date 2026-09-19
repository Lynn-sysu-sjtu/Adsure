import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document

try:
    from src import import_sector_docx
except ModuleNotFoundError:  # Support direct execution: python3 src/import_case_library_docx.py
    import import_sector_docx


DEFAULT_INPUTS = [
    Path("保健食品广告合规案例库.docx"),
    Path("美妆广告合规案例库.docx"),
    Path("游戏广告合规案例库.docx"),
]
DEFAULT_CANDIDATES_DIR = Path("data/structured_candidates")
DEFAULT_RAW_TEXT_DIR = Path("data/raw_text")
DEFAULT_REPORT_PATH = Path("data/reports/case_library_docx_import_report.json")
SOURCE_TYPE = "manual_docx_sector_report_pending_source_verification"

SECTOR_BY_FILENAME = {
    "保健食品广告合规案例库.docx": "health",
    "美妆广告合规案例库.docx": "beauty",
    "游戏广告合规案例库.docx": "game",
}

LABELS = {
    "案号": "case_number",
    "审理法院": "court",
    "审理日期": "decision_date",
    "基本事实": "facts_summary",
    "违反条款": "legal_basis_text",
    "处罚结果": "penalty_result",
    "裁判要点": "decision_highlights",
}


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_case_number(value: str) -> str:
    normalized = normalize(value)
    match = re.match(r"([(（]20\d{2}[)）].*?号)", normalized)
    return match.group(1) if match else normalized


def parse_cases(path: Path) -> list[dict]:
    cases: list[dict] = []
    current: dict | None = None
    for paragraph in Document(path).paragraphs:
        text = normalize(paragraph.text)
        if not text:
            continue
        title_match = re.match(r"案例\d+[:：](.+)", text)
        if title_match:
            if current:
                cases.append(current)
            current = {"title": normalize(title_match.group(1)), "raw_lines": [text]}
            continue
        if current is None:
            continue
        current["raw_lines"].append(text)
        for label, field in LABELS.items():
            prefix = f"{label}："
            if text.startswith(prefix):
                value = normalize(text[len(prefix) :])
                current[field] = normalize_case_number(value) if field == "case_number" else value
                break
    if current:
        cases.append(current)
    return cases


def violation_type(sector: str, case: dict) -> str:
    text = f"{case.get('facts_summary', '')} {case.get('legal_basis_text', '')}"
    if sector == "game":
        return "游戏广告允诺不准确"
    if sector == "beauty":
        if "专利" in text:
            return "化妆品虚假广告及医疗用语"
        return "化妆品广告使用医疗用语"
    if "反不正当竞争法" in text or "会议" in text or "讲座" in text:
        return "会销保健品虚假宣传"
    return "保健食品广告涉及疾病预防治疗功能"


def parse_date(value: str) -> str | None:
    match = re.search(r"(20\d{2})年(\d{1,2})月(\d{1,2})日", value or "")
    if not match:
        return value or None
    return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def parse_penalty_amount(value: str) -> int | None:
    amounts = [int(item) for item in re.findall(r"罚款\s*(\d+)元", value or "")]
    return max(amounts) if amounts else None


def raw_archive(path: Path, raw_text_dir: Path) -> Path:
    doc = Document(path)
    paragraphs = [normalize(item.text) for item in doc.paragraphs if normalize(item.text)]
    tables = [
        [[normalize(cell.text) for cell in row.cells] for row in table.rows]
        for table in doc.tables
    ]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    output = raw_text_dir / f"case_library_docx__{path.stem}.json"
    payload = {
        "source_name": path.name,
        "source_type": SOURCE_TYPE,
        "sha256": digest,
        "paragraphs": paragraphs,
        "tables": tables,
        "full_text": "\n".join(paragraphs),
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def build_case(sector: str, source_path: Path, raw_path: Path, parsed: dict) -> dict:
    violation = violation_type(sector, parsed)
    row = {
        "违法类型": violation,
        "案例名称": parsed["title"],
        "案号": parsed.get("case_number", ""),
        "审理法院": parsed.get("court", ""),
        "审理日期": parse_date(parsed.get("decision_date", "")) or "",
        "关键违法事实": parsed.get("facts_summary", ""),
        "处罚依据": parsed.get("legal_basis_text", ""),
        "处罚结果": parsed.get("penalty_result", ""),
    }
    case = import_sector_docx.build_case(sector, row)
    case.update(
        {
            "source_name": source_path.name,
            "raw_text_path": str(raw_path),
            "penalty_amount": parse_penalty_amount(parsed.get("penalty_result", "")),
            "decision_highlights": parsed.get("decision_highlights"),
            "notes": "由用户提供的分赛道DOCX案例库导入；已归档原文，但缺少官方source_url，需来源核验和法务复核后方可进入production RAG。",
            "source_documents": [
                {
                    "source_name": source_path.name,
                    "raw_text_path": str(raw_path),
                    "verification_status": "pending_source_lookup",
                }
            ],
        }
    )
    case["audit"]["review_notes"] = "Imported from user-provided sector DOCX case library; official source_url pending verification."
    case["vector_text"] = import_sector_docx.vector_text(case)
    return case


def merge_provenance(existing: dict, source_path: Path, raw_path: Path) -> dict:
    provenance = {
        "source_name": source_path.name,
        "raw_text_path": str(raw_path),
        "verification_status": "pending_source_lookup",
    }
    documents = existing.get("source_documents") or []
    if not any(item.get("source_name") == source_path.name for item in documents):
        documents.append(provenance)
    existing["source_documents"] = documents
    existing["raw_text_path"] = existing.get("raw_text_path") or str(raw_path)
    return existing


def run(
    inputs: list[Path] = DEFAULT_INPUTS,
    candidates_dir: Path = DEFAULT_CANDIDATES_DIR,
    raw_text_dir: Path = DEFAULT_RAW_TEXT_DIR,
    report_path: Path | None = DEFAULT_REPORT_PATH,
) -> dict:
    candidates_dir.mkdir(parents=True, exist_ok=True)
    raw_text_dir.mkdir(parents=True, exist_ok=True)
    report = {
        "documents": 0,
        "parsed_cases": 0,
        "new_cases_created_this_run": 0,
        "existing_cases_updated_this_run": 0,
        "library_owned_records_present": 0,
        "preexisting_records_linked": 0,
        "by_sector": {},
        "case_ids": [],
    }
    for input_path in inputs:
        sector = SECTOR_BY_FILENAME.get(input_path.name)
        if not sector:
            raise ValueError(f"Unsupported case-library filename: {input_path.name}")
        archive = raw_archive(input_path, raw_text_dir)
        parsed_cases = parse_cases(input_path)
        report["documents"] += 1
        report["parsed_cases"] += len(parsed_cases)
        report["by_sector"][sector] = len(parsed_cases)
        for parsed in parsed_cases:
            case = build_case(sector, input_path, archive, parsed)
            report["case_ids"].append(case["case_id"])
            output = candidates_dir / f"{case['case_id']}.json"
            if output.exists():
                existing = json.loads(output.read_text(encoding="utf-8"))
                if existing.get("source_name") == input_path.name:
                    report["library_owned_records_present"] += 1
                else:
                    report["preexisting_records_linked"] += 1
                merged = merge_provenance(existing, input_path, archive)
                output.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")
                report["existing_cases_updated_this_run"] += 1
            else:
                output.write_text(json.dumps(case, ensure_ascii=False, indent=2), encoding="utf-8")
                report["new_cases_created_this_run"] += 1
                report["library_owned_records_present"] += 1
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Import standalone sector case-library DOCX files.")
    parser.add_argument("--input", type=Path, action="append", dest="inputs")
    parser.add_argument("--candidates-dir", type=Path, default=DEFAULT_CANDIDATES_DIR)
    parser.add_argument("--raw-text-dir", type=Path, default=DEFAULT_RAW_TEXT_DIR)
    parser.add_argument("--report-path", type=Path, default=DEFAULT_REPORT_PATH)
    args = parser.parse_args()
    report = run(args.inputs or DEFAULT_INPUTS, args.candidates_dir, args.raw_text_dir, args.report_path)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
