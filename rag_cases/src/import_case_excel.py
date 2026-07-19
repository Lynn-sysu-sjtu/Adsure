import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_INPUT_PATH = Path("data/imports/40个十大类违法广告行政处罚案例汇总表.xlsx")
DEFAULT_CANDIDATES_DIR = Path("data/structured_candidates")
DEFAULT_REPORTS_DIR = Path("data/reports")
DEFAULT_ANALYSIS_DIR = Path("data/analysis")

SOURCE_NAME = "40个十大类违法广告行政处罚案例汇总表.xlsx"
SOURCE_TYPE = "manual_compilation_pending_source_verification"
NOTES = (
    "Excel人工检索候选案例，缺少source_url，需后续根据案号或官方公开页面补充来源链接后方可进入production RAG。"
)

HEADERS = {
    "序号",
    "违法类型",
    "案例名称",
    "案号",
    "审理法院",
    "审理日期",
    "关键违法事实",
    "处罚依据",
    "处罚结果",
}

CATEGORY_PATTERN = re.compile(r"^[一二三四五六七八九十]+、\s*(.+)$")
STOP_SECTIONS = {"执法趋势分析", "合规建议"}

CANONICAL_TYPES = [
    "普通食品宣称疾病预防、治疗功能",
    "保健食品功效保证、替代药物、超批准功能宣传",
    "化妆品宣称医疗作用",
    "医疗美容/医疗器械功效保证",
    "绝对化用语",
    "虚假宣传、引人误解宣传",
    "广告引证内容不真实、不准确",
    "房地产升值承诺、规划误导",
    "直播带货违法广告",
    "软文、健康科普变相广告",
]

SLUGS = {
    "普通食品宣称疾病预防、治疗功能": "food_medical_claim",
    "保健食品功效保证、替代药物、超批准功能宣传": "health_food_overclaim",
    "化妆品宣称医疗作用": "cosmetic_medical_claim",
    "医疗美容/医疗器械功效保证": "medical_beauty_device_guarantee",
    "绝对化用语": "absolute_terms",
    "虚假宣传、引人误解宣传": "false_misleading_claim",
    "广告引证内容不真实、不准确": "improper_citation",
    "房地产升值承诺、规划误导": "real_estate_misleading",
    "直播带货违法广告": "live_commerce_ad",
    "软文、健康科普变相广告": "native_health_ad",
}

RISK_MAPPING = {
    "普通食品宣称疾病预防、治疗功能": ["涉医疗宣传", "普通食品疾病治疗功效宣传", "虚假宣传"],
    "保健食品功效保证、替代药物、超批准功能宣传": ["保健食品违规宣传", "涉医疗宣传", "虚假宣传"],
    "化妆品宣称医疗作用": ["化妆品医疗化宣传", "涉医疗宣传"],
    "医疗美容/医疗器械功效保证": ["医疗广告违规", "医疗器械广告违规", "功效保证"],
    "绝对化用语": ["绝对化用语"],
    "虚假宣传、引人误解宣传": ["虚假宣传", "引人误解宣传"],
    "广告引证内容不真实、不准确": ["广告引证内容不规范"],
    "房地产升值承诺、规划误导": ["房地产广告误导", "升值承诺", "虚假宣传"],
    "直播带货违法广告": ["直播带货违法广告", "虚假宣传", "平台广告合规"],
    "软文、健康科普变相广告": ["广告可识别性不足", "变相广告", "涉医疗宣传"],
}

INDUSTRY_MAPPING = {
    "普通食品宣称疾病预防、治疗功能": "普通食品",
    "保健食品功效保证、替代药物、超批准功能宣传": "保健食品",
    "化妆品宣称医疗作用": "化妆品",
    "医疗美容/医疗器械功效保证": "医疗美容/医疗器械",
    "绝对化用语": "一般商品或服务",
    "虚假宣传、引人误解宣传": "一般商品或服务",
    "广告引证内容不真实、不准确": "一般商品或服务",
    "房地产升值承诺、规划误导": "房地产",
    "直播带货违法广告": "直播带货/网络营销",
    "软文、健康科普变相广告": "医疗健康/保健服务",
}

CHANNEL_TERMS = [
    "官网",
    "天猫",
    "阿里巴巴网店",
    "微信朋友圈",
    "微信公众号",
    "抖音",
    "广播",
    "电视",
    "宣传单",
    "宣传单页",
    "会销",
    "会议",
    "售楼处",
    "网页",
    "网站",
    "海报",
    "横幅",
    "报纸",
    "期刊",
    "直播",
    "网店",
    "朋友圈",
    "PPT",
]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return re.sub(r"\s+", " ", value).strip()
    return str(value).strip()


def normalize_violation_type(value: str) -> str:
    text = normalize_text(value)
    for canonical in CANONICAL_TYPES:
        if text == canonical or canonical.startswith(text) or text.startswith(canonical[:8]):
            return canonical
    if text.startswith("保健食品功效保证、替代药物"):
        return "保健食品功效保证、替代药物、超批准功能宣传"
    return text


def excel_date_to_iso(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
    text = normalize_text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text or None


def stable_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:8]


def make_case_id(violation_type: str, case_number: str, title: str, court: str, decision_date: str | None) -> str:
    slug = SLUGS.get(violation_type, re.sub(r"[^a-z0-9]+", "_", violation_type.lower()).strip("_") or "unknown")
    basis = case_number or f"{title}{court}{decision_date or ''}"
    return f"excel_candidate__{slug}__{stable_hash(basis)}"


def first_match(pattern: str, text: str) -> str | None:
    match = re.search(pattern, text)
    return match.group(1).strip() if match else None


def extract_penalty_authority(title: str) -> str | None:
    return first_match(r"([\u4e00-\u9fa5]{2,}(?:市场监督管理局|工商行政管理局|食品药品监督管理局))", title)


def extract_party_name(title: str, authority: str | None) -> str | None:
    text = title
    if authority:
        text = text.replace(authority, "")
    text = re.sub(r"(?:与|、|申请执行)", " ", text)
    text = re.sub(r"(?:非诉执行审查案|非诉审查案|行政处罚案|行政案|买卖合同纠纷案|产品责任纠纷案|再审案)$", "", text)
    parts = [part.strip(" ，、") for part in text.split() if part.strip(" ，、")]
    if not parts:
        return None
    return max(parts, key=len)[:80]


def infer_region(title: str, court: str) -> str | None:
    text = court or title
    match = re.search(r"([\u4e00-\u9fa5]{2,8}(?:省|市|县|区))", text)
    if match:
        return match.group(1)
    for suffix in ("省", "市", "县", "区"):
        idx = text.find(suffix)
        if idx >= 1:
            return text[: idx + 1]
    return None


def extract_product_or_service(facts: str) -> str | None:
    quoted = re.findall(r"[“\"]([^”\"]{2,40})[”\"]", facts)
    if quoted:
        return quoted[0]
    match = re.search(r"(?:销售|宣传|发布|经营)([^，。；;]{2,30})", facts)
    return match.group(1).strip() if match else None


def extract_ad_channel(facts: str) -> str | None:
    found = [term for term in CHANNEL_TERMS if term in facts]
    return "、".join(dict.fromkeys(found)) if found else None


def extract_illegal_claims(facts: str) -> list[str]:
    claims = [item.strip() for item in re.findall(r"[“\"]([^”\"]{2,80})[”\"]", facts) if item.strip()]
    if claims:
        return list(dict.fromkeys(claims))
    risky_terms = [
        "降血糖",
        "治疗",
        "治愈",
        "预防",
        "最",
        "第一",
        "唯一",
        "升值",
        "回报",
        "有效率",
    ]
    sentences = [part.strip() for part in re.split(r"[。；;]", facts) if part.strip()]
    return [sentence for sentence in sentences if any(term in sentence for term in risky_terms)][:3]


def extract_penalty_amount(penalty_result: str) -> float | None:
    match = re.search(r"罚款\s*([0-9]+(?:\.[0-9]+)?)\s*万?元", penalty_result)
    if not match:
        return None
    value = float(match.group(1))
    if "万元" in match.group(0):
        value *= 10000
    return int(value) if value.is_integer() else value


def extract_keywords(violation_type: str, facts: str) -> list[str]:
    candidates = [violation_type, *RISK_MAPPING.get(violation_type, [])]
    candidates.extend(term for term in CHANNEL_TERMS if term in facts)
    candidates.extend(re.findall(r"[“\"]([^”\"]{2,12})[”\"]", facts))
    return list(dict.fromkeys(item for item in candidates if item))


def regulatory_logic(violation_type: str, facts: str, legal_basis: str) -> str:
    basis = f"，涉及{legal_basis}" if legal_basis else ""
    return f"该候选案例显示，广告内容属于{violation_type}场景，相关宣传可能影响消费者判断{basis}，需核验官方来源后确认监管认定。"


def vector_text_for_case(
    violation_type: str,
    facts: str,
    legal_basis: str,
    penalty_result: str,
) -> str:
    return (
        f"这是一起人工整理的{violation_type}候选案例。案件事实显示，经营者{facts}"
        f" 处罚依据记载为{legal_basis or '待补充'}，处理结果为{penalty_result or '待补充'}；"
        "该内容仅用于候选检索，需补充可核验 source_url 后才能进入生产 RAG。"
    )


def is_header_row(values: list[str]) -> bool:
    return len(HEADERS.intersection(values)) >= 6


def row_to_mapping(values: list[Any], header_map: dict[str, int]) -> dict[str, Any]:
    return {name: values[idx] if idx < len(values) else None for name, idx in header_map.items()}


def build_case(row: dict[str, Any], current_category: str | None) -> dict:
    violation_type = normalize_violation_type(normalize_text(row.get("违法类型")) or (current_category or ""))
    title = normalize_text(row.get("案例名称"))
    case_number = normalize_text(row.get("案号"))
    court = normalize_text(row.get("审理法院"))
    decision_date = excel_date_to_iso(row.get("审理日期"))
    facts = normalize_text(row.get("关键违法事实"))
    basis_text = normalize_text(row.get("处罚依据"))
    penalty_result = normalize_text(row.get("处罚结果"))
    authority = extract_penalty_authority(title)
    case_id = make_case_id(violation_type, case_number, title, court, decision_date)
    return {
        "case_id": case_id,
        "title": title,
        "case_number": case_number or None,
        "source_type": SOURCE_TYPE,
        "source_name": SOURCE_NAME,
        "source_url": None,
        "source_verification_status": "pending_source_lookup",
        "publish_date": None,
        "decision_date": decision_date,
        "penalty_authority": authority,
        "party_name": extract_party_name(title, authority),
        "region": infer_region(title, court),
        "court": court or None,
        "industry": INDUSTRY_MAPPING.get(violation_type),
        "product_or_service": extract_product_or_service(facts),
        "ad_channel": extract_ad_channel(facts),
        "risk_dimensions": RISK_MAPPING.get(violation_type, ["其他"]),
        "illegal_claims": extract_illegal_claims(facts),
        "facts_summary": facts,
        "legal_basis": [basis_text] if basis_text else [],
        "penalty_result": penalty_result,
        "penalty_amount": extract_penalty_amount(penalty_result),
        "regulatory_logic": regulatory_logic(violation_type, facts, basis_text),
        "mapped_rule_ids": [],
        "keywords": extract_keywords(violation_type, facts),
        "vector_text": vector_text_for_case(violation_type, facts, basis_text, penalty_result),
        "rag_chunk_type": "case_summary",
        "review_status": "pending_review",
        "approved_for_rag": False,
        "notes": NOTES,
        "audit": {
            "review_status": "pending_review",
            "reviewer": "",
            "review_date": "",
            "review_notes": "Imported from Excel manual compilation; source_url pending verification.",
            "approved_for_rag": False,
        },
    }


def missing_key_fields(case: dict) -> list[str]:
    fields = ["title", "case_number", "court", "decision_date", "facts_summary", "legal_basis", "penalty_result"]
    return [field for field in fields if not case.get(field)]


def write_report(report: dict, reports_dir: Path) -> None:
    lines = [
        "# Excel Import Report",
        "",
        f"- Rows read: {report['rows_read']}",
        f"- Categories detected: {report['category_count']}",
        f"- Successfully imported candidate cases: {report['imported_count']}",
        f"- New files written: {report['new_files_count']}",
        f"- Existing duplicates kept: {report['deduplicated_existing_count']}",
        f"- Skipped rows: {report['skipped_count']}",
        "",
        "## Imported By Violation Type",
        "",
    ]
    for violation_type, count in sorted(report["by_violation_type"].items()):
        lines.append(f"- {violation_type}: {count}")
    lines.extend(["", "## Missing Key Fields", ""])
    if report["missing_key_fields"]:
        lines.append("| case_id | title | missing_fields |")
        lines.append("| --- | --- | --- |")
        for item in report["missing_key_fields"]:
            lines.append(f"| {item['case_id']} | {item['title']} | {', '.join(item['missing_fields'])} |")
    else:
        lines.append("- None")
    lines.append("")
    (reports_dir / "excel_import_report.md").write_text("\n".join(lines), encoding="utf-8")


def write_lookup_todo(cases: list[dict], reports_dir: Path) -> None:
    path = reports_dir / "source_lookup_todo.csv"
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "case_id",
                "title",
                "case_number",
                "court",
                "decision_date",
                "source_lookup_query",
                "status",
            ],
        )
        writer.writeheader()
        for case in cases:
            query = (
                f"{case.get('case_number') or ''} {case.get('title') or ''}".strip()
                or f"{case.get('case_number') or ''} {case.get('court') or ''} 广告法 行政处罚".strip()
            )
            writer.writerow(
                {
                    "case_id": case["case_id"],
                    "title": case.get("title") or "",
                    "case_number": case.get("case_number") or "",
                    "court": case.get("court") or "",
                    "decision_date": case.get("decision_date") or "",
                    "source_lookup_query": query,
                    "status": "pending_source_lookup",
                }
            )


def write_analysis(lines: list[str], analysis_dir: Path) -> None:
    if not lines:
        return
    analysis_dir.mkdir(parents=True, exist_ok=True)
    content = ["# 执法趋势分析与合规建议", "", *lines, ""]
    (analysis_dir / "enforcement_trends.md").write_text("\n".join(content), encoding="utf-8")


def run(
    input_path: Path = DEFAULT_INPUT_PATH,
    candidates_dir: Path = DEFAULT_CANDIDATES_DIR,
    reports_dir: Path = DEFAULT_REPORTS_DIR,
    analysis_dir: Path = DEFAULT_ANALYSIS_DIR,
) -> dict:
    candidates_dir.mkdir(parents=True, exist_ok=True)
    reports_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(input_path, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active

    current_category = None
    header_map: dict[str, int] | None = None
    analysis_lines: list[str] = []
    imported_cases: list[dict] = []
    missing_rows: list[dict] = []
    by_type: Counter[str] = Counter()
    detected_categories: set[str] = set()
    skipped_count = 0
    new_files_count = 0
    duplicate_count = 0
    in_analysis = False

    for row in ws.iter_rows(values_only=True):
        raw_values = list(row)
        values = [normalize_text(value) for value in raw_values]
        nonempty = [value for value in values if value]
        if not nonempty:
            skipped_count += 1
            continue

        first = nonempty[0]
        if first in STOP_SECTIONS:
            in_analysis = True
            analysis_lines.append(f"## {first}")
            skipped_count += 1
            continue
        if in_analysis:
            analysis_lines.append(" ".join(nonempty))
            skipped_count += 1
            continue

        category_match = CATEGORY_PATTERN.match(first)
        if category_match:
            current_category = normalize_violation_type(category_match.group(1))
            detected_categories.add(current_category)
            header_map = None
            skipped_count += 1
            continue

        if is_header_row(nonempty):
            header_map = {value: idx for idx, value in enumerate(values) if value in HEADERS}
            skipped_count += 1
            continue

        if not header_map:
            skipped_count += 1
            continue

        row_data = row_to_mapping(raw_values, header_map)
        title = normalize_text(row_data.get("案例名称"))
        facts = normalize_text(row_data.get("关键违法事实"))
        if not title or not facts:
            skipped_count += 1
            continue

        case = build_case(row_data, current_category)
        output_path = candidates_dir / f"{case['case_id']}.json"
        if output_path.exists():
            duplicate_count += 1
        else:
            output_path.write_text(json.dumps(case, ensure_ascii=False, indent=2), encoding="utf-8")
            new_files_count += 1
        imported_cases.append(case)
        by_type[normalize_violation_type(normalize_text(row_data.get("违法类型")) or current_category or "其他")] += 1
        missing = missing_key_fields(case)
        if missing:
            missing_rows.append({"case_id": case["case_id"], "title": case.get("title", ""), "missing_fields": missing})

    write_analysis(analysis_lines, analysis_dir)
    write_lookup_todo(imported_cases, reports_dir)
    report = {
        "rows_read": ws.max_row,
        "category_count": len(detected_categories),
        "imported_count": len(imported_cases),
        "new_files_count": new_files_count,
        "deduplicated_existing_count": duplicate_count,
        "skipped_count": skipped_count,
        "by_violation_type": dict(by_type),
        "missing_key_fields": missing_rows,
    }
    write_report(report, reports_dir)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Import manually compiled Excel candidate cases.")
    parser.add_argument("--input-path", type=Path, default=DEFAULT_INPUT_PATH)
    parser.add_argument("--candidates-dir", type=Path, default=DEFAULT_CANDIDATES_DIR)
    parser.add_argument("--reports-dir", type=Path, default=DEFAULT_REPORTS_DIR)
    parser.add_argument("--analysis-dir", type=Path, default=DEFAULT_ANALYSIS_DIR)
    args = parser.parse_args()
    report = run(args.input_path, args.candidates_dir, args.reports_dir, args.analysis_dir)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
