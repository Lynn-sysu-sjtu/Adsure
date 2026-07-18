import csv
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from src import build_chunks, import_sector_docx, test_retrieval, validate_cases


HEADERS = [
    "序号",
    "违法类型",
    "案例名称",
    "案号",
    "审理法院",
    "审理日期",
    "关键违法事实",
    "处罚依据",
    "处罚结果",
]


class ImportSectorDocxTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        for rel in [
            "data/imports",
            "data/structured",
            "data/structured_candidates",
            "data/structured_samples",
            "data/chunks",
            "data/reports",
            "data/analysis",
            "data/mappings",
            "prompts",
        ]:
            (self.root / rel).mkdir(parents=True, exist_ok=True)
        self.prompt_path = self.root / "prompts/clean_case_prompt.md"
        self.prompt_path.write_text(
            """
risk_dimensions 枚举：
- 游戏抽奖概率公示
- 盲盒概率虚假宣传
- 广告内容真实性
- 虚假宣传
- 游戏福利承诺
- 奖励获取条件不清楚
- 化妆品医疗化宣传
- 涉医疗宣传
- 医疗用语
- 功效夸大
- 保健品违规宣传
- 会销虚假宣传
- 疾病治疗功效宣传
- 老年人营销风险
- 普通食品疾病治疗功效宣传
- 食品标签违法
- 保健品功能虚假宣传

输出 JSON schema：
""",
            encoding="utf-8",
        )

    def tearDown(self):
        self.tmp.cleanup()

    def add_case_table(self, doc: Document, rows: list[list[str]]) -> None:
        table = doc.add_table(rows=1, cols=len(HEADERS))
        for idx, header in enumerate(HEADERS):
            table.rows[0].cells[idx].text = header
        for row_values in rows:
            row = table.add_row()
            for idx, value in enumerate(row_values):
                row.cells[idx].text = value

    def make_docx(self) -> Path:
        path = self.root / "data/imports/sector.docx"
        doc = Document()
        doc.add_paragraph("2023-2026年游戏、美妆、保健品三大领域违法广告行政处罚案例汇总")
        doc.add_paragraph("一、游戏领域违法广告案例（2023-2026年）")
        self.add_case_table(
            doc,
            [
                [
                    "1",
                    "盲盒概率虚假宣传",
                    "四川某某科技有限公司诉眉山市市场监督管理局行政处罚案",
                    "（2024）川14行终29号",
                    "眉山市中级人民法院",
                    "2024-04-15",
                    "盲盒经营者在APP和小程序主页面标识“商品概率经司法鉴定真实有效，请放心购买”，但鉴定范围未覆盖完整运行代码。",
                    "《反不正当竞争法》第八条第一款、第二十条第一款",
                    "眉山市市监局罚款700000元（二审维持）",
                ],
                [
                    "2",
                    "游戏福利码宣传争议",
                    "许某诉广州某公司等网络服务合同纠纷案",
                    "（2025）粤0192民初9279号",
                    "广州互联网法院",
                    "2025-11-20",
                    "游戏广告宣称输入福利码可领取“888888仙玉+经验丹+六阶仙器*8”，法院认为广告表达不够清楚但未构成欺诈。",
                    "《消费者权益保护法》第五十五条",
                    "驳回原告诉讼请求",
                ],
            ],
        )
        doc.add_paragraph("二、美妆领域违法广告案例（2023-2026年）")
        self.add_case_table(
            doc,
            [
                [
                    "1",
                    "化妆品虚假广告及医疗用语",
                    "绍兴市柯桥区市场监督管理局申请执行绍兴金乌玉农业科技有限公司行政处罚案",
                    "（2024）浙0603行审7号",
                    "绍兴市柯桥区人民法院",
                    "2024-01-18",
                    "在手册、挂纸、视频、微信公众号上宣传“镇静镇痛、活血化瘀”“杀菌、消炎、镇痛”等医疗用语。",
                    "《广告法》第十七条、第二十八条、第五十五条第一款",
                    "罚款500000元",
                ]
            ],
        )
        doc.add_paragraph("三、保健品领域违法广告案例（2023-2026年）")
        self.add_case_table(
            doc,
            [
                [
                    "1",
                    "保健品功能虚假宣传（与美妆案例1为同一案件）",
                    "绍兴市柯桥区市场监督管理局申请执行绍兴金乌玉农业科技有限公司行政处罚案",
                    "（2024）浙0603行审7号",
                    "绍兴市柯桥区人民法院",
                    "2024-01-18",
                    "在手册、挂纸、视频、微信公众号上宣传产品改善亚健康、提升免疫力，并使用“杀菌、消炎、镇痛”等医疗用语。",
                    "《广告法》第十七条、第二十八条、第五十五条第一款",
                    "罚款500000元",
                ],
                [
                    "2",
                    "会销保健品虚假宣传",
                    "宜兴市市场监督管理局申请执行宜兴市某某服务部行政处罚案",
                    "（2025）苏0282行审17号",
                    "宜兴市人民法院",
                    "2025-04-02",
                    "以会销形式销售“三七人参胶囊”，播放视频宣传“降血压血脂”“头疼脑梗”“心脑血管病”等治疗效果。",
                    "《反不正当竞争法》第八条第一款、第二十条第一款",
                    "罚款200000元",
                ],
                [
                    "3",
                    "食用农产品标签涉及疾病治疗功能",
                    "成都市青羊区市场监督管理局申请执行成都某有限公司行政处罚案",
                    "（2024）川0105行审11号",
                    "成都市青羊区人民法院",
                    "2024-04-03",
                    "包装标签宣传“预防心脑血管硬化，抑制肿瘤发生和生长等作用”。",
                    "《食品安全法》第七十一条第一款、第一百二十五条第一款第（二）项",
                    "没收违法所得5433.2元、罚款44424.4元",
                ],
            ],
        )
        doc.add_paragraph("四、三大领域执法趋势分析")
        doc.add_paragraph("（一）游戏领域：监管重点聚焦概率公示与广告内容真实性")
        doc.add_paragraph("1. 概率公示合规成为核心监管要求。")
        doc.add_paragraph("（二）美妆领域：重点打击医疗用语与虚假功效宣传")
        doc.add_paragraph("1. 医疗用语是美妆广告红线。")
        doc.add_paragraph("（三）保健品领域：会销和网络营销是重点监管对象")
        doc.add_paragraph("1. 会销模式是重点监管对象。")
        doc.add_paragraph("（四）合规建议")
        doc.add_paragraph("1. 游戏企业应建立概率公示合规审查机制。")
        doc.add_paragraph("2. 美妆企业应严格审查广告文案。")
        doc.add_paragraph("3. 保健品企业应避免会销虚假宣传。")
        doc.save(path)
        return path

    def test_import_recognizes_sectors_skips_analysis_and_handles_duplicates(self):
        report = import_sector_docx.run(
            input_path=self.make_docx(),
            candidates_dir=self.root / "data/structured_candidates",
            reports_dir=self.root / "data/reports",
            analysis_dir=self.root / "data/analysis",
            mappings_dir=self.root / "data/mappings",
        )

        self.assertEqual(report["imported_count"], 6)
        self.assertEqual(report["by_sector"], {"game": 2, "beauty": 1, "health": 3})
        self.assertEqual(report["duplicate_group_count"], 1)
        cases = [json.loads(p.read_text(encoding="utf-8")) for p in (self.root / "data/structured_candidates").glob("*.json")]
        self.assertEqual({case["sector_cn"] for case in cases}, {"游戏", "美妆", "保健品"})
        game_civil = next(case for case in cases if case["violation_type"] == "游戏福利码宣传争议")
        self.assertNotEqual(game_civil["case_nature"], "administrative_penalty")
        self.assertFalse(game_civil["is_admin_penalty_candidate"])
        duplicate_cases = [case for case in cases if case["case_number"] == "（2024）浙0603行审7号"]
        self.assertEqual(len(duplicate_cases), 2)
        self.assertEqual(len({case["duplicate_group_id"] for case in duplicate_cases}), 1)
        self.assertEqual(set(duplicate_cases[0]["related_sectors"]), {"beauty", "health"})
        trend_text = (self.root / "data/analysis/sector_enforcement_trends.md").read_text(encoding="utf-8")
        suggestion_text = (self.root / "data/analysis/sector_compliance_suggestions.md").read_text(encoding="utf-8")
        self.assertIn("三大领域执法趋势分析", trend_text)
        self.assertIn("游戏企业", suggestion_text)
        with (self.root / "data/reports/sector_source_lookup_todo.csv").open(encoding="utf-8") as f:
            self.assertEqual(len(list(csv.DictReader(f))), 6)

    def test_sector_candidates_validate_chunk_and_retrieve_without_production(self):
        import_sector_docx.run(
            input_path=self.make_docx(),
            candidates_dir=self.root / "data/structured_candidates",
            reports_dir=self.root / "data/reports",
            analysis_dir=self.root / "data/analysis",
            mappings_dir=self.root / "data/mappings",
        )
        validation = validate_cases.run(
            structured_dir=self.root / "data/structured",
            structured_candidates_dir=self.root / "data/structured_candidates",
            structured_samples_dir=self.root / "data/structured_samples",
            reports_dir=self.root / "data/reports",
            prompt_path=self.prompt_path,
        )
        paths = build_chunks.run(
            structured_dir=self.root / "data/structured",
            structured_candidates_dir=self.root / "data/structured_candidates",
            structured_samples_dir=self.root / "data/structured_samples",
            chunks_dir=self.root / "data/chunks",
            reports_dir=self.root / "data/reports",
        )
        sector_chunks = json.loads(paths["sector_candidate"].read_text(encoding="utf-8"))
        production_chunks = json.loads(paths["production"].read_text(encoding="utf-8"))

        self.assertEqual(validation["sector_candidate_count"], 6)
        self.assertEqual(validation["invalid_count"], 0)
        self.assertGreaterEqual(len(sector_chunks), 12)
        self.assertEqual(production_chunks, [])
        for query, expected_sector in [
            ("游戏 抽奖 概率 公示 虚假宣传", "game"),
            ("化妆品 宣传 杀菌 消炎 医疗用语", "beauty"),
            ("保健品 会销 降血压 心脑血管", "health"),
            ("普通食品 包装 宣传 抑制肿瘤", "health"),
            ("游戏 福利码 奖励 条件不清楚", "game"),
        ]:
            result = test_retrieval.search(query=query, chunks_path=paths["sector_candidate"], top_k=1)[0]
            self.assertEqual(result["sector"], expected_sector)

        mappings = (self.root / "data/mappings/sector_case_rule_candidates.jsonl").read_text(encoding="utf-8").splitlines()
        self.assertGreaterEqual(len(mappings), 6)


if __name__ == "__main__":
    unittest.main()
