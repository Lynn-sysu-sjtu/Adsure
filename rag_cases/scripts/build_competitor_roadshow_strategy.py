from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from build_defense_qa import (
    BLUE,
    DARK_BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    PALE_GOLD,
    add_callout,
    add_markdown_inline,
    add_numbered_paragraph,
    add_page_number,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_row_cant_split,
    set_run_font,
    set_table_geometry,
    setup_styles,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "参赛提交材料" / "Adsure_竞品差异与路演制胜策略.md"
OUTPUT = ROOT / "参赛提交材料" / "Adsure_竞品差异与路演制胜策略.docx"
RED = RGBColor(155, 28, 28)


def setup_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("ADSURE  |  复赛路演制胜策略")
    set_run_font(run, size=9, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def table_widths(headers):
    if len(headers) == 4:
        if headers[0] == "评委关注点":
            return [1500, 2300, 2350, 3210]
        if headers[0] == "页码":
            return [720, 2300, 3340, 3000]
        return [1500, 2400, 2400, 3060]
    if len(headers) == 3:
        return [1600, 3900, 3860]
    return [9360 // len(headers)] * len(headers)


def add_table(doc, rows):
    headers = rows[0]
    body = rows[2:] if len(rows) > 1 and all(set(cell) <= {"-", ":"} for cell in rows[1]) else rows[1:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    widths = table_widths(headers)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.2, bold=True, color=DARK_BLUE)
    for row_data in body:
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for index, text in enumerate(row_data):
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (headers[0] == "页码" and index == 0) else WD_ALIGN_PARAGRAPH.LEFT
            add_markdown_inline(p, text)
            for run in p.runs:
                run.font.size = Pt(8.8)
            set_cell_margins(cells[index], top=90, bottom=90, start=100, end=100)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def parse_markdown(doc, lines):
    in_table = False
    rows = []

    def flush_table():
        nonlocal in_table, rows
        if rows:
            add_table(doc, rows)
        in_table = False
        rows = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|") and line.endswith("|"):
            in_table = True
            rows.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        if in_table:
            flush_table()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
            continue
        if line.startswith("> "):
            add_callout(doc, "核心口径", line[2:].replace("**", ""), fill=PALE_GOLD)
            continue
        if re.match(r"^\d+\. ", line):
            number = int(line.split(".", 1)[0])
            add_numbered_paragraph(doc, re.sub(r"^\d+\. ", "", line), number)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(27)
            p.paragraph_format.first_line_indent = Pt(-13.5)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            add_markdown_inline(p, line[2:])
            continue
        p = doc.add_paragraph()
        add_markdown_inline(p, line)
    flush_table()


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_styles(doc)
    setup_header_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Adsure 竞品差异与路演制胜策略")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("对标广审通 LexAd｜从功能展示升级为证据型叙事")

    meta = doc.add_table(rows=1, cols=3)
    set_table_geometry(meta, [3120, 3120, 3120])
    for index, value in enumerate(("复赛内部作战稿", "2026 年 7 月 22 日", "定位 · 演示 · 答辩")):
        cell = meta.cell(0, index)
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    parse_markdown(doc, lines[1:])

    doc.core_properties.title = "Adsure 竞品差异与路演制胜策略"
    doc.core_properties.subject = "对标广审通 LexAd 的复赛路演、演示和答辩作战稿"
    doc.core_properties.author = "Adsure 项目组"
    doc.core_properties.keywords = "Adsure, 广告合规, 竞品分析, 路演, LexAd"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
