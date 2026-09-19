from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "参赛提交材料" / "Adsure_开发部分答辩Q&A.md"
OUTPUT = ROOT / "参赛提交材料" / "Adsure_开发部分答辩Q&A.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(36, 46, 56)
MUTED = RGBColor(92, 101, 110)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF7E0"
BORDER = "B8C4D0"


def set_run_font(run, size=11, bold=None, color=INK, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    settings = {
        "Title": (28, DARK_BLUE, 0, 8),
        "Subtitle": (13, MUTED, 0, 18),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def setup_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("ADSURE  |  开发答辩预备")
    set_run_font(run, size=9, bold=True, color=MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}：")
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_markdown_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=10, color=DARK_BLUE)
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def create_numbering_instance(doc, start_value):
    style = doc.styles["List Number"]
    style_num_id = style._element.pPr.numPr.numId.val
    numbering = doc.part.numbering_part.element
    base_num = numbering.find(f".//{{{qn('w:num').split('}')[0][1:]}}}num[@{{{qn('w:numId').split('}')[0][1:]}}}numId='{style_num_id}']")
    if base_num is None:
        base_num = next(iter(numbering.findall(qn("w:num"))))
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), str(start_value))
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def add_numbered_paragraph(doc, text, start_value):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = create_numbering_instance(doc, start_value)
    add_markdown_inline(p, text)


def parse_markdown(doc, lines):
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            return
        headers = table_rows[0]
        body = table_rows[2:] if len(table_rows) > 1 and all(set(cell) <= {"-", ":"} for cell in table_rows[1]) else table_rows[1:]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        widths = [2500, 1700, 5160] if len(headers) == 3 else [9360 // len(headers)] * len(headers)
        set_table_geometry(table, widths)
        set_repeat_table_header(table.rows[0])
        for index, text in enumerate(headers):
            cell = table.rows[0].cells[index]
            set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, size=9.5, bold=True, color=DARK_BLUE)
        for row_data in body:
            cells = table.add_row().cells
            for index, text in enumerate(row_data):
                p = cells[index].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
                add_markdown_inline(p, text)
                for run in p.runs:
                    run.font.size = Pt(9.5)
        set_table_geometry(table, widths)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        table_rows = []
        in_table = False

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|") and line.endswith("|"):
            in_table = True
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        if in_table:
            flush_table()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            continue
        if line.startswith("### "):
            title = line[4:]
            p = doc.add_heading(title, level=2)
            if title.startswith("Q"):
                p.paragraph_format.keep_with_next = True
            continue
        if line.startswith("> "):
            add_callout(doc, "使用提示", line[2:], fill=PALE_GOLD)
            continue
        if re.match(r"^\d+\. ", line):
            number = int(line.split(".", 1)[0])
            add_numbered_paragraph(doc, re.sub(r"^\d+\. ", "", line), number)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.188)
            p.paragraph_format.space_after = Pt(4)
            add_markdown_inline(p, line[2:])
            continue
        if line.startswith("**建议回答：**"):
            add_callout(doc, "建议回答", line[len("**建议回答：**"):].strip())
            continue
        if line.startswith("**必须同时说明：**"):
            add_callout(doc, "必须同时说明", line[len("**必须同时说明：**"):].strip(), fill=PALE_GOLD)
            continue
        p = doc.add_paragraph()
        add_markdown_inline(p, line)
    flush_table()


def build():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    setup_styles(doc)
    setup_header_footer(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Adsure 广告合规 AI")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("开发部分答辩 Q&A｜可直接背诵与追问展开")

    meta = doc.add_table(rows=1, cols=3)
    set_table_geometry(meta, [3120, 3120, 3120])
    for index, value in enumerate(("答辩预备稿 v1.1", "2026 年 7 月 23 日", "案例库 · RAG · 竞品对打")):
        cell = meta.cell(0, index)
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    parse_markdown(doc, lines[1:])

    doc.core_properties.title = "Adsure 广告合规 AI 开发部分答辩 Q&A"
    doc.core_properties.subject = "案例库、RAG、接口、安全、测试与部署答辩预备"
    doc.core_properties.author = "Adsure 项目组"
    doc.core_properties.keywords = "Adsure, 广告合规, RAG, 答辩, Q&A"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
