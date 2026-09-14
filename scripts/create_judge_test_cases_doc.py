from pathlib import Path
import json

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "评委测试用例-广告合规AI案例库RAG.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9E2F3", sz="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:color"), color)


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    node.set(qn("w:val"), "1" if value else "0")


def add_run(paragraph, text, bold=False, color=None, size=None, italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial Unicode MS"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        add_run(p, bold_prefix, bold=True, color="17365D")
        add_run(p, text[len(bold_prefix):])
    else:
        add_run(p, text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    add_run(p, text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    add_run(p, text)
    return p


def add_callout(doc, label, text, fill="EAF2F8"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(16.4)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="9FBAD0", sz="8")
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    add_run(p, label + "  ", bold=True, color="17365D")
    add_run(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], "F3F6FA")
        p0 = cells[0].paragraphs[0]
        add_run(p0, key, bold=True, color="17365D", size=9.5)
        p1 = cells[1].paragraphs[0]
        add_run(p1, value, size=9.5)
    set_table_widths(table, [3.3, 13.1])
    return table


def add_output_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    for c, text in zip(hdr, ("输出字段", "期望值")):
        set_cell_shading(c, "17365D")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, text, bold=True, color="FFFFFF", size=9.5)
    set_repeat_table_header(table.rows[0])
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], "F3F6FA")
        add_run(cells[0].paragraphs[0], key, bold=True, color="17365D", size=9)
        add_run(cells[1].paragraphs[0], value, size=9)
    set_table_widths(table, [4.0, 12.4])
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in [
        ("Title", 24, "17365D", 0, 10),
        ("Heading 1", 16, "17365D", 14, 6),
        ("Heading 2", 12.5, "2F5597", 10, 4),
        ("Heading 3", 11, "17365D", 8, 3),
    ]:
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[list_name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.1


def add_footer(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, "广告合规 AI｜评委测试用例｜仅供测试验收", color="6B7280", size=8.5)


def add_case(doc, number, title, background, input_material, operation, expected, source_note, acceptance):
    add_heading(doc, f"案例 {number}｜{title}", level=1)
    add_callout(doc, "测试目标", background)
    add_heading(doc, "一、输入材料", level=2)
    add_key_value_table(doc, input_material)
    add_heading(doc, "二、预期操作", level=2)
    for step in operation:
        add_number(doc, step)
    add_heading(doc, "三、期望输出", level=2)
    add_body(doc, expected["summary"])
    add_output_table(doc, expected["fields"])
    add_heading(doc, "四、验收要点", level=2)
    for item in acceptance:
        add_bullet(doc, item)
    add_callout(doc, "数据背景", source_note, fill="FFF7E6")


def main():
    doc = Document()
    configure_styles(doc)

    # Title block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, "广告合规 AI 案例库 RAG", bold=True, color="17365D", size=24)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    add_run(p, "评委测试用例提交材料（3 个完整案例）", color="2F5597", size=15)
    add_callout(doc, "提交说明", "本文档可直接发送给评委。评委可依次录入三个案例的输入材料，按照预期操作执行，并以期望输出和验收要点判断案例库 RAG 是否正确完成召回与展示。")

    add_heading(doc, "一、测试背景与统一口径", level=1)
    add_body(doc, "本次测试覆盖三类典型广告合规场景：普通食品/保健食品疾病功效宣传、绝对化用语与国家机关工作人员形象、房地产升值承诺及规划配套误导。测试重点是验证案例库能否根据物料场景召回具有监管参考价值的公开处罚案例，并返回可回溯的来源和法律依据边界。")
    add_body(doc, "案例库 RAG 的职责是提供相似监管场景、处罚事实、监管逻辑和来源，不直接替代规则引擎判断产品事实，也不直接生成最终法律结论。")
    add_key_value_table(doc, [
        ("建议测试接口", "POST /cases/retrieve"),
        ("索引范围", "production；只返回已核验、批准进入 RAG 的正式案例"),
        ("建议检索模式", "hybrid；若评委只支持词法检索，使用 lexical 也应命中首条案例"),
        ("统一 top_k", "3"),
        ("结果边界", "risk_level 不根据相似度推导；具体法条若为推定映射，必须保留待法律复核标识"),
    ])
    add_heading(doc, "统一验收标准", level=2)
    for item in [
        "首条召回案例的 case_id 与本材料一致，且同一案例不重复展示。",
        "结果包含 title、violation_type、risk_dimensions、ruling、legal_basis、regulatory_logic。",
        "结果包含 source_url、raw_text_path、review_status、source_verification_status 和 candidate_data。",
        "正式案例返回 candidate_data=false；不得用候选案例静默补位。",
        "案例处罚事实可以作为监管参考，但不能被表述为当前待审物料已经构成同等违法。",
    ]:
        add_bullet(doc, item)

    add_case(
        doc, 1, "普通食品/保健食品疾病功效宣传（可命中类案）",
        "验证保健食品场景能够跨兼容行业召回普通食品疾病功效宣传类案。本案例是本组中可直接验证实际命中的案例。",
        [
            ("案例编号", "TC-01"),
            ("行业", "保健食品"),
            ("产品品类", "普通食品"),
            ("投放平台", "抖音、线下会销（测试请求可传空数组）"),
            ("广告内容", "驼乳粉改善肺结节、降血糖、降脂降压"),
            ("claim_spans", "改善肺结节；降血糖；降脂降压"),
            ("risk_dimensions", "普通食品疾病治疗功效宣传"),
            ("matched_rule_ids", "ADLAW-017"),
            ("top_k", "3"),
        ],
        [
            "将输入材料转换为标准 JSON，请求 POST /cases/retrieve。",
            "industry 传入“保健食品”，product_category 传入“普通食品”，claim_spans 和 matched_rule_ids 按上表传入。",
            "使用 production 索引执行混合检索，并读取返回的第一条案例卡片。",
            "核对案例原始行业、监管逻辑、处罚结果及 source_url，不把“保健食品”请求标签改写成案例原始行业。",
        ],
        {
            "summary": "预期首条召回为市场监管总局公开的那拉尊驼案。该案原始行业为“普通食品”，但保健食品请求允许召回普通食品功效宣传类案，因此属于行业兼容召回。",
            "fields": [
                ("case_id", "samr_2025_typical_ads_02"),
                ("title", "广州市海珠区市场监管局查处那拉尊驼（广州）乳业有限公司违法广告案"),
                ("violation_type", "普通食品疾病治疗功效及虚假宣传"),
                ("risk_dimensions", "虚假宣传；普通食品疾病治疗功效宣传；广告引证内容不规范"),
                ("content_snippet", "改善肺结节、降血糖、降脂降压等原案例宣称"),
                ("regulatory_logic", "普通食品广告不得宣称疾病预防、治疗功效；虚构项目和优惠名额、引证内容不真实准确也具有合规风险。"),
                ("ruling", "罚款45万元"),
                ("mapped_rule_ids", "ADLAW-011-02；ADLAW-017；ADLAW-028"),
                ("legal_basis", "《中华人民共和国广告法》有关规定"),
                ("source / status", "市场监管总局；source_verified；approved；candidate_data=false"),
            ],
        },
        "本案来源为市场监管总局公开典型案例页面，原始案例行业标记为“普通食品”。具体条款映射来自官方摘要事实的适用性推定，法律依据详情中的 mapping_review_status 应保持 pending_legal_review。",
        [
            "首条 case_id 为 samr_2025_typical_ads_02。",
            "返回行业应保留为普通食品，不能因为请求行业为保健食品而篡改。",
            "不得输出“该待审广告必然违法”或将45万元处罚直接推导为当前物料处罚金额。",
        ],
    )

    add_case(
        doc, 2, "绝对化用语与国家机关工作人员形象（预期类案输出）",
        "验证系统在当前实际无法稳定生成完整业务答案时，仍能提供明确的 RAG 目标输出：命中相关正式案例，并返回可供上游审核模型使用的监管逻辑和规则映射。",
        [
            ("案例编号", "TC-02"),
            ("行业", "通用"),
            ("产品品类", "普通食品或未填写"),
            ("投放平台", "未指定"),
            ("广告内容", "世界最强，使用国家机关工作人员形象"),
            ("claim_spans", "世界最强；国家机关工作人员形象"),
            ("risk_dimensions", "绝对化用语"),
            ("matched_rule_ids", "ADLAW-009-03"),
            ("top_k", "3"),
        ],
        [
            "将广告内容和风险维度作为标准检索请求发送至 POST /cases/retrieve。",
            "industry 传入“通用”，不强制要求上游先确认产品属于普通食品。",
            "优先观察第一条案例是否为 samr_2025_typical_ads_05，并核对结果中绝对化用语与国家机关形象两类监管逻辑。",
            "若页面暂时不能生成最终审核结论，应展示“相关案例参考”卡片，而不是返回空泛的“无法判断”。",
        ],
        {
            "summary": "预期首条召回为重庆林晖泰案。该案同时覆盖普通食品疾病功效宣传、虚假权威背书、绝对化用语以及国家机关及工作人员名义或形象等风险。",
            "fields": [
                ("case_id", "samr_2025_typical_ads_05"),
                ("title", "重庆市两江新区市场监管局查处重庆林晖泰商务信息咨询有限公司违法广告案"),
                ("violation_type", "普通食品疾病治疗功效、虚假宣传及绝对化用语"),
                ("risk_dimensions", "虚假宣传；普通食品疾病治疗功效宣传；绝对化用语"),
                ("content_snippet", "世界最强；国宝级密药；诺贝尔奖得主背书；治疗癌症、心脏病和高血压等"),
                ("regulatory_logic", "不得使用“世界最强”等绝对化用语，不得使用国家机关及工作人员名义或形象，也不得虚构权威背书和疾病治疗功效。"),
                ("ruling", "罚款16万元"),
                ("mapped_rule_ids", "ADLAW-009-02；ADLAW-009-03；ADLAW-017；ADLAW-028"),
                ("legal_basis", "《中华人民共和国广告法》有关规定"),
                ("source / status", "市场监管总局；source_verified；approved；candidate_data=false"),
            ],
        },
        "该案例作为评委测试的目标类案输出。RAG 只说明待审内容与历史监管场景相似，不直接确认当前物料中的“国家机关工作人员形象”是否真实存在、是否构成违法。",
        [
            "首条 case_id 为 samr_2025_typical_ads_05。",
            "输出至少覆盖 ADLAW-009-03；若展示完整案例映射，还应保留 ADLAW-009-02、ADLAW-017、ADLAW-028。",
            "risk_level 应为 null；score 只能作为检索排序分，不能展示为风险百分比。",
        ],
    )

    add_case(
        doc, 3, "房地产升值承诺与规划配套误导（预期类案输出）",
        "验证系统在无法直接输出最终法律意见时，能否提供房地产监管场景的目标类案、规则关联和处罚参考，供上游审核链路继续生成答案。",
        [
            ("案例编号", "TC-03"),
            ("行业", "通用"),
            ("产品品类", "房地产"),
            ("投放平台", "未指定"),
            ("广告内容", "楼盘承诺升值，宣传未纳入规划的交通和医疗配套"),
            ("claim_spans", "承诺升值；未纳入规划的交通和医疗配套"),
            ("risk_dimensions", "房地产广告误导；升值承诺"),
            ("matched_rule_ids", "ADLAW-026-01；ADLAW-026-04"),
            ("top_k", "3"),
        ],
        [
            "将房地产广告内容转换为标准 JSON，发送至 POST /cases/retrieve。",
            "industry 传入“通用”，product_category 传入“房地产”，risk_dimensions 和 matched_rule_ids 按上表传入。",
            "执行 production 检索并读取第一条案例卡片；不要求 RAG 直接给出最终风险等级。",
            "将召回结果中的两类规则关联分别展示：升值承诺、规划或建设中的配套设施误导。",
        ],
        {
            "summary": "预期首条召回为安徽康桥置业案。该案与输入材料在“房地产广告、未纳入规划的配套设施、升值承诺”三个关键场景上高度一致。",
            "fields": [
                ("case_id", "samr_2025_typical_ads_07"),
                ("title", "安徽省六安市市场监管局查处安徽康桥置业有限公司违法广告案"),
                ("violation_type", "房地产广告误导及升值承诺"),
                ("risk_dimensions", "房地产广告误导；升值承诺"),
                ("content_snippet", "将未纳入规划的交通、商业、医疗等设施作为宣传内容；升值承诺"),
                ("regulatory_logic", "房地产广告不得把未纳入规划的配套设施作为确定宣传内容，也不得作出升值承诺误导购房者。"),
                ("ruling", "罚款15万元"),
                ("mapped_rule_ids", "ADLAW-026-01；ADLAW-026-04"),
                ("legal_basis", "《中华人民共和国广告法》有关规定"),
                ("source / status", "市场监管总局；source_verified；approved；candidate_data=false"),
            ],
        },
        "该案例来源为市场监管总局公开典型案例页面。具体条款及责任条款映射均应保留来源证据边界，不得将类案处罚金额直接作为当前房地产广告的预期处罚。",
        [
            "首条 case_id 为 samr_2025_typical_ads_07。",
            "ADLAW-026-01 和 ADLAW-026-04 均应保留，不应只返回其中一项。",
            "RAG 输出应说明“可参考该监管场景”，而不是直接代替对项目规划真实性和广告证明材料的核验。",
        ],
    )

    add_heading(doc, "附录：评委可直接复制的测试请求", level=1)
    add_body(doc, "以下请求体可分别用于三个案例。若接口地址、鉴权方式或字段名由现场环境提供，以现场接口文档为准；本材料只冻结测试业务输入和预期结果。")
    requests = [
        '{"content":"驼乳粉改善肺结节、降血糖、降脂降压","industry":"保健食品","platform":["抖音"],"claim_spans":["改善肺结节","降血糖","降脂降压"],"product_category":"普通食品","risk_dimensions":["普通食品疾病治疗功效宣传"],"matched_rule_ids":["ADLAW-017"],"top_k":3}',
        '{"content":"世界最强，使用国家机关工作人员形象","industry":"通用","platform":[],"claim_spans":["世界最强","国家机关工作人员形象"],"product_category":"普通食品","risk_dimensions":["绝对化用语"],"matched_rule_ids":["ADLAW-009-03"],"top_k":3}',
        '{"content":"楼盘承诺升值，宣传未纳入规划的交通和医疗配套","industry":"通用","platform":[],"claim_spans":["承诺升值","未纳入规划的交通和医疗配套"],"product_category":"房地产","risk_dimensions":["房地产广告误导","升值承诺"],"matched_rule_ids":["ADLAW-026-01","ADLAW-026-04"],"top_k":3}',
    ]
    for idx, request in enumerate(requests, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.right_indent = Cm(0.3)
        p.paragraph_format.space_after = Pt(7)
        add_run(p, f"案例 {idx} 请求体：\n", bold=True, color="17365D", size=9)
        add_run(p, request, size=8.5)

    add_heading(doc, "最终说明", level=1)
    add_body(doc, "三个案例均使用已进入 production RAG 的公开典型案例作为目标输出。案例中的处罚事实、来源链接和原文路径用于人工回溯；法律条款映射中的推定部分仍需法务复核。评委验收时，应重点关注召回正确性、字段完整性、来源可追溯性和事实边界，而不是要求案例库直接承担完整的最终审核意见生成。")

    add_footer(doc)
    doc.core_properties.title = "广告合规 AI 案例库 RAG 评委测试用例"
    doc.core_properties.subject = "三个完整提交案例：输入材料、预期操作和期望输出"
    doc.core_properties.author = "广告合规 AI"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
