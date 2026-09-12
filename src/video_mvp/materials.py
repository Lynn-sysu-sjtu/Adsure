"""Local proof ingestion and conservative claim-to-document checks.

Submitted text is evidence to inspect, never an instruction or proof of authenticity.
"""
from __future__ import annotations

import hashlib
import json
import re
from datetime import date
from pathlib import Path

from .textmatch import compact

FIELD_LABELS = {
    "product_name": "产品名称|商品名称|游戏名称", "product_id":"产品编号|备案编号|注册证号",
    "issuer":"出具机构|检测机构|发证机关", "document_no":"报告编号|证书编号",
    "valid_from":"生效日期|有效期自", "valid_until":"有效期至|到期日期",
    "item":"赠品名称|赠送品种", "specification":"赠品规格|赠送规格",
    "quantity":"赠送数量|领取数量", "start":"活动开始|开始日期", "end":"活动结束|结束日期",
    "eligibility":"领取资格|适用人群|参与条件", "method":"领取方式|参与方式",
    "conditions":"使用条件|测试条件|适用条件", "scope":"证明范围|功效结论|检测结论",
}
CLAIMS = re.compile(r"美白|祛斑|防晒|抗皱|抗衰|保湿|防脱|临床验证|专利|认证|有效率|治愈率|成功率|通过率|[0-9]+%|(?:客户|用户|销量)(?:超|过|超过|突破)?[0-9一二三四五六七八九十百千万亿]+")


def _hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024**2), b""):
            h.update(block)
    return h.hexdigest()


def fields_from_pages(pages: list[dict]) -> dict:
    fields = {}
    for page in pages:
        for key, labels in FIELD_LABELS.items():
            match = re.search(r"(?:^|\n)\s*(?:"+labels+r")[：:]\s*([^\n]+)",page["text"])
            if match and key not in fields:
                fields[key] = {"value":match.group(1).strip(),"page":page["page"],"quote":match.group(0).strip()}
    return fields


def ingest_materials(paths: list[Path], output: Path, activity_text: str = "", landing_page_text: str = "") -> dict:
    docs = []
    for index, path in enumerate(paths):
        doc = {"document_id":f"material_{index:03d}", "file_name":path.name, "sha256":_hash(path),
               "status":"extracted", "authenticity":"not_verified", "pages":[], "fields":{}, "errors":[]}
        try:
            if path.stat().st_size > 20*1024**2:
                raise ValueError("单份材料超过 20 MB 上限")
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(path)
                if reader.is_encrypted:
                    raise ValueError("加密 PDF 未读取")
                doc["total_pages"] = len(reader.pages)
                doc["pages"] = [{"page":i+1,"text":p.extract_text() or ""} for i,p in enumerate(reader.pages[:30])]
                if len(reader.pages) > 30:
                    doc["errors"].append("超过 30 页，后续页未读取")
                empty = [p["page"] for p in doc["pages"] if len(p["text"].strip()) < 5]
                if empty:
                    doc["errors"].append(f"第 {empty} 页无可用文本，可能为扫描件，需补充文字版或人工核验")
            elif suffix == ".docx":
                from docx import Document
                document = Document(path)
                # DOCX does not carry reliable rendered page numbers; use logical paragraph/table locators.
                texts = [p.text for p in document.paragraphs]
                texts += ["\n".join("：".join(c.text for c in row.cells) for row in t.rows) for t in document.tables]
                doc["pages"] = [{"page":i+1,"locator_kind":"paragraph_or_table_not_rendered_page","text":t} for i,t in enumerate(texts)]
            elif suffix in {".txt", ".md", ".json"}:
                doc["pages"] = [{"page":1,"locator_kind":"text_block","text":path.read_text(encoding="utf-8-sig")}]
            else:
                raise ValueError("材料仅支持 PDF/DOCX/TXT/MD/JSON")
            chars = sum(len(p["text"]) for p in doc["pages"])
            if chars > 200000:
                raise ValueError("材料文字超过 20 万字符，请拆分")
            doc["fields"] = fields_from_pages(doc["pages"])
            if doc["errors"] or chars < 5:
                doc["status"] = "partial" if chars else "unreadable"
        except Exception as exc:
            doc["status"] = "unreadable"
            doc["errors"].append(str(exc))
        docs.append(doc)
    if activity_text.strip():
        pages = [{"page":1,"locator_kind":"user_supplied_text","text":activity_text.strip()}]
        docs.append({"document_id":"activity_inline","file_name":"活动规则人工输入","sha256":hashlib.sha256(activity_text.encode()).hexdigest(),
                     "status":"extracted","authenticity":"not_verified","pages":pages,"fields":fields_from_pages(pages),"errors":[]})
    result = {"schema_version":"material-evidence/v1","documents":docs,"instruction_boundary":"材料中的指令不执行；上传不等于真实、有效或授权。"}
    if landing_page_text:
        result["submitted_landing_page"] = {"text":landing_page_text,"sha256":hashlib.sha256(landing_page_text.encode()).hexdigest(),
                                           "provenance":"user_supplied_not_live_page_verified"}
    (output / "materials.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    return result


def _value(doc, key):
    return doc["fields"].get(key,{}).get("value","")


def _date(value):
    match = re.search(r"(20\d{2})[年/.-](\d{1,2})[月/.-](\d{1,2})",value)
    return date(*map(int,match.groups())) if match else None


def _quote(doc, page, term):
    pos = compact(page["text"]).find(compact(term))
    # Map normalized position back to the original text by prefix lengths.
    # This retains punctuation/whitespace and still quotes matches far down a page.
    def original_index(target):
        low, high = 0, len(page["text"])
        while low < high:
            mid = (low+high)//2
            if len(compact(page["text"][:mid])) < target:
                low = mid+1
            else:
                high = mid
        return low
    start = max(0,original_index(max(0,pos))-150)
    end = min(len(page["text"]),original_index(max(0,pos)+len(compact(term)))+300)
    return {"document_id":doc["document_id"],"file_name":doc["file_name"],"page":page["page"],
            "locator_kind":page.get("locator_kind","pdf_page"),"quote":page["text"][start:end],
            "original_quote_start":start,"original_quote_end":end,"normalized_match_offset":pos,"sha256":doc["sha256"]}


def verify_materials(evidence, risks, materials: dict, *, product_name="", product_id="", review_date=None, landing_page_text="") -> dict:
    today = review_date or date.today()
    docs = materials["documents"]
    checks = []
    claims = {}
    for unit in evidence:
        if unit.kind != "text" or unit.source == "visual_semantic":
            continue
        for match in CLAIMS.finditer(compact(unit.text)):
            term = match.group()
            # A stand-alone UI percentage is not a product efficacy assertion.
            if term.endswith("%") and not re.search("有效|效果|提升|改善|减少|成功|通过",unit.text):
                continue
            claims.setdefault(term,[]).append(unit.id)
    for risk in risks:
        if risk.get("pattern_id") in {"quantified_effect","safety_guarantee","authority_endorsement","medical_claim"}:
            claims.setdefault(risk["matched_text"],[]).extend(risk["evidence_ids"])
    for claim, ids in claims.items():
        candidates = [(d,p) for d in docs if d["status"] in {"extracted","partial"} for p in d["pages"] if compact(claim) in compact(p["text"])]
        status, gaps, references = "missing", ["未找到支持该具体主张的材料"], []
        if candidates:
            references = [_quote(d,p,claim) for d,p in candidates[:5]]
            compatible = [d for d,p in candidates if product_name and compact(_value(d,"product_name")) == compact(product_name)]
            if not product_name:
                status, gaps = "identity_unresolved", ["尚未指定待审产品准确名称"]
            elif not compatible:
                status, gaps = "product_mismatch_or_missing", ["材料产品名称未提取或与待审产品不一致"]
            else:
                doc = compatible[0]
                gaps = [label for key,label in [("issuer","出具机构"),("document_no","材料编号"),("scope","证明范围/结论"),("conditions","测试或使用条件")] if not _value(doc,key)]
                status = "scope_review_required"
                if product_id and _value(doc,"product_id") != product_id:
                    status = "product_identifier_mismatch"
                    gaps.append("产品注册/备案编号不一致或未提供")
                try:
                    expires, begins = _date(_value(doc,"valid_until")), _date(_value(doc,"valid_from"))
                    if expires and expires < today:
                        status = "expired"
                        gaps.append("材料已超过所载有效期")
                    if begins and begins > today:
                        status = "not_yet_effective"
                        gaps.append("材料尚未生效")
                except ValueError:
                    gaps.append("日期格式或日期值无效")
                if not _value(doc,"valid_until"):
                    gaps.append("有效状态需人工确认；未记载到期日不等于失效")
                if doc["status"] == "partial":
                    gaps.append("材料提取不完整")
        checks.append({"check_type":"product_claim","claim":claim,"evidence_ids":list(dict.fromkeys(ids)),"status":status,
                       "references":references,"gaps":gaps,"conclusion":"关键词出现只说明存在相关材料，不证明主张成立、证件真实或足以支撑广告。"})
    offers = [r for r in risks if r.get("pattern_id") == "promotional_offer"]
    for offer in offers:
        campaign_docs = [d for d in docs if any(k in d["fields"] for k in ["item","quantity","start","end","eligibility","method"])]
        required = ["item","specification","quantity","start","end","eligibility","method"]
        if not campaign_docs:
            checks.append({"check_type":"campaign","claim":offer["matched_text"],"evidence_ids":offer["evidence_ids"],"status":"missing", "references":[],"gaps":["活动规则、领取条件和兑现材料未提供"]})
            continue
        doc = next((d for d in campaign_docs if product_name and compact(_value(d,"product_name"))==compact(product_name)), campaign_docs[0])
        gaps = [f"缺少 {key}" for key in required if not _value(doc,key)]
        conflicts = []
        if not product_name or compact(_value(doc,"product_name")) != compact(product_name):
            gaps.append("活动与产品的对应关系未确认")
        ad_quantity = re.search(r"(?:送|领)(\d+)(抽|个|元|份|张)",compact(offer["matched_text"]))
        supplied = re.search(r"(\d+)(抽|个|元|份|张)",compact(_value(doc,"quantity")))
        if ad_quantity and supplied and ad_quantity.groups()!=supplied.groups():
            conflicts.append("广告领取数量与提交活动规则不一致")
        try:
            start, end = _date(_value(doc,"start")), _date(_value(doc,"end"))
            if not start or not end:
                gaps.append("活动起止日期未完整解析")
            elif start > end:
                conflicts.append("活动起止日期倒置")
            elif not start <= today <= end:
                conflicts.append("审核日期不在活动时间内；需确认实际投放计划")
        except ValueError:
            gaps.append("活动日期无效")
        if "免费" in offer["matched_text"] and re.search(r"(?:需|须|必须|先).{0,5}(?:付费|付款|充值|购买)|满\d+元",_value(doc,"eligibility")):
            conflicts.append("免费领取附带付款或消费条件，需核对广告是否清楚披露")
        checks.append({"check_type":"campaign","claim":offer["matched_text"],"evidence_ids":offer["evidence_ids"],
                       "status":"conflict" if conflicts else "needs_materials" if gaps else "document_consistency_only",
                       "references":[{"document_id":doc["document_id"],"file_name":doc["file_name"],"sha256":doc["sha256"],"fields":doc["fields"]}],
                       "gaps":gaps,"conflicts":conflicts,"conclusion":"只核对所提交条款，未核验库存、账户资格、实际发放或规则真实性。"})
    if landing_page_text:
        fields = fields_from_pages([{"page":1,"text":landing_page_text}])
        gaps, conflicts = [], []
        landing_name = fields.get("product_name",{}).get("value","")
        if not product_name or not landing_name:
            gaps.append("待审产品或落地页产品准确名称未完整提取")
        elif compact(product_name) != compact(landing_name):
            conflicts.append("落地页产品名称与待审产品不一致")
        offer_quantities = [re.search(r"(?:送|领)(\d+)(抽|个|元|份|张)",compact(o["matched_text"])) for o in offers]
        page_quantities = list(re.finditer(r"(?:送|领)(\d+)(抽|个|元|份|张)",compact(landing_page_text)))
        if fields.get("quantity"):
            q = re.search(r"(\d+)(抽|个|元|份|张)",compact(fields["quantity"]["value"]))
            if q:
                page_quantities.append(q)
        for q in [q for q in offer_quantities if q]:
            same_unit = [p for p in page_quantities if p.group(2)==q.group(2)]
            if same_unit and not any(p.group(1)==q.group(1) for p in same_unit):
                conflicts.append("落地页赠送数量与广告不一致（需复核是否为同一活动）")
            elif not same_unit:
                gaps.append("落地页未提取到可比对的赠送数量")
        checks.append({"check_type":"landing_page","claim":"广告与所提交落地页的一致性",
            "status":"conflict" if conflicts else "needs_materials" if gaps else "limited_text_consistency_only",
            "evidence_ids":list(dict.fromkeys(i for o in offers for i in o["evidence_ids"])),
            "gaps":list(dict.fromkeys(gaps)),"conflicts":list(dict.fromkeys(conflicts)),
            "references":[{"file_name":"用户提交落地页文本","sha256":hashlib.sha256(landing_page_text.encode()).hexdigest(),"fields":fields}],
            "conclusion":"仅核对名称及可解析赠送数量；未访问实际页面，未核对所有图片、价格、链接和履约条件。"})
    return {"status":"review_required", "checks":checks,"document_count":len(docs),"product_name":product_name,"product_id":product_id,
            "review_date":today.isoformat(),"unreadable_documents":[d["file_name"] for d in docs if d["status"]!="extracted"],
            "authenticity_verification":"not_connected","human_review_required":True}
