# -*- coding: utf-8 -*-

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from export_parallel_rule_review_workbench_v01 import (
    build_parallel_review_rows,
    export_workbench,
)


ROOT = Path(__file__).resolve().parents[1]


class ParallelRuleReviewWorkbenchTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.result = build_parallel_review_rows(ROOT)
        cls.rows = cls.result["rows"]

    def test_each_uid_is_one_task(self):
        uids = [row["规则UID"] for row in self.rows]
        self.assertEqual(len(uids), len(set(uids)))
        self.assertGreater(len(uids), 100)

    def test_only_pending_review_rules_are_included(self):
        self.assertTrue(all(row["风险标签"] for row in self.rows))
        self.assertTrue(all(row["规则原文"] for row in self.rows))
        self.assertFalse(
            self.result["removed_uids"] & {row["规则UID"] for row in self.rows}
        )

    def test_task_packages_are_balanced_by_workload(self):
        scores = self.result["package_scores"]
        self.assertEqual(set(scores), {"任务包A", "任务包B", "任务包C"})
        self.assertLessEqual(max(scores.values()) - min(scores.values()), 4)

    def test_workload_score_matches_risk_tags(self):
        weights = {
            "跨目录重叠": 3,
            "问题错配": 1,
            "主动补资料": 2,
            "移出普通文案树": 2,
            "中低置信度": 1,
        }
        for row in self.rows:
            expected = sum(weights[tag] for tag in row["风险标签"].split("、"))
            self.assertEqual(row["工作量积分"], expected)

    def test_export_has_parallel_review_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "workbench.xlsx"
            export_workbench(self.result, output)
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(workbook.sheetnames[1], "并行审核主表")
            self.assertIn("任务包A", workbook.sheetnames)
            self.assertIn("任务包B", workbook.sheetnames)
            self.assertIn("任务包C", workbook.sheetnames)
            self.assertIn("争议池", workbook.sheetnames)
            self.assertIn("校验结果", workbook.sheetnames)
            self.assertEqual(
                workbook["并行审核主表"].max_row - 1,
                len(self.rows),
            )
            workbook.close()


if __name__ == "__main__":
    unittest.main()
