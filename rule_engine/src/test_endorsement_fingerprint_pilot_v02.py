# -*- coding: utf-8 -*-

import unittest
import json
import tempfile
from pathlib import Path

from endorsement_fingerprint_pilot_v02 import (
    assemble_draft_assets,
    build_candidate_groups,
    build_canonical_rule_groups,
    build_singleton_cluster,
    call_with_checkpoint,
    cluster_messages,
    detect_source_scope_conflict,
    export_review_html,
    export_review_workbook,
    fingerprint_messages,
    infer_object_scope_from_legal_text,
    legal_source_details,
    select_endorsement_records,
    should_route_to_proactive,
    sort_rule_uids_by_legal_hierarchy,
    validate_cluster_payload,
    validate_fingerprint,
)


def record(uid, title, text, track="保健食品", level=1, source_type="法律"):
    return {
        "track": track,
        "source_file": f"{track}/source.json",
        "legal_sources": [{
            "id": "SRC", "name": "中华人民共和国广告法",
            "type": source_type, "legal_level": level,
        }],
        "rule": {
            "rule_uid": uid,
            "rule_id": uid,
            "title": title,
            "source_type": source_type,
            "legal_basis": [{
                "source_id": "SRC", "article": "第三十八条",
                "legal_level": level, "text": text,
            }],
        },
    }


def fingerprint(uid, *, action="use_recommender", effect="prohibited", family="endorsement_eligibility"):
    return {
        "rule_uid": uid,
        "issue_family": family,
        "regulated_action": action,
        "legal_effect": effect,
        "actor_scope": ["expert"],
        "object_scope": ["health_food"],
        "channel_scope": [],
        "platform_scope": [],
        "source_type": "law",
        "legal_level": 1,
        "legal_basis_key": ["SRC:第三十八条"],
        "evidence_requirement": "content",
        "source_scope_conflict": False,
        "recommended_parent": "ENDORSEMENT.PROHIBITED_RECOMMENDER",
        "canonical_problem_name": "特定品类禁止使用推荐人或证明人",
        "aliases": [],
        "confidence": "high",
        "reason": "规则原文直接禁止专家作推荐证明。",
    }


class EndorsementFingerprintPilotTests(unittest.TestCase):
    def test_selector_requires_endorsement_role_not_generic_proof_material(self):
        records = [
            record("keep", "保健食品广告不得利用专家作推荐证明", "不得利用专家、用户的名义作推荐、证明。"),
            record("drop", "数据引证需提供证明材料", "广告使用数据应当提供证明材料。"),
            record("drop2", "食品检测机构证明提及要求", "涉及食品检测机构认证，需提供相关证明材料。"),
            record("keep_consent", "广告中使用他人名义或形象需书面同意", "使用他人名义或形象应当取得书面同意。"),
            record("drop3", "农药广告不得违反安全使用规程", "本条还规定不得利用专家作推荐、证明。"),
        ]
        selected = select_endorsement_records(records)
        self.assertEqual([item["rule"]["rule_uid"] for item in selected], ["keep", "keep_consent"])

    def test_validate_fingerprint_requires_all_structured_fields(self):
        item = fingerprint("r1")
        self.assertEqual(validate_fingerprint(item, "r1")["legal_effect"], "prohibited")
        del item["actor_scope"]
        with self.assertRaisesRegex(ValueError, "missing fingerprint fields"):
            validate_fingerprint(item, "r1")

    def test_groups_cross_old_skeleton_branches_but_separate_legal_effects(self):
        first = fingerprint("r1")
        first["old_skeleton_leaf_key"] = "OLD.A"
        second = fingerprint("r2", action="use_expert_for_testimonial")
        second["recommended_parent"] = "ENDORSEMENT.PROHIBITED_SUBJECT"
        second["old_skeleton_leaf_key"] = "OLD.B"
        consent = fingerprint("r3", effect="consent_required")
        groups = build_candidate_groups([first, second, consent])
        memberships = sorted(sorted(group["rule_uids"]) for group in groups)
        self.assertEqual(memberships, [["r1", "r2"], ["r3"]])

    def test_cluster_validator_blocks_mixed_legal_effect_merge(self):
        group = {
            "group_id": "g1",
            "rule_uids": ["r1", "r2"],
            "fingerprints": [
                fingerprint("r1", effect="prohibited"),
                fingerprint("r2", effect="joint_liability"),
            ],
        }
        payload = {"clusters": [{
            "cluster_key": "BAD",
            "canonical_name": "错误合并",
            "member_rule_uids": ["r1", "r2"],
            "relation": "merge",
            "recommended_parent": "ENDORSEMENT.PROHIBITED_SUBJECT",
            "confidence": "high",
            "reason": "名称相似",
        }]}
        with self.assertRaisesRegex(ValueError, "mixed legal_effect"):
            validate_cluster_payload(payload, group)

    def test_source_scope_conflict_uses_legal_text_over_folder_track(self):
        item = record(
            "r1",
            "禁止使用推荐人",
            "农药、兽药、饲料和饲料添加剂广告不得利用科研单位、专家作推荐、证明。",
            track="保健食品",
        )
        self.assertTrue(detect_source_scope_conflict(item, ["pesticide", "veterinary_drug", "feed"]))
        self.assertTrue(detect_source_scope_conflict(item, ["health_food"]))

    def test_legal_source_details_are_mapped_from_record(self):
        item = record("r1", "禁止专家推荐", "不得利用专家作推荐、证明。")
        details = legal_source_details(item)
        self.assertEqual(details[0]["source_name"], "中华人民共和国广告法")
        self.assertEqual(details[0]["article"], "第三十八条")
        self.assertEqual(details[0]["original_text"], "不得利用专家作推荐、证明。")

    def test_rule_order_follows_legal_hierarchy_then_general_before_specific(self):
        records = {
            "platform": record("platform", "平台规则", "平台不得推荐。", level=6, source_type="平台规则"),
            "specific": record("specific", "行业法规", "保健食品不得推荐。", level=1),
            "general": record("general", "一般法律", "广告不得推荐。", level=1),
        }
        ordered = sort_rule_uids_by_legal_hierarchy(
            ["platform", "specific", "general"], records,
            {"platform": ["platform"], "specific": ["health_food"], "general": ["general"]},
        )
        self.assertEqual(ordered, ["general", "specific", "platform"])


    def test_prompt_marks_track_as_weak_hint_and_requires_source_type(self):
        messages = fingerprint_messages([record("r1", "\u7981\u6b62\u4e13\u5bb6\u63a8\u8350", "\u4e0d\u5f97\u5229\u7528\u4e13\u5bb6\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002")])
        payload = json.loads(messages[1]["content"])
        self.assertEqual(payload["rules"][0]["track_hint_only"], "\u4fdd\u5065\u98df\u54c1")
        self.assertIn("source_type", payload["rules"][0])
        self.assertIn("legal_effect", payload["required_fields"])
        self.assertIsInstance(payload["output"]["fingerprints"][0]["actor_scope"], list)
        self.assertIsInstance(payload["output"]["fingerprints"][0]["legal_basis_key"], list)

    def test_cluster_prompt_preserves_exact_rule_uid_coverage(self):
        group = build_candidate_groups([fingerprint("r1"), fingerprint("r2")])[0]
        payload = json.loads(cluster_messages(group)[1]["content"])
        self.assertEqual([item["rule_uid"] for item in payload["fingerprints"]], ["r1", "r2"])

    def test_checkpoint_reuses_valid_deepseek_payload(self):
        class Client:
            calls = 0
            def create_chat_completion(self, **kwargs):
                self.calls += 1
                return {"choices": [{"message": {"content": json.dumps({"value": 1})}}]}

        with tempfile.TemporaryDirectory() as folder:
            client = Client()
            checkpoint = Path(folder) / "checkpoint.json"
            first = call_with_checkpoint(client, [{"role": "user", "content": "x"}], checkpoint)
            second = call_with_checkpoint(client, [{"role": "user", "content": "x"}], checkpoint)
            self.assertEqual(first, second)
            self.assertEqual(client.calls, 1)


    def test_assembled_tree_has_no_source_less_issue_and_preserves_rules(self):
        records = {
            "r1": record("r1", "\u7981\u6b62\u4e13\u5bb6\u63a8\u8350", "\u4e0d\u5f97\u5229\u7528\u4e13\u5bb6\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002"),
            "r2": record("r2", "\u7981\u6b62\u673a\u6784\u63a8\u8350", "\u4e0d\u5f97\u5229\u7528\u673a\u6784\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002"),
        }
        fingerprints = [fingerprint("r1"), fingerprint("r2")]
        clusters = [{
            "cluster_key": "PROHIBITED_RECOMMENDER",
            "canonical_name": "\u7279\u5b9a\u54c1\u7c7b\u7981\u6b62\u4f7f\u7528\u63a8\u8350\u4eba\u6216\u8bc1\u660e\u4eba",
            "member_rule_uids": ["r1", "r2"],
            "relation": "merge", "confidence": "high", "reason": "\u6cd5\u5f8b\u6548\u679c\u76f8\u540c",
            "legal_effect": "prohibited",
        }]
        assets = assemble_draft_assets(fingerprints, clusters, records)
        issue_nodes = [n for n in assets["taxonomy"]["nodes"] if n["node_type"] == "legal_issue"]
        self.assertTrue(all(n["rule_uids"] for n in issue_nodes))
        self.assertEqual(
            {m["rule_uid"] for m in assets["mapping"]["mappings"]},
            {"r1", "r2"},
        )


    def test_singleton_cluster_carries_final_parent(self):
        item = fingerprint("r1")
        cluster = build_singleton_cluster(item)
        self.assertEqual(cluster["recommended_parent"], item["recommended_parent"])


    def test_legal_text_scope_overrides_misleading_track_and_title(self):
        item = record(
            "r1",
            "\u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u7981\u6b62\u4e13\u5bb6\u63a8\u8350",
            "\u519c\u836f\u3001\u517d\u836f\u3001\u9972\u6599\u548c\u9972\u6599\u6dfb\u52a0\u5242\u5e7f\u544a\u4e0d\u5f97\u5229\u7528\u4e13\u5bb6\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002",
            track="\u4fdd\u5065\u98df\u54c1",
        )
        scopes = infer_object_scope_from_legal_text(item)
        self.assertEqual(scopes, ["pesticide", "veterinary_drug", "feed", "feed_additive"])
        self.assertTrue(detect_source_scope_conflict(item, ["health_food"]))

    def test_canonical_group_preserves_occurrences_and_selects_correct_track_record(self):
        correct = record(
            "correct", "\u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u5229\u7528\u4ee3\u8a00\u4eba",
            "\u7b2c\u5341\u516b\u6761 \u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u5229\u7528\u4ee3\u8a00\u4eba\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002",
            track="\u4fdd\u5065\u98df\u54c1",
        )
        wrong = record(
            "wrong", "\u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u5229\u7528\u4ee3\u8a00\u4eba",
            "\u7b2c\u5341\u516b\u6761  \u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u5229\u7528\u4ee3\u8a00\u4eba\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002",
            track="\u6e38\u620f",
        )
        groups, uid_to_key = build_canonical_rule_groups(
            {"correct": correct, "wrong": wrong},
            {"correct": fingerprint("correct"), "wrong": fingerprint("wrong")},
        )
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["canonical_rule_uid"], "correct")
        self.assertEqual(set(groups[0]["source_occurrence_uids"]), {"correct", "wrong"})
        self.assertEqual(uid_to_key["correct"], uid_to_key["wrong"])

    def test_third_party_doctor_proof_routes_to_proactive_check(self):
        item = record(
            "doctor", "\u666e\u901a\u98df\u54c1\u6d89\u53ca\u533b\u751f\u63a8\u8350\u9700\u63d0\u4ea4\u533b\u751f\u8bc1\u660e",
            "\u666e\u901a\u98df\u54c1\u6d89\u53ca\u533b\u751f\u63a8\u8350\u9700\u8981\u63d0\u4ea4\u533b\u751f\u8bc1\u660e\u3002",
        )
        fp = fingerprint("doctor", effect="qualification_check")
        fp["evidence_requirement"] = "fact"
        self.assertTrue(should_route_to_proactive(fp, item))

    def test_prohibition_branches_share_subject_eligibility_parent(self):
        records = {
            "subject": record("subject", "\u672a\u6210\u5e74\u4eba\u4e0d\u5f97\u4ee3\u8a00", "\u4e0d\u5f97\u4f7f\u7528\u672a\u6210\u5e74\u4eba\u4ee3\u8a00\u3002"),
            "category": record("category", "\u4fdd\u5065\u98df\u54c1\u4e0d\u5f97\u4ee3\u8a00", "\u4fdd\u5065\u98df\u54c1\u4e0d\u5f97\u4f7f\u7528\u4ee3\u8a00\u4eba\u3002"),
        }
        first = fingerprint("subject")
        first["recommended_parent"] = "ENDORSEMENT.PROHIBITED_SUBJECT"
        second = fingerprint("category")
        second["recommended_parent"] = "ENDORSEMENT.PROHIBITED_RECOMMENDER_BY_CATEGORY"
        clusters = [
            {"cluster_key": "SUBJECT", "canonical_name": "\u7981\u6b62\u7279\u5b9a\u4e3b\u4f53\u4ee3\u8a00", "member_rule_uids": ["subject"], "relation": "separate", "confidence": "high", "reason": "", "recommended_parent": first["recommended_parent"], "legal_effect": "prohibited"},
            {"cluster_key": "CATEGORY", "canonical_name": "\u7279\u5b9a\u54c1\u7c7b\u7981\u6b62\u4ee3\u8a00", "member_rule_uids": ["category"], "relation": "separate", "confidence": "high", "reason": "", "recommended_parent": second["recommended_parent"], "legal_effect": "prohibited"},
        ]
        assets = assemble_draft_assets([first, second], clusters, records)
        nodes = {node["issue_id"]: node for node in assets["taxonomy"]["nodes"]}
        self.assertEqual(nodes["ENDORSEMENT.PROHIBITED_SUBJECT"]["parent_issue_id"], "ENDORSEMENT.SUBJECT_ELIGIBILITY")
        self.assertEqual(nodes["ENDORSEMENT.PROHIBITED_RECOMMENDER_BY_CATEGORY"]["parent_issue_id"], "ENDORSEMENT.SUBJECT_ELIGIBILITY")


    def test_html_proactive_check_preserves_source_rule_uid_and_version(self):
        item = record(
            "doctor", "\u666e\u901a\u98df\u54c1\u6d89\u53ca\u533b\u751f\u63a8\u8350\u9700\u63d0\u4ea4\u533b\u751f\u8bc1\u660e",
            "\u666e\u901a\u98df\u54c1\u6d89\u53ca\u533b\u751f\u63a8\u8350\u9700\u8981\u63d0\u4ea4\u533b\u751f\u8bc1\u660e\u3002",
        )
        fp = fingerprint("doctor", effect="qualification_check")
        fp["evidence_requirement"] = "fact"
        assets = assemble_draft_assets([fp], [build_singleton_cluster(fp)], {"doctor": item})
        with tempfile.TemporaryDirectory() as directory:
            output = export_review_html(Path(directory) / "review.html", assets)
            html = output.read_text(encoding="utf-8")
        self.assertIn("doctor", html)
        self.assertIn("v0.3", html)



    def test_focus_review_sheet_is_self_contained_and_editable(self):
        first = record(
            "r1", "\u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u4f7f\u7528\u4ee3\u8a00\u4eba",
            "\u7b2c\u5341\u516b\u6761 \u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u5229\u7528\u5e7f\u544a\u4ee3\u8a00\u4eba\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002",
            track="\u4fdd\u5065\u98df\u54c1",
        )
        second = record(
            "r2", "\u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u4f7f\u7528\u4ee3\u8a00\u4eba",
            "\u7b2c\u5341\u516b\u6761 \u4fdd\u5065\u98df\u54c1\u5e7f\u544a\u4e0d\u5f97\u5229\u7528\u5e7f\u544a\u4ee3\u8a00\u4eba\u4f5c\u63a8\u8350\u3001\u8bc1\u660e\u3002",
            track="\u6e38\u620f",
        )
        fp1 = fingerprint("r1")
        fp2 = fingerprint("r2")
        fp2["source_scope_conflict"] = True
        clusters = [{
            "cluster_key": "HEALTH_FOOD_ENDORSER_PROHIBITED",
            "canonical_name": "\u4fdd\u5065\u98df\u54c1\u7981\u6b62\u4ee3\u8a00",
            "member_rule_uids": ["r1", "r2"],
            "relation": "merge", "confidence": "high", "reason": "",
            "recommended_parent": "ENDORSEMENT.PROHIBITED_RECOMMENDER_BY_CATEGORY",
            "legal_effect": "prohibited",
        }]
        assets = assemble_draft_assets([fp1, fp2], clusters, {"r1": first, "r2": second})
        with tempfile.TemporaryDirectory() as directory:
            output = export_review_workbook(Path(directory) / "review.xlsx", assets)
            from openpyxl import load_workbook
            workbook = load_workbook(output)
        sheet = workbook["\u5f85\u4eba\u5de5\u91cd\u70b9\u5ba1\u6838"]
        headers = [cell.value for cell in sheet[1]]
        for expected in (
            "\u89c4\u5219\u6807\u9898", "\u6765\u6e90\u7c7b\u578b", "\u6cd5\u89c4/\u5e73\u53f0\u89c4\u5219\u540d\u79f0", "\u6761\u6b3e",
            "\u89c4\u5219\u539f\u6587", "\u6a21\u578b\u8bc6\u522b\u9002\u7528\u5bf9\u8c61", "\u6807\u51c6rule_uid",
            "\u5176\u4ed6\u91cd\u590d\u6765\u6e90rule_uid", "\u5efa\u8bae\u5224\u65ad", "\u8c03\u6574\u540e\u7684\u5f52\u5c5e",
        ):
            self.assertIn(expected, headers)
        original_text_col = headers.index("\u89c4\u5219\u539f\u6587") + 1
        duplicate_col = headers.index("\u5176\u4ed6\u91cd\u590d\u6765\u6e90rule_uid") + 1
        suggestion_col = headers.index("\u5efa\u8bae\u5224\u65ad") + 1
        self.assertTrue(all(sheet.cell(row, original_text_col).value for row in range(2, sheet.max_row + 1)))
        self.assertIn("r2", sheet.cell(2, duplicate_col).value)
        self.assertIn("\u5408\u5e76", sheet.cell(2, suggestion_col).value)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.auto_filter.ref)
        self.assertTrue(sheet.data_validations.dataValidation)
        self.assertIn("\u5408\u5e76\u4e3a\u540c\u4e00\u6807\u51c6\u89c4\u8303", sheet.data_validations.dataValidation[0].formula1)


if __name__ == "__main__":
    unittest.main()
