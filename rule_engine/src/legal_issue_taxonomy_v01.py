# -*- coding: utf-8 -*-
"""Build a review-only legal issue taxonomy from rule-level DeepSeek drafts."""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROMPT_VERSION = "legal_issue_taxonomy_v01_20260912"
CONFIDENCE = {"high", "medium", "low"}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def atomic_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def _dedupe(values):
    return list(dict.fromkeys(str(value).strip() for value in values if str(value or "").strip()))


def _source_details(record):
    rule = record.get("rule") or {}
    names = {str(item.get("id") or ""): str(item.get("name") or "") for item in record.get("legal_sources") or []}
    result = []
    for basis in rule.get("legal_basis") or []:
        source_id = str(basis.get("source_id") or "")
        result.append({
            "source_name": names.get(source_id) or source_id,
            "article": basis.get("article"),
            "original_text": basis.get("text") or basis.get("quote") or "",
        })
    return result or [{"source_name": "", "article": "", "original_text": ""}]


def build_issue_inputs(issue_asset, mapping_asset, records_by_uid):
    mapping_by_uid = {item.get("rule_uid"): item for item in mapping_asset.get("mappings") or []}
    rows = []
    for issue in sorted(issue_asset.get("issues") or [], key=lambda item: item.get("issue_id") or ""):
        rule_uids = _dedupe(issue.get("candidate_rule_uids") or [])
        source_rules = []
        for uid in rule_uids:
            record = records_by_uid.get(uid)
            if not record:
                raise ValueError(f"unknown rule_uid: {uid}")
            rule = record.get("rule") or {}
            details = _source_details(record)
            source_rules.append({
                "rule_uid": uid,
                "rule_id": rule.get("rule_id") or (mapping_by_uid.get(uid) or {}).get("rule_id"),
                "title": rule.get("title"),
                "track": record.get("track"),
                "source_type": rule.get("source_type"),
                "platform": rule.get("platform"),
                "applies_to": rule.get("applies_to") or {},
                "source_name": "\n".join(_dedupe(item["source_name"] for item in details)),
                "article": "\n".join(_dedupe(item["article"] for item in details)),
                "original_text": "\n".join(_dedupe(item["original_text"] for item in details)),
                "source_file": record.get("source_file"),
            })
        rows.append({
            "issue_id": issue.get("issue_id"),
            "name": issue.get("name"),
            "track": issue.get("track"),
            "definition": issue.get("definition"),
            "in_scope": issue.get("in_scope") or [],
            "out_of_scope": issue.get("out_of_scope") or [],
            "claim_types": issue.get("claim_types") or [],
            "rule_uids": rule_uids,
            "source_rules": source_rules,
        })
    return rows


def _response_json(response):
    try:
        content = response["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as exc:
        raise ValueError("DeepSeek response missing message content") from exc
    if isinstance(content, dict):
        return content
    text = str(content).strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I | re.S)
    return json.loads(text)


def _call_json(client, messages):
    response = client.create_chat_completion(
        messages=messages,
        model="deepseek-chat",
        temperature=0.0,
        response_format={"type": "json_object"},
    )
    try:
        return _response_json(response)
    except json.JSONDecodeError:
        raw = str(response.get("choices", [{}])[0].get("message", {}).get("content") or "")
        repair_messages = list(messages) + [
            {"role": "assistant", "content": raw},
            {"role": "user", "content": "上一个响应不是完整、合法的JSON。请只修复JSON语法和截断结构，不新增事实、不改变已表达内容，并返回一个完整JSON对象。"},
        ]
        repaired = client.create_chat_completion(
            messages=repair_messages,
            model="deepseek-chat",
            temperature=0.0,
            response_format={"type": "json_object"},
        )
        return _response_json(repaired)


def parse_fingerprint_batch(payload, expected_issue_ids):
    records = payload.get("fingerprints") or []
    actual = [item.get("issue_id") for item in records]
    unknown = sorted(set(actual) - set(expected_issue_ids))
    if unknown:
        raise ValueError("unknown issue_id: " + ", ".join(unknown))
    if set(actual) != set(expected_issue_ids) or len(actual) != len(set(actual)):
        raise ValueError("fingerprint batch must contain each expected issue_id exactly once")
    required = {
        "issue_id", "regulated_behavior", "claim_object", "legal_test",
        "evidence_requirement", "applicability", "skeleton_leaf_key",
        "canonical_problem_name", "aliases", "confidence", "reason",
    }
    for item in records:
        missing = required - set(item)
        if missing:
            raise ValueError(f"fingerprint missing fields: {sorted(missing)}")
        if item.get("confidence") not in CONFIDENCE:
            raise ValueError("invalid fingerprint confidence")
    return records


def _normalize(value):
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", str(value or "")).lower()


def build_candidate_groups(fingerprints, max_group_size=12):
    buckets = defaultdict(list)
    for item in fingerprints:
        key = str(item.get("skeleton_leaf_key") or "UNRESOLVED").strip()
        buckets[key].append(item)
    groups = []
    for leaf_key, members in sorted(buckets.items()):
        members.sort(key=lambda item: (
            _normalize(item.get("regulated_behavior")),
            _normalize(item.get("claim_object")),
            item.get("issue_id") or "",
        ))
        for index in range(0, len(members), max_group_size):
            chunk = members[index:index + max_group_size]
            groups.append({
                "group_id": f"{leaf_key}#{index // max_group_size + 1}",
                "skeleton_leaf_key": leaf_key,
                "issue_ids": [item["issue_id"] for item in chunk],
                "fingerprints": chunk,
            })
    return groups


def _review():
    return {"status": "pending", "reviewer": None, "reviewed_at": None, "notes": None}


def assemble_taxonomy(skeleton, proposals, issue_asset, mapping_asset):
    old_issues = {item.get("issue_id"): item for item in issue_asset.get("issues") or []}
    old_to_new = {}
    taxonomy_nodes = []
    for node in skeleton.get("nodes") or []:
        taxonomy_nodes.append({
            "issue_id": node["node_id"],
            "parent_issue_id": node.get("parent_node_id"),
            "level": node.get("level"),
            "node_type": node.get("node_type") or "directory",
            "name": node.get("name"),
            "definition": node.get("definition") or "",
            "aliases": node.get("aliases") or [],
            "applicability": node.get("applicability") or {"industries": [], "platforms": []},
            "candidate_rule_uids": [],
            "source_issue_ids": [],
            "generation": node.get("generation") or {"model": "deepseek-chat", "prompt_version": PROMPT_VERSION, "confidence": node.get("confidence") or "medium", "reason": node.get("reason") or "骨架节点"},
            "review": _review(),
        })
    known_ids = {item["issue_id"] for item in taxonomy_nodes}
    for proposal in proposals:
        base = str(proposal.get("parent_node_id") or "UNRESOLVED")
        key = re.sub(r"[^A-Z0-9_]+", "_", str(proposal.get("canonical_issue_key") or "ISSUE").upper()).strip("_")
        issue_id = f"{base}.{key}"
        suffix = 2
        while issue_id in known_ids:
            issue_id = f"{base}.{key}_{suffix}"
            suffix += 1
        known_ids.add(issue_id)
        members = _dedupe(proposal.get("member_issue_ids") or [])
        for old_id in members:
            if old_id in old_to_new:
                raise ValueError(f"old issue mapped more than once: {old_id}")
            old_to_new[old_id] = issue_id
        taxonomy_nodes.append({
            "issue_id": issue_id,
            "parent_issue_id": base if base in known_ids else None,
            "level": next((item.get("level", 2) + 1 for item in taxonomy_nodes if item["issue_id"] == base), 3),
            "node_type": "legal_issue",
            "name": proposal.get("canonical_name"),
            "definition": proposal.get("definition") or "",
            "aliases": _dedupe(proposal.get("aliases") or [old_issues.get(old_id, {}).get("name") for old_id in members]),
            "applicability": proposal.get("applicability") or {"industries": [], "platforms": []},
            "candidate_rule_uids": [],
            "source_issue_ids": members,
            "generation": {"model": "deepseek-chat", "prompt_version": PROMPT_VERSION, "confidence": proposal.get("confidence") or "medium", "reason": proposal.get("reason") or ""},
            "review": _review(),
        })
    projected = []
    rules_by_new = defaultdict(list)
    for mapping in mapping_asset.get("mappings") or []:
        source_ids = [mapping.get("primary_issue_id")] + list(mapping.get("secondary_issue_ids") or [])
        target_ids = _dedupe(old_to_new.get(old_id) for old_id in source_ids)
        if not target_ids:
            continue
        projected.append({
            "rule_uid": mapping.get("rule_uid"),
            "rule_id": mapping.get("rule_id"),
            "primary_issue_id": target_ids[0],
            "secondary_issue_ids": target_ids[1:],
            "source_issue_ids": _dedupe(source_ids),
            "review": _review(),
        })
        for target_id in target_ids:
            rules_by_new[target_id].append(mapping.get("rule_uid"))
    for node in taxonomy_nodes:
        node["candidate_rule_uids"] = _dedupe(rules_by_new.get(node["issue_id"], []))
    now = datetime.now(timezone.utc).isoformat()
    return (
        {"schema_version": "0.2", "asset_status": "draft", "generated_at": now, "issues": taxonomy_nodes},
        {"schema_version": "0.2", "asset_status": "draft", "generated_at": now, "mappings": projected},
    )


def validate_taxonomy_result(taxonomy, mappings, expected_rule_uids):
    errors = []
    issues = taxonomy.get("issues") or []
    ids = [item.get("issue_id") for item in issues]
    known = set(ids)
    if len(ids) != len(known):
        errors.append("duplicate issue_id")
    parents = {item.get("issue_id"): item.get("parent_issue_id") for item in issues}
    for issue_id, parent in parents.items():
        if parent and parent not in known:
            errors.append(f"unknown parent: {issue_id} -> {parent}")
        seen = set()
        current = issue_id
        while current:
            if current in seen:
                errors.append(f"cycle: {issue_id}")
                break
            seen.add(current)
            current = parents.get(current)
    rows = mappings.get("mappings") or []
    mapped = [item.get("rule_uid") for item in rows]
    if len(mapped) != len(set(mapped)):
        errors.append("duplicate rule mapping")
    missing = set(expected_rule_uids) - set(mapped)
    extra = set(mapped) - set(expected_rule_uids)
    if missing:
        errors.append(f"missing rules: {len(missing)}")
    if extra:
        errors.append(f"unknown rules: {len(extra)}")
    for item in rows:
        if item.get("primary_issue_id") not in known:
            errors.append(f"unknown primary issue: {item.get('rule_uid')}")
    return {"valid": not errors, "errors": errors}


def _sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def export_merge_review_workbook(path, skeleton, taxonomy, mappings, proposals, old_issue_asset, records_by_uid, warnings):
    workbook = Workbook()
    workbook.remove(workbook.active)
    issue_by_id = {item.get("issue_id"): item for item in taxonomy.get("issues") or []}
    old_by_id = {item.get("issue_id"): item for item in old_issue_asset.get("issues") or []}
    tree_rows = []
    for item in sorted(issue_by_id.values(), key=lambda value: (value.get("issue_id") or "")):
        tree_rows.append([item.get("level"), item.get("issue_id"), item.get("parent_issue_id"), item.get("node_type"), item.get("name"), item.get("definition"), "；".join(item.get("aliases") or []), len(item.get("source_issue_ids") or []), len(item.get("candidate_rule_uids") or []), (item.get("generation") or {}).get("confidence"), (item.get("generation") or {}).get("reason"), "", ""])
    _sheet(workbook, "标准问题树", ["层级", "标准问题ID", "父节点", "节点类型", "标准问题名称", "定义", "别名", "旧问题数", "规则数", "模型置信度", "模型理由", "人工结论", "审核备注"], tree_rows)
    proposal_by_member = {old_id: item for item in proposals for old_id in item.get("member_issue_ids") or []}
    merge_rows = []
    for old_id, old in sorted(old_by_id.items()):
        proposal = proposal_by_member.get(old_id) or {}
        target = next((item for item in issue_by_id.values() if old_id in (item.get("source_issue_ids") or [])), {})
        merge_rows.append([old_id, old.get("name"), old.get("definition"), target.get("issue_id"), target.get("name"), len(proposal.get("member_issue_ids") or []), proposal.get("relation") or "merge", proposal.get("confidence"), proposal.get("reason"), "", ""])
    _sheet(workbook, "归并审核", ["旧问题ID", "旧问题名称", "旧问题定义", "标准问题ID", "标准问题名称", "簇内问题数", "关系结论", "模型置信度", "模型理由", "人工结论", "审核备注"], merge_rows)
    source_rows = []
    for mapping in mappings.get("mappings") or []:
        record = records_by_uid.get(mapping.get("rule_uid")) or {}
        rule = record.get("rule") or {}
        details = _source_details(record)
        target_ids = _dedupe(
            [mapping.get("primary_issue_id")] + list(mapping.get("secondary_issue_ids") or [])
        )
        for target_id in target_ids:
            target = issue_by_id.get(target_id) or {}
            source_rows.append([target.get("issue_id"), target.get("name"), mapping.get("rule_uid"), mapping.get("rule_id") or rule.get("rule_id"), rule.get("title"), "\n".join(_dedupe(item["source_name"] for item in details)), "\n".join(_dedupe(item["article"] for item in details)), "\n".join(_dedupe(item["original_text"] for item in details)), record.get("source_file"), ";".join(mapping.get("source_issue_ids") or []), "", ""])
    _sheet(workbook, "问题-规则原文对照", ["标准问题ID", "标准问题名称", "rule_uid", "rule_id", "规则标题", "规则来源名称", "法条定位", "原始规则原文", "JSON源文件", "原问题ID", "人工结论", "审核备注"], source_rows)
    singleton_rows = [row for row in merge_rows if row[5] == 1]
    _sheet(workbook, "未归并单例", ["旧问题ID", "旧问题名称", "旧问题定义", "标准问题ID", "标准问题名称", "簇内问题数", "关系结论", "模型置信度", "模型理由", "人工结论", "审核备注"], singleton_rows)
    warning_rows = [[item.get("severity"), item.get("code"), item.get("record_id"), item.get("message"), ""] for item in warnings]
    _sheet(workbook, "质量告警", ["严重度", "代码", "记录", "说明", "处理结论"], warning_rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _checkpoint_call(client, messages, checkpoint, validator=None):
    digest = hashlib.sha256(json.dumps(messages, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    checkpoint = Path(checkpoint)
    if checkpoint.exists():
        saved = read_json(checkpoint)
        if saved.get("digest") == digest:
            payload = saved["payload"]
            return validator(payload) if validator else payload
    payload = _call_json(client, messages)
    result = validator(payload) if validator else payload
    atomic_json(checkpoint, {"digest": digest, "payload": payload})
    return result



def build_seed_skeleton(issue_count, scan_fragment_count):
    branches = {
        "SCOPE_ACCESS": ("适用主体与准入", [
            ("SUBJECT_QUALIFICATION", "主体资格"), ("PRODUCT_REGISTRATION", "产品注册备案"),
            ("AD_REVIEW_ACCESS", "广告审查与投放准入"), ("SCOPE_CLASSIFICATION", "产品属性与适用范围"),
        ]),
        "PROHIBITED_CONTENT": ("禁止投放内容", [
            ("PROHIBITED_PRODUCTS", "禁投商品与服务"), ("POLITICAL_SECURITY", "政治与国家安全"),
            ("VULGAR_SEXUAL_VIOLENT", "低俗、色情与暴力内容"), ("GAMBLING_SUPERSTITION", "赌博与迷信内容"),
        ]),
        "TRUTHFULNESS": ("宣传真实性", [
            ("FALSE_FACT", "虚构或不实事实"), ("MISLEADING_OMISSION", "误导与重大遗漏"),
            ("MATERIAL_MISMATCH", "广告素材与商品服务不一致"), ("FABRICATED_EFFECT", "虚构效果与体验"),
        ]),
        "CLAIM_EXPRESSION": ("宣称表达方式", [
            ("ABSOLUTE", "绝对化用语"), ("COMPARATIVE_RANKING", "比较、排名与唯一性宣称"),
            ("GUARANTEE_COMMITMENT", "保证与承诺性表达"), ("AMBIGUOUS_EXAGGERATED", "模糊、夸张与易误解表达"),
        ]),
        "EFFICACY_PERFORMANCE": ("功效与性能宣称", [
            ("DISEASE_MEDICAL", "疾病预防治疗与医疗化宣称"), ("HEALTH_COSMETIC_EFFECT", "保健及化妆品功效宣称"),
            ("SAFETY_SIDE_EFFECT", "安全性与副作用宣称"), ("PERFORMANCE_RESULT", "性能、效果与结果宣称"),
        ]),
        "PRICE_PROMOTION": ("价格与促销", [
            ("PRICE_ACCURACY", "价格真实性与准确标示"), ("DISCOUNT_LIMIT", "折扣优惠与限制条件"),
            ("GIFT_FREE", "赠送、免费与福利活动"), ("PRIZE_PROBABILITY", "抽奖、概率与奖励机制"),
        ]),
        "EVIDENCE_FACT": ("证明材料与事实核验", [
            ("DATA_CITATION", "数据、统计与引证来源"), ("TEST_RESEARCH", "测试、实验与研究证明"),
            ("QUALIFICATION_FILING", "资质、许可与备案材料"), ("AUTHORIZATION_PROVENANCE", "授权、来源与权属证明"),
        ]),
        "DISCLOSURE_WARNING": ("信息披露与警示语", [
            ("AD_IDENTIFIABILITY", "广告可识别性"), ("STATUTORY_WARNING", "法定警示语"),
            ("CONDITIONS_LIMITS", "适用条件与限制披露"), ("LABEL_SOURCE", "标签、出处与必要信息"),
        ]),
        "ENDORSEMENT_REVIEW": ("代言、推荐与用户评价", [
            ("ENDORSER_ELIGIBILITY", "代言人资格"), ("PERSONAL_EXPERIENCE", "真实使用体验"),
            ("EXPERT_ORGANIZATION", "专家及机构推荐证明"), ("UGC_BUYER", "用户评价与买手推广"),
        ]),
        "IP_PERSONALITY": ("知识产权与人格权益", [
            ("COPYRIGHT_TRADEMARK", "著作权与商标权"), ("PORTRAIT_NAME", "肖像、姓名与人格权益"),
            ("PATENT_AWARD", "专利、奖项与荣誉"), ("THIRD_PARTY_AUTH", "第三方内容授权"),
        ]),
        "MINORS_PUBLIC_ORDER": ("未成年人及公序良俗", [
            ("MINOR_INDUCEMENT", "未成年人诱导与付费保护"), ("PUBLIC_MORALITY", "公序良俗"),
            ("DISCRIMINATION_INSULT", "歧视、侮辱与网络暴力"), ("DANGEROUS_BEHAVIOR", "危险行为与不良导向"),
        ]),
        "MATERIAL_PLATFORM": ("素材形式与平台规范", [
            ("MEDIA_QUALITY", "图片、音频与视频质量"), ("LAYOUT_SAFE_AREA", "版式、字体与安全区域"),
            ("FORM_SYMBOL", "素材形式与特殊符号"), ("PLACEMENT_JUMP", "广告位、跳转与落地页"),
        ]),
        "WORKFLOW_DUTY": ("投放流程与持续义务", [
            ("PRE_REVIEW", "投放前审查"), ("ARCHIVE_RECORD", "档案与记录保存"),
            ("PLATFORM_ALGORITHM", "平台与算法责任"), ("POST_LAUNCH_FULFILLMENT", "投放后履约与持续核验"),
        ]),
    }
    nodes = []
    for root_id, (root_name, children) in branches.items():
        nodes.append({"node_id": root_id, "parent_node_id": None, "level": 1, "name": root_name, "node_type": "directory"})
        for child_id, child_name in children:
            nodes.append({"node_id": f"{root_id}.{child_id}", "parent_node_id": root_id, "level": 2, "name": child_name, "node_type": "directory"})
    return {
        "schema_version": "0.1", "asset_status": "draft",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "prompt_version": PROMPT_VERSION,
        "generation_method": "approved_legal_domain_seed_with_deepseek_full_scan",
        "source_issue_count": issue_count,
        "scan_fragment_count": scan_fragment_count,
        "nodes": nodes,
    }

def skeleton_messages(issue_rows):
    compact = [{"issue_id": x["issue_id"], "name": x["name"], "track": x["track"], "definition": x["definition"], "rule_titles": [r.get("title") for r in x["source_rules"]]} for x in issue_rows]
    system = """你是广告合规法律问题目录架构师。请从输入问题中归纳稳定的三级目录骨架，不逐规则复制问题。L1是审核领域，L2是问题族，L3是可供具体法律问题挂载的标准分支。行业和平台通常属于适用范围，只有法律要件确实不同才拆分。不得新增法条。仅输出JSON。"""
    user = {"task": "生成本批问题的目录骨架片段", "requirements": {"merge_synonyms": True, "levels": [1, 2, 3], "node_id_format": "大写英文点分路径", "output": {"nodes": [{"node_id": "CLAIM.ABSOLUTE", "parent_node_id": "CLAIM", "level": 2, "name": "绝对化用语", "definition": "", "aliases": [], "confidence": "high/medium/low", "reason": ""}]}}, "issues": compact}
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(user, ensure_ascii=False)}]


def consolidate_skeleton_messages(fragments, max_total_nodes=60):
    system = """你是广告合规法律问题目录架构师。合并骨架片段，消除同义节点，形成一棵紧凑稳定的三级目录。不得逐规则建节点；行业和平台优先放适用范围。所有非根节点必须引用存在的父节点。每个节点只能包含四个字段：node_id、parent_node_id、level、name；严禁输出定义、别名、置信度、理由或其他字段。必须严格遵守节点数量上限，宁可保留较宽的标准分支，也不要输出大量细碎叶节点。仅输出JSON。"""
    compact_fragments = []
    for fragment in fragments:
        compact_fragments.append({"nodes": [
            {"node_id": node.get("node_id"), "parent_node_id": node.get("parent_node_id"), "level": node.get("level"), "name": node.get("name"), "aliases": (node.get("aliases") or [])[:3]}
            for node in (fragment.get("nodes") or [])
        ]})
    user = {
        "task": "合并目录骨架片段为V0.1",
        "constraints": {"max_level_1_nodes": 14, "max_total_nodes": max_total_nodes, "max_definition_chars": 60, "max_reason_chars": 40, "max_aliases_per_node": 5},
        "output": {"nodes": [{"node_id": "CLAIM", "parent_node_id": None, "level": 1, "name": "宣称表达"}]},
        "fragments": compact_fragments,
    }
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(user, ensure_ascii=False)}]



def reduce_skeleton_fragments(client, fragments, checkpoint_dir, chunk_size=7):
    checkpoint_dir = Path(checkpoint_dir)
    intermediate = []
    for index in range(0, len(fragments), chunk_size):
        chunk = fragments[index:index + chunk_size]
        payload = _checkpoint_call(
            client,
            consolidate_skeleton_messages(chunk, max_total_nodes=35),
            checkpoint_dir / "skeleton_intermediate" / f"{index // chunk_size:03d}.json",
        )
        intermediate.append(payload)
    return _checkpoint_call(
        client,
        consolidate_skeleton_messages(intermediate, max_total_nodes=60),
        checkpoint_dir / "skeleton_v01.json",
        parse_skeleton,
    )

def fingerprint_messages(issue_rows, skeleton):
    clipped = []
    for issue in issue_rows:
        item = dict(issue)
        item["source_rules"] = [{**rule, "original_text": str(rule.get("original_text") or "")[:700]} for rule in issue["source_rules"]]
        clipped.append(item)
    leaves = [{"node_id": x["node_id"], "name": x["name"], "definition": x.get("definition", "")} for x in skeleton.get("nodes") or [] if x.get("level") in {2, 3}]
    system = """你是广告合规规则资产分析员。逐个问题根据其定义和通过rule_uid提供的规则原文生成结构化问题指纹。必须忠于原文，不新增法条；每个输入issue_id恰好输出一次。skeleton_leaf_key必须选择给定骨架节点。仅输出JSON。"""
    output = {"fingerprints": [{"issue_id": "原ID", "regulated_behavior": "规制行为", "claim_object": "宣称或行为对象", "legal_test": "判断成立所需核心要件", "evidence_requirement": "content/fact/context/mixed", "applicability": {"industries": [], "platforms": []}, "skeleton_leaf_key": "给定节点ID", "canonical_problem_name": "建议标准问题名", "aliases": [], "confidence": "high/medium/low", "reason": "原文依据"}]}
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps({"task": "生成问题指纹", "skeleton_nodes": leaves, "issues": clipped, "output": output}, ensure_ascii=False)}]


def cluster_messages(group):
    system = """你是广告合规法律问题归并审核员。只在输入候选组内判断哪些旧问题法律要件相同可合并，哪些需保持独立或构成上下位关系。名称相似不是充分条件；规制行为、对象、法律要件和证据要求均相同才可合并。每个issue_id必须且只能出现一次，不得新增ID或法条。仅输出JSON。"""
    output = {"clusters": [{"canonical_issue_key": "大写英文或下划线", "canonical_name": "中文标准问题名", "parent_node_id": group["skeleton_leaf_key"], "member_issue_ids": ["原issue_id"], "relation": "merge/parent_child/related/separate", "definition": "标准定义", "aliases": [], "applicability": {"industries": [], "platforms": []}, "confidence": "high/medium/low", "reason": "归并或保持独立的理由"}]}
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps({"task": "候选组内归并", "group_id": group["group_id"], "fingerprints": group["fingerprints"], "output": output}, ensure_ascii=False)}]


def parse_skeleton(payload):
    nodes = payload.get("nodes") or []
    ids = [item.get("node_id") for item in nodes]
    if not nodes or len(ids) != len(set(ids)) or any(not item for item in ids):
        raise ValueError("invalid skeleton node identities")
    known = set(ids)
    for item in nodes:
        if item.get("level") not in {1, 2, 3}:
            raise ValueError("invalid skeleton level")
        if item.get("level") == 1 and item.get("parent_node_id") is not None:
            raise ValueError("level-1 node must be root")
        if item.get("level") > 1 and item.get("parent_node_id") not in known:
            raise ValueError("unknown skeleton parent")
    return {"schema_version": "0.1", "asset_status": "draft", "generated_at": datetime.now(timezone.utc).isoformat(), "prompt_version": PROMPT_VERSION, "nodes": nodes}


def parse_clusters(payload, group):
    clusters = payload.get("clusters") or []
    expected = set(group["issue_ids"])
    actual = [issue_id for item in clusters for issue_id in item.get("member_issue_ids") or []]
    if set(actual) != expected or len(actual) != len(set(actual)):
        raise ValueError(f"cluster coverage mismatch: {group['group_id']}")
    for item in clusters:
        if item.get("confidence") not in CONFIDENCE:
            raise ValueError("invalid cluster confidence")
        item["parent_node_id"] = group["skeleton_leaf_key"]
    return clusters



def cluster_with_coverage_retry(client, group, checkpoint):
    messages = cluster_messages(group)
    digest = hashlib.sha256(json.dumps(messages, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    checkpoint = Path(checkpoint)
    if checkpoint.exists():
        saved = read_json(checkpoint)
        if saved.get("digest") == digest:
            return parse_clusters(saved["payload"], group)
    payload = _call_json(client, messages)
    try:
        result = parse_clusters(payload, group)
    except ValueError as exc:
        actual = [issue_id for item in payload.get("clusters") or [] for issue_id in item.get("member_issue_ids") or []]
        expected = set(group["issue_ids"])
        missing = sorted(expected - set(actual))
        unexpected = sorted(set(actual) - expected)
        duplicates = sorted({item for item in actual if actual.count(item) > 1})
        repair_request = {
            "task": "只修复候选组成员覆盖，不改变有充分依据的归并判断",
            "validation_error": str(exc),
            "required_issue_ids": group["issue_ids"],
            "missing_issue_ids": missing,
            "unexpected_issue_ids": unexpected,
            "duplicate_issue_ids": duplicates,
            "requirements": [
                "每个required_issue_id必须且只能出现一次",
                "不得新增required_issue_ids之外的ID",
                "缺失问题可以加入已有簇；若法律要件不同则建立独立簇",
                "返回完整JSON对象，不要解释",
            ],
        }
        repaired_messages = messages + [
            {"role": "assistant", "content": json.dumps(payload, ensure_ascii=False)},
            {"role": "user", "content": json.dumps(repair_request, ensure_ascii=False)},
        ]
        payload = _call_json(client, repaired_messages)
        result = parse_clusters(payload, group)
    atomic_json(checkpoint, {"digest": digest, "payload": payload})
    return result


def _proposal_id(parent_node_id, proposal):
    raw = json.dumps({"parent": parent_node_id, "members": sorted(proposal.get("member_issue_ids") or [])}, ensure_ascii=False, sort_keys=True)
    return "P-" + hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]


def parent_consolidation_messages(parent_node_id, proposals):
    compact = []
    for proposal in proposals:
        compact.append({
            "proposal_id": _proposal_id(parent_node_id, proposal),
            "name": proposal.get("canonical_name"),
            "definition": str(proposal.get("definition") or "")[:180],
            "applicability": proposal.get("applicability") or {},
            "source_issue_count": len(proposal.get("member_issue_ids") or []),
        })
    system = """你是广告合规标准问题目录审核员。请在同一个二级问题分支内，把第一轮提案收束为少量稳定、可重复使用的标准审核问题。行业、平台、商品名称、具体示例和措辞差异不是独立建问题的理由，应合并并体现在适用范围或别名中；只有规制行为、核心法律要件或处置逻辑实质不同才保留独立问题。每个proposal_id必须且只能出现一次，不得新增ID。原则上本分支标准问题不超过12个。仅输出JSON。"""
    output = {"standard_issues": [{"canonical_issue_key": "大写英文或下划线", "canonical_name": "中文标准问题名", "member_proposal_ids": ["输入proposal_id"], "definition": "标准定义", "aliases": [], "applicability": {"industries": [], "platforms": []}, "confidence": "high/medium/low", "reason": "合并或分立理由"}]}
    user = {"task": "二级分支内标准问题收束", "parent_node_id": parent_node_id, "max_standard_issues": 12, "proposals": compact, "output": output}
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(user, ensure_ascii=False)}]


def _parse_parent_consolidation(payload, parent_node_id, proposals):
    proposal_by_id = {_proposal_id(parent_node_id, item): item for item in proposals}
    expected = set(proposal_by_id)
    standard = payload.get("standard_issues") or []
    actual = [proposal_id for item in standard for proposal_id in item.get("member_proposal_ids") or []]
    if set(actual) != expected or len(actual) != len(set(actual)):
        raise ValueError(f"proposal coverage mismatch: {parent_node_id}")
    result = []
    for item in standard:
        if item.get("confidence") not in CONFIDENCE:
            raise ValueError("invalid consolidation confidence")
        members = [proposal_by_id[proposal_id] for proposal_id in item.get("member_proposal_ids") or []]
        result.append({
            "canonical_issue_key": item.get("canonical_issue_key"),
            "canonical_name": item.get("canonical_name"),
            "parent_node_id": parent_node_id,
            "member_issue_ids": _dedupe(old_id for proposal in members for old_id in proposal.get("member_issue_ids") or []),
            "relation": "merge",
            "definition": item.get("definition") or "",
            "aliases": _dedupe(alias for proposal in members for alias in ([proposal.get("canonical_name")] + list(proposal.get("aliases") or []))),
            "applicability": item.get("applicability") or {"industries": [], "platforms": []},
            "confidence": item.get("confidence"),
            "reason": item.get("reason") or "",
            "member_proposal_ids": item.get("member_proposal_ids") or [],
        })
    return result


def consolidate_parent_proposals(client, parent_node_id, proposals, checkpoint):
    messages = parent_consolidation_messages(parent_node_id, proposals)
    digest = hashlib.sha256(json.dumps(messages, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    checkpoint = Path(checkpoint)
    if checkpoint.exists():
        saved = read_json(checkpoint)
        if saved.get("digest") == digest:
            return _parse_parent_consolidation(saved["payload"], parent_node_id, proposals)
    payload = _call_json(client, messages)
    try:
        result = _parse_parent_consolidation(payload, parent_node_id, proposals)
    except ValueError as exc:
        proposal_ids = [_proposal_id(parent_node_id, proposal) for proposal in proposals]
        repair = messages + [
            {"role": "assistant", "content": json.dumps(payload, ensure_ascii=False)},
            {"role": "user", "content": json.dumps({"task": "修复proposal_id覆盖", "error": str(exc), "required_proposal_ids": proposal_ids, "requirements": ["每个ID恰好出现一次", "不得新增ID", "返回完整JSON"]}, ensure_ascii=False)},
        ]
        payload = _call_json(client, repair)
        result = _parse_parent_consolidation(payload, parent_node_id, proposals)
    atomic_json(checkpoint, {"digest": digest, "payload": payload})
    return result


def consolidate_all_proposals(client, proposals, checkpoint_dir, workers=6):
    by_parent = defaultdict(list)
    for proposal in proposals:
        by_parent[proposal.get("parent_node_id")].append(proposal)
    results = []
    def job(item):
        parent_node_id, parent_proposals = item
        filename = hashlib.sha256(str(parent_node_id).encode("utf-8")).hexdigest()[:16] + ".json"
        return consolidate_parent_proposals(client, parent_node_id, parent_proposals, Path(checkpoint_dir) / filename)
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = [executor.submit(job, item) for item in sorted(by_parent.items())]
        for future in as_completed(futures):
            results.extend(future.result())
    return sorted(results, key=lambda item: (item.get("parent_node_id") or "", item.get("canonical_name") or ""))

def run_pipeline(client, issue_asset, mapping_asset, records_by_uid, output_dir, assets_dir, batch_size=10, workers=6):
    output_dir = Path(output_dir)
    assets_dir = Path(assets_dir)
    checkpoints = output_dir / "checkpoints"
    issue_rows = build_issue_inputs(issue_asset, mapping_asset, records_by_uid)
    fragments = []
    for index in range(0, len(issue_rows), 25):
        batch = issue_rows[index:index + 25]
        payload = _checkpoint_call(client, skeleton_messages(batch), checkpoints / "skeleton_fragments" / f"{index // 25:03d}.json")
        fragments.append(payload)
    skeleton = build_seed_skeleton(len(issue_rows), len(fragments))
    atomic_json(assets_dir / "legal_issue_taxonomy_skeleton_v0.1.json", skeleton)

    fingerprint_batches = [(index, issue_rows[index:index + batch_size]) for index in range(0, len(issue_rows), batch_size)]
    fingerprints = []
    def fingerprint_job(job):
        index, batch = job
        expected = {item["issue_id"] for item in batch}
        return _checkpoint_call(client, fingerprint_messages(batch, skeleton), checkpoints / "fingerprints" / f"{index // batch_size:04d}.json", lambda payload: parse_fingerprint_batch(payload, expected))
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {executor.submit(fingerprint_job, job): job[0] for job in fingerprint_batches}
        for future in as_completed(futures):
            fingerprints.extend(future.result())
    fingerprints.sort(key=lambda item: item["issue_id"])
    fingerprint_asset = {"schema_version": "0.1", "asset_status": "draft", "generated_at": datetime.now(timezone.utc).isoformat(), "prompt_version": PROMPT_VERSION, "fingerprints": fingerprints}
    atomic_json(assets_dir / "issue_fingerprints_draft_v0.1.json", fingerprint_asset)

    groups = build_candidate_groups(fingerprints)
    proposals = []
    def cluster_job(group):
        checkpoint = checkpoints / "clusters_v2" / (hashlib.sha256(group["group_id"].encode()).hexdigest()[:16] + ".json")
        return cluster_with_coverage_retry(client, group, checkpoint)
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {executor.submit(cluster_job, group): group["group_id"] for group in groups}
        for future in as_completed(futures):
            proposals.extend(future.result())
    proposal_asset = {"schema_version": "0.1", "asset_status": "draft", "generated_at": datetime.now(timezone.utc).isoformat(), "prompt_version": PROMPT_VERSION, "candidate_group_count": len(groups), "proposals": proposals}
    atomic_json(assets_dir / "issue_cluster_proposals_v0.1.json", proposal_asset)
    consolidated = consolidate_all_proposals(client, proposals, checkpoints / "consolidation_v1", workers=workers)
    consolidation_asset = {"schema_version": "0.1", "asset_status": "draft", "generated_at": datetime.now(timezone.utc).isoformat(), "prompt_version": PROMPT_VERSION, "input_proposal_count": len(proposals), "standard_issue_count": len(consolidated), "proposals": consolidated}
    atomic_json(assets_dir / "issue_cluster_consolidation_v0.1.json", consolidation_asset)
    taxonomy, mappings = assemble_taxonomy(skeleton, consolidated, issue_asset, mapping_asset)
    validation = validate_taxonomy_result(taxonomy, mappings, set(records_by_uid))
    warnings = [] if validation["valid"] else [{"severity": "high", "code": "TAXONOMY_VALIDATION", "record_id": "taxonomy", "message": message} for message in validation["errors"]]
    atomic_json(assets_dir / "legal_issue_taxonomy_draft_v0.2.json", taxonomy)
    atomic_json(assets_dir / "rule_issue_mapping_draft_v0.2.json", mappings)
    export_merge_review_workbook(output_dir / "问题目录归并人工审核表_v0.2.xlsx", skeleton, taxonomy, mappings, consolidated, issue_asset, records_by_uid, warnings)
    summary = {
        "input_issue_count": len(issue_rows),
        "input_rule_count": len(records_by_uid),
        "skeleton_node_count": len(skeleton.get("nodes") or []),
        "candidate_group_count": len(groups),
        "first_stage_proposal_count": len(proposals),
        "canonical_issue_count": sum(item.get("node_type") == "legal_issue" for item in taxonomy.get("issues") or []),
        "projected_mapping_count": len(mappings.get("mappings") or []),
        "validation": validation,
    }
    atomic_json(output_dir / "generation_summary.json", summary)
    return summary
