# -*- coding: utf-8 -*-

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from export_specialized_rule_audit_v02 import build_audit_rows, export_workbook


ROOT = Path(__file__).resolve().parents[1]


class SpecializedRuleAuditV02Tests(unittest.TestCase):
    def test_build_rows_uses_v02_mapping_and_keeps_historical_signals(self):
        audit = build_audit_rows(ROOT)
        self.assertGreater(len(audit["all_rules"]), 800)
        self.assertGreater(len(audit["exact_duplicates"]), 0)
        self.assertGreater(len(audit["track_conflicts"]), 0)
        self.assertGreater(len(audit["excluded_candidates"]), 0)
        self.assertTrue(all(row.get("规则原文") for row in audit["all_rules"]))
        self.assertTrue(all("当前问题" in row for row in audit["track_conflicts"]))

    def test_export_has_specialized_sheets_and_review_columns(self):
        required = {
            "审核说明", "全部规则索引", "完全重复规范", "同一规则跨目录重叠",
            "赛道污染", "问题错配复核", "移出普通文案树", "移出规则库候选",
            "主动补资料候选", "中低置信度", "校验结果",
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "audit.xlsx"
            export_workbook(ROOT, output)
            workbook = load_workbook(output, read_only=False, data_only=True)
            self.assertTrue(required.issubset(set(workbook.sheetnames)))
            for name in required - {"审核说明", "校验结果"}:
                headers = [cell.value for cell in workbook[name][1]]
                self.assertIn("人工结论", headers)
                self.assertIn("备注", headers)
            self.assertIn("重复目标rule_uid", [cell.value for cell in workbook["完全重复规范"][1]])
            self.assertIn("建议赛道", [cell.value for cell in workbook["赛道污染"][1]])
            self.assertIn("调整后问题", [cell.value for cell in workbook["问题错配复核"][1]])


if __name__ == "__main__":
    unittest.main()
