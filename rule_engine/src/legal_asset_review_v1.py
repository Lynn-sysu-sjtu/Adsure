# -*- coding: utf-8 -*-
"""Machine quality checks and human-review workbook export for draft assets."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


DOMAIN_MARKERS = (
    "概率", "抽卡", "掉落", "爆率", "版号", "未成年人", "充值", "付费",
    "美白", "祛斑", "防脱", "特殊化妆品", "注册证", "备案",
    "疾病", "治疗", "预防", "药物", "降血糖", "降血压", "警示语",
    "数据", "实验", "检测", "代言", "赠送", "福利", "奖励品",
)


def _json_text(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _authoritative_rule_text(rule):
    return _json_text({"title": rule.get("title"), "dimension": rule.get("dimension"), "legal_basis": rule.get("legal_basis"), "rule_applicability": rule.get("rule_applicability"), "fact_check": rule.get("fact_check")})


def assess_candidate_quality(issue_asset, mapping_asset, check_asset, records_by_uid):
    warnings = []
    for issue in issue_asset.get("issues") or []:
        if issue.get("track") == "通用" or str(issue.get("issue_id") or "").startswith("GEN."):
            for uid in issue.get("candidate_rule_uids") or []:
                record = records_by_uid.get(uid) or {}
                if record.get("track") in {"游戏", "美妆", "保健食品"}:
                    warnings.append({"severity": "warning", "code": "TRACK_RULE_MAPPED_TO_GEN", "record_id": issue.get("issue_id"), "rule_uid": uid, "message": f"{record.get('track')}目录规则被映射为通用问题，需人工确认是否确属跨赛道规则。"})
    for check in check_asset.get("checks") or []:
        check_text = _json_text({"name": check.get("name"), "requirement": check.get("requirement"), "required_materials": check.get("required_materials")})
        for uid in check.get("basis_rule_uids") or []:
            record = records_by_uid.get(uid) or {}; source_text = _authoritative_rule_text(record.get("rule") or {})
            unsupported = [marker for marker in DOMAIN_MARKERS if marker in check_text and marker not in source_text]
            if unsupported:
                warnings.append({"severity": "high", "code": "UNSUPPORTED_PROACTIVE_TOPIC", "record_id": check.get("check_id"), "rule_uid": uid, "message": "主动核查出现未被规则标题、法条、适用条件或事实核验字段直接支持的主题：" + "、".join(unsupported)})
    return warnings


def _sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    return sheet



def _dedupe(values):
    seen = set()
    result = []
    for value in values:
        value = str(value or "").strip()
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def _rule_source_details(record):
    rule = record.get("rule") or {}
    sources = record.get("legal_sources") or []
    source_names = {
        str(item.get("id") or "").strip(): str(item.get("name") or "").strip()
        for item in sources
        if item.get("id")
    }
    names = []
    locations = []
    original_texts = []
    for basis in rule.get("legal_basis") or []:
        source_id = str(basis.get("source_id") or "").strip()
        source_name = source_names.get(source_id) or source_id
        if source_name:
            names.append(source_name)
        article = str(basis.get("article") or "").strip()
        location = " ".join(part for part in (source_name, article) if part)
        if location:
            locations.append(location)
        original = basis.get("text")
        if original in (None, ""):
            original = basis.get("quote")
        if original not in (None, ""):
            original_texts.append(str(original).strip())
    if not names:
        names.extend(item.get("name") for item in sources)
    return {
        "source_names": "\n".join(_dedupe(names)),
        "locations": "\n".join(_dedupe(locations)),
        "original_texts": "\n".join(_dedupe(original_texts)),
    }


def _build_issue_rule_review_rows(issue_asset, mapping_asset, records_by_uid):
    issues = {item.get("issue_id"): item for item in issue_asset.get("issues") or [] if item.get("issue_id")}
    mappings = {item.get("rule_uid"): item for item in mapping_asset.get("mappings") or [] if item.get("rule_uid")}
    pairs = set()
    for issue_id, issue in issues.items():
        for rule_uid in issue.get("candidate_rule_uids") or []:
            pairs.add((issue_id, rule_uid))
    for rule_uid, mapping in mappings.items():
        if mapping.get("primary_issue_id"):
            pairs.add((mapping.get("primary_issue_id"), rule_uid))
        for issue_id in mapping.get("secondary_issue_ids") or []:
            pairs.add((issue_id, rule_uid))

    rows = []
    for issue_id, rule_uid in sorted(pairs, key=lambda item: (str(item[0]), str(item[1]))):
        issue = issues.get(issue_id) or {}
        mapping = mappings.get(rule_uid) or {}
        record = records_by_uid.get(rule_uid) or {}
        rule = record.get("rule") or {}
        declared_by_issue = rule_uid in (issue.get("candidate_rule_uids") or [])
        if mapping.get("primary_issue_id") == issue_id:
            relationship = "主问题"
            declared_by_mapping = True
        elif issue_id in (mapping.get("secondary_issue_ids") or []):
            relationship = "次问题"
            declared_by_mapping = True
        else:
            relationship = "未在规则映射声明"
            declared_by_mapping = False
        if declared_by_issue and declared_by_mapping:
            consistency = "一致"
        elif declared_by_issue:
            consistency = "仅问题目录声明"
        else:
            consistency = "仅规则映射声明"
        source = _rule_source_details(record)
        rows.append([
            issue_id,
            issue.get("track") or record.get("track"),
            issue.get("name"),
            " > ".join(issue.get("issue_path") or []),
            issue.get("definition"),
            "；".join(issue.get("in_scope") or []),
            "；".join(issue.get("out_of_scope") or []),
            rule_uid,
            mapping.get("rule_id") or rule.get("rule_id"),
            relationship,
            consistency,
            rule.get("title"),
            source["source_names"],
            source["locations"],
            source["original_texts"],
            record.get("source_file"),
            _json_text(mapping.get("elements") or []),
            _json_text(mapping.get("evidence_policy") or {}),
            mapping.get("default_terminal_outcome"),
            mapping.get("mapping_confidence"),
            "",
            "",
        ])
    return rows

def export_review_workbook(path, issue_asset, mapping_asset, check_asset, warnings, records_by_uid):
    workbook = Workbook(); workbook.remove(workbook.active)
    issue_rows = [[item.get("issue_id"), item.get("track"), item.get("name"), item.get("parent_issue_id"), " > ".join(item.get("issue_path") or []), item.get("definition"), "；".join(item.get("in_scope") or []), "；".join(item.get("out_of_scope") or []), "；".join(item.get("candidate_rule_uids") or []), (item.get("review") or {}).get("status"), "", ""] for item in issue_asset.get("issues") or []]
    _sheet(workbook, "问题目录", ["issue_id", "赛道", "问题名称", "父节点", "问题路径", "定义", "纳入情形", "排除情形", "rule_uid", "当前状态", "人工结论", "审核备注"], issue_rows)
    mapping_rows = []
    for item in mapping_asset.get("mappings") or []:
        record = records_by_uid.get(item.get("rule_uid")) or {}; rule = record.get("rule") or {}
        mapping_rows.append([item.get("rule_uid"), item.get("rule_id"), record.get("track"), rule.get("title"), item.get("primary_issue_id"), "；".join(item.get("secondary_issue_ids") or []), _json_text(item.get("elements") or []), _json_text(item.get("evidence_policy") or {}), item.get("default_terminal_outcome"), _json_text(rule.get("legal_basis") or []), item.get("mapping_confidence"), "", ""])
    _sheet(workbook, "规则映射", ["rule_uid", "rule_id", "赛道", "规则标题", "主问题", "次问题", "成立要件", "证据权限", "默认终止结果", "原始法律依据", "模型置信度", "人工结论", "审核备注"], mapping_rows)
    merged_rows = _build_issue_rule_review_rows(issue_asset, mapping_asset, records_by_uid)
    _sheet(workbook, "问题目录-规则原文对照", ["issue_id", "赛道", "问题名称", "问题路径", "定义", "纳入情形", "排除情形", "rule_uid", "rule_id", "映射关系", "映射一致性", "规则标题", "规则来源名称", "法条定位", "原始规则原文", "JSON源文件", "成立要件", "证据权限", "默认终止结果", "模型置信度", "人工结论", "审核备注"], merged_rows)
    check_rows = [[item.get("check_id"), item.get("track"), item.get("name"), item.get("check_type"), "；".join(item.get("trigger_issue_ids") or []), item.get("requirement"), "；".join(item.get("required_materials") or []), "；".join(item.get("basis_rule_uids") or []), _json_text(item.get("applicability") or {}), _json_text(item.get("trigger_conditions") or {}), "", ""] for item in check_asset.get("checks") or []]
    _sheet(workbook, "主动核查", ["check_id", "赛道", "名称", "类型", "触发问题", "核查要求", "所需材料", "依据规则", "适用范围", "触发条件", "人工结论", "审核备注"], check_rows)
    warning_rows = [[item.get("severity"), item.get("code"), item.get("record_id"), item.get("rule_uid"), item.get("message"), ""] for item in warnings]
    _sheet(workbook, "质检告警", ["严重度", "代码", "记录", "rule_uid", "说明", "处理结论"], warning_rows)
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True); workbook.save(path)
    return path
