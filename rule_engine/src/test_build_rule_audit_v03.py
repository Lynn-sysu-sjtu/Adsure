# -*- coding: utf-8 -*-

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from build_rule_audit_v03 import build_v03_review, export_v03_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "approved_issue_tree_v02" / "全部规则专项人工审核表_v0.2.xlsx"


class RuleAuditV03Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.result = build_v03_review(ROOT, SOURCE)

    def test_human_decisions_are_loaded(self):
        self.assertEqual(len(self.result["decisions"]["remove_uids"]), 60)
        self.assertEqual(len(self.result["decisions"]["track_updates"]), 17)
        self.assertEqual(len(self.result["decisions"]["duplicate_groups"]), 103)

    def test_removed_rules_do_not_enter_cleaned_assets(self):
        removed = self.result["decisions"]["remove_uids"]
        current = self.result["active"] + self.result["excluded"]
        self.assertFalse(removed & {item["rule_uid"] for item in current})

    def test_duplicate_groups_resolve_to_one_uid(self):
        current_uids = {
            item["rule_uid"] for item in self.result["active"] + self.result["excluded"]
        }
        for group in self.result["decisions"]["duplicate_groups"].values():
            survivors = set(group["uids"]) - self.result["decisions"]["remove_uids"]
            self.assertLessEqual(len(survivors & current_uids), 1)
        self.assertEqual(self.result["residual_duplicate_groups"], [])

    def test_track_updates_are_applied_to_surviving_uid(self):
        by_uid = {}
        for item in self.result["active"] + self.result["excluded"]:
            by_uid.setdefault(item["rule_uid"], set()).add(item["track"])
        for record in self.result["track_records"]:
            if record["处理结果"] == "已删除":
                continue
            self.assertEqual(by_uid[record["最终规则UID"]], {record["调整后赛道"]})

    def test_export_contains_trace_and_validation_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "audit.xlsx"
            export_v03_workbook(self.result, output)
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(workbook.sheetnames[2], "完全重复规范")
            self.assertEqual(workbook.sheetnames[4], "赛道污染")
            self.assertIn("不纳入规则库", workbook.sheetnames)
            self.assertIn("校验结果", workbook.sheetnames)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
