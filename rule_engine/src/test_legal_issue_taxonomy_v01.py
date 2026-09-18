# -*- coding: utf-8 -*-
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from legal_issue_taxonomy_v01 import (
    _call_json,
    assemble_taxonomy,
    build_candidate_groups,
    cluster_with_coverage_retry,
    consolidate_parent_proposals,
    build_issue_inputs,
    build_seed_skeleton,
    export_merge_review_workbook,
    parse_fingerprint_batch,
    consolidate_skeleton_messages,
    reduce_skeleton_fragments,
    validate_taxonomy_result,
)


class LegalIssueTaxonomyV01Tests(unittest.TestCase):
    def setUp(self):
        self.issue_asset = {"issues": [
            {"issue_id": "COSM.ABS.ONE", "name": "功效绝对化用语", "definition": "使用绝对化措辞宣传功效", "in_scope": [], "out_of_scope": [], "candidate_rule_uids": ["R1"]},
            {"issue_id": "COSM.ABS.TWO", "name": "绝对化功效宣传", "definition": "对产品功效作绝对保证", "in_scope": [], "out_of_scope": [], "candidate_rule_uids": ["R2"]},
        ]}
        self.mapping_asset = {"mappings": [
            {"rule_uid": "R1", "rule_id": "C1", "primary_issue_id": "COSM.ABS.ONE", "secondary_issue_ids": []},
            {"rule_uid": "R2", "rule_id": "C2", "primary_issue_id": "COSM.ABS.TWO", "secondary_issue_ids": []},
        ]}
        self.records = {
            "R1": {"track": "美妆", "source_file": "美妆/a.json", "legal_sources": [{"id": "AL", "name": "中华人民共和国广告法"}], "rule": {"rule_uid": "R1", "title": "不得使用绝对化用语", "legal_basis": [{"source_id": "AL", "article": "第九条", "text": "不得使用国家级、最高级、最佳等用语。"}]}},
            "R2": {"track": "美妆", "source_file": "美妆/b.json", "legal_sources": [{"id": "P", "name": "平台美妆规则"}], "rule": {"rule_uid": "R2", "title": "不得绝对保证功效", "legal_basis": [{"source_id": "P", "article": "第二条", "text": "不得绝对保证商品功效。"}]}},
        }

    def test_build_issue_inputs_hydrates_sources_by_rule_uid(self):
        rows = build_issue_inputs(self.issue_asset, self.mapping_asset, self.records)
        self.assertEqual(2, len(rows))
        self.assertEqual(["R1"], rows[0]["rule_uids"])
        self.assertIn("不得使用国家级", rows[0]["source_rules"][0]["original_text"])
        self.assertEqual("中华人民共和国广告法", rows[0]["source_rules"][0]["source_name"])

    def test_parse_fingerprint_batch_rejects_unknown_issue(self):
        payload = {"fingerprints": [{"issue_id": "UNKNOWN"}]}
        with self.assertRaisesRegex(ValueError, "unknown issue_id"):
            parse_fingerprint_batch(payload, {"COSM.ABS.ONE"})

    def test_call_json_asks_deepseek_to_repair_malformed_json_once(self):
        class RepairClient:
            def __init__(self):
                self.calls = []

            def create_chat_completion(self, messages, **kwargs):
                self.calls.append(messages)
                content = '{"nodes": [' if len(self.calls) == 1 else '{"nodes": []}'
                return {"choices": [{"message": {"content": content}}]}

        client = RepairClient()
        result = _call_json(client, [{"role": "user", "content": "生成JSON"}])
        self.assertEqual({"nodes": []}, result)
        self.assertEqual(2, len(client.calls))
        self.assertEqual("assistant", client.calls[1][-2]["role"])
        self.assertIn("修复", client.calls[1][-1]["content"])

    def test_skeleton_consolidation_prompt_enforces_compact_tree(self):
        messages = consolidate_skeleton_messages([{"nodes": []}])
        payload = json.loads(messages[1]["content"])
        self.assertEqual(60, payload["constraints"]["max_total_nodes"])
        self.assertEqual(14, payload["constraints"]["max_level_1_nodes"])
        self.assertIn("不得逐规则建节点", messages[0]["content"])
        node_contract = payload["output"]["nodes"][0]
        self.assertEqual({"node_id", "parent_node_id", "level", "name"}, set(node_contract))
        self.assertIn("只能包含四个字段", messages[0]["content"])

    def test_skeleton_fragments_are_reduced_hierarchically(self):
        class SkeletonClient:
            def __init__(self):
                self.calls = 0

            def create_chat_completion(self, messages, **kwargs):
                self.calls += 1
                return {"choices": [{"message": {"content": json.dumps({"nodes": [{"node_id": "CLAIM", "parent_node_id": None, "level": 1, "name": "宣称表达"}]}, ensure_ascii=False)}}]}

        with tempfile.TemporaryDirectory() as tmp:
            client = SkeletonClient()
            result = reduce_skeleton_fragments(client, [{"nodes": []}] * 15, Path(tmp), chunk_size=7)
            self.assertEqual(4, client.calls)
            self.assertEqual("0.1", result["schema_version"])

    def test_seed_skeleton_uses_legal_domains_not_tracks_as_roots(self):
        skeleton = build_seed_skeleton(issue_count=1214, scan_fragment_count=49)
        roots = [node for node in skeleton["nodes"] if node["level"] == 1]
        names = {node["name"] for node in roots}
        self.assertIn("宣称表达方式", names)
        self.assertIn("证明材料与事实核验", names)
        self.assertNotIn("游戏", names)
        self.assertNotIn("美妆", names)
        self.assertTrue(all(node["parent_node_id"] is None for node in roots))
        self.assertEqual(1214, skeleton["source_issue_count"])

    def test_candidate_groups_put_similar_fingerprints_together(self):
        fingerprints = [
            {"issue_id": "COSM.ABS.ONE", "skeleton_leaf_key": "CLAIM.ABSOLUTE", "canonical_problem_name": "功效绝对化宣称", "regulated_behavior": "绝对化宣传", "claim_object": "功效"},
            {"issue_id": "COSM.ABS.TWO", "skeleton_leaf_key": "CLAIM.ABSOLUTE", "canonical_problem_name": "功效绝对化宣传", "regulated_behavior": "绝对化宣传", "claim_object": "功效"},
        ]
        groups = build_candidate_groups(fingerprints)
        self.assertEqual(1, len(groups))
        self.assertEqual({"COSM.ABS.ONE", "COSM.ABS.TWO"}, set(groups[0]["issue_ids"]))

    def test_cluster_retry_repairs_missing_issue_coverage(self):
        group = {
            "group_id": "CLAIM.ABS#1",
            "skeleton_leaf_key": "CLAIM.ABS",
            "issue_ids": ["A", "B"],
            "fingerprints": [{"issue_id": "A"}, {"issue_id": "B"}],
        }

        class CoverageClient:
            def __init__(self):
                self.calls = []

            def create_chat_completion(self, messages, **kwargs):
                self.calls.append(messages)
                members = ["A"] if len(self.calls) == 1 else ["A", "B"]
                payload = {"clusters": [{"canonical_issue_key": "ABS", "canonical_name": "绝对化用语", "parent_node_id": "CLAIM.ABS", "member_issue_ids": members, "relation": "merge", "definition": "", "aliases": [], "applicability": {"industries": [], "platforms": []}, "confidence": "high", "reason": ""}]}
                return {"choices": [{"message": {"content": json.dumps(payload, ensure_ascii=False)}}]}

        with tempfile.TemporaryDirectory() as tmp:
            client = CoverageClient()
            result = cluster_with_coverage_retry(client, group, Path(tmp) / "cluster.json")
            self.assertEqual(["A", "B"], result[0]["member_issue_ids"])
            self.assertEqual(2, len(client.calls))
            self.assertIn("缺失", client.calls[1][-1]["content"])

    def test_candidate_groups_cap_deepseek_batch_at_twelve(self):
        fingerprints = [
            {"issue_id": f"I{index:02d}", "skeleton_leaf_key": "DISCLOSURE.LABEL", "regulated_behavior": "披露", "claim_object": "标签"}
            for index in range(25)
        ]
        groups = build_candidate_groups(fingerprints)
        self.assertEqual([12, 12, 1], [len(group["issue_ids"]) for group in groups])

    def test_second_stage_consolidates_proposals_and_preserves_old_issues(self):
        proposals = [
            {"canonical_issue_key": "A", "canonical_name": "美妆广告绝对化用语", "parent_node_id": "CLAIM.ABS", "member_issue_ids": ["OLD1"], "definition": "绝对化宣传", "confidence": "high", "reason": ""},
            {"canonical_issue_key": "B", "canonical_name": "保健食品使用极限词", "parent_node_id": "CLAIM.ABS", "member_issue_ids": ["OLD2"], "definition": "绝对化宣传", "confidence": "high", "reason": ""},
        ]

        class ConsolidationClient:
            def create_chat_completion(self, messages, **kwargs):
                request = json.loads(messages[1]["content"])
                ids = [item["proposal_id"] for item in request["proposals"]]
                payload = {"standard_issues": [{"canonical_issue_key": "ABSOLUTE_TERMS", "canonical_name": "使用绝对化用语", "member_proposal_ids": ids, "definition": "广告使用绝对化用语", "aliases": [], "applicability": {"industries": [], "platforms": []}, "confidence": "high", "reason": "核心规制行为相同"}]}
                return {"choices": [{"message": {"content": json.dumps(payload, ensure_ascii=False)}}]}

        with tempfile.TemporaryDirectory() as tmp:
            result = consolidate_parent_proposals(ConsolidationClient(), "CLAIM.ABS", proposals, Path(tmp) / "result.json")
            self.assertEqual(1, len(result))
            self.assertEqual({"OLD1", "OLD2"}, set(result[0]["member_issue_ids"]))
            self.assertEqual("使用绝对化用语", result[0]["canonical_name"])

    def test_assembly_merges_old_issues_and_preserves_rules(self):
        skeleton = {"nodes": [
            {"node_id": "CLAIM", "parent_node_id": None, "level": 1, "name": "宣称表达"},
            {"node_id": "CLAIM.ABSOLUTE", "parent_node_id": "CLAIM", "level": 2, "name": "绝对化用语"},
        ]}
        proposals = [{"canonical_issue_key": "EFFICACY", "canonical_name": "功效绝对化宣称", "parent_node_id": "CLAIM.ABSOLUTE", "member_issue_ids": ["COSM.ABS.ONE", "COSM.ABS.TWO"], "confidence": "high", "reason": "规制行为和对象相同"}]
        taxonomy, mappings = assemble_taxonomy(skeleton, proposals, self.issue_asset, self.mapping_asset)
        leaves = [x for x in taxonomy["issues"] if x.get("node_type") == "legal_issue"]
        self.assertEqual(1, len(leaves))
        self.assertEqual({"R1", "R2"}, set(leaves[0]["candidate_rule_uids"]))
        self.assertEqual(2, len(mappings["mappings"]))
        self.assertTrue(validate_taxonomy_result(taxonomy, mappings, {"R1", "R2"})["valid"])

    def test_workbook_contains_tree_and_merge_review(self):
        skeleton = {"nodes": [{"node_id": "CLAIM", "parent_node_id": None, "level": 1, "name": "claim"}]}
        taxonomy = {"issues": [
            {"issue_id": "CLAIM.ABS", "parent_issue_id": "CLAIM", "level": 2, "node_type": "legal_issue", "name": "absolute", "candidate_rule_uids": ["R1"]},
            {"issue_id": "CLAIM.PROOF", "parent_issue_id": "CLAIM", "level": 2, "node_type": "legal_issue", "name": "proof", "candidate_rule_uids": ["R1"]},
        ]}
        mappings = {"mappings": [{"rule_uid": "R1", "primary_issue_id": "CLAIM.ABS", "secondary_issue_ids": ["CLAIM.PROOF"], "source_issue_ids": ["COSM.ABS.ONE"]}]}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            export_merge_review_workbook(path, skeleton, taxonomy, mappings, [], self.issue_asset, self.records, [])
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(5, len(workbook.sheetnames))
                source_rows = list(workbook.worksheets[2].iter_rows(min_row=2, values_only=True))
                self.assertEqual({"CLAIM.ABS", "CLAIM.PROOF"}, {row[0] for row in source_rows})
                self.assertEqual(["R1", "R1"], sorted(row[2] for row in source_rows))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
