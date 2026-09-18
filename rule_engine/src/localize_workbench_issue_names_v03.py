# -*- coding: utf-8 -*-
"""Replace issue IDs in human-facing workbook columns with strict Chinese names."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports" / "approved_issue_tree_v03"
DEFAULT_SOURCE = REPORT_DIR / "规则并行审核工作台_DeepSeek预审_v0.2.xlsx"
DEFAULT_OUTPUT = REPORT_DIR / "规则并行审核工作台_DeepSeek预审_中文问题名_v0.3.xlsx"
CURRENT_TAXONOMY = REPORT_DIR / "approved_issue_taxonomy_v0.2.json"
LEGACY_ASSETS = ROOT / "assets" / "all_primary_issue_review_draft_v0.1"
REVIEW_SHEETS = ("并行审核主表", "任务包A", "任务包B", "任务包C")


def _node_names(asset):
    names = {}
    for node in asset.get("nodes", []):
        issue_id = str(node.get("issue_id") or "").strip()
        name = str(node.get("name") or "").strip()
        if not issue_id or not name:
            continue
        previous = names.get(issue_id)
        if previous and previous != name:
            raise ValueError(f"conflicting names for {issue_id}: {previous!r} / {name!r}")
        names[issue_id] = name
    return names


def build_issue_name_map(current_taxonomy, legacy_assets):
    """Use current approved names first; legacy nodes may only fill removed IDs."""
    current = _node_names(current_taxonomy)
    legacy_candidates = {}
    for asset in legacy_assets:
        for issue_id, name in _node_names(asset).items():
            if issue_id in current:
                continue
            legacy_candidates.setdefault(issue_id, set()).add(name)
    conflicts = {
        issue_id: sorted(names)
        for issue_id, names in legacy_candidates.items()
        if len(names) > 1
    }
    if conflicts:
        raise ValueError(f"conflicting legacy issue names: {conflicts}")
    result = dict(current)
    result.update({issue_id: next(iter(names)) for issue_id, names in legacy_candidates.items()})
    return result


def parse_issue_ids(value):
    return [part.strip() for part in str(value or "").splitlines() if part.strip()]


def map_issue_ids(value, issue_names):
    issue_ids = parse_issue_ids(value)
    missing = [issue_id for issue_id in issue_ids if issue_id not in issue_names]
    if missing:
        raise ValueError("unmapped issue IDs: " + ", ".join(sorted(set(missing))))
    return "\n".join(issue_names[issue_id] for issue_id in issue_ids)


def localize_workbook(source, output, issue_names):
    source = Path(source)
    output = Path(output)
    workbook = load_workbook(source)
    changed_rows = 0
    mapped_ids = set()
    for sheet_name in REVIEW_SHEETS:
        sheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
        required = {
            "当前全部问题", "当前全部问题ID",
            "AI建议保留问题ID", "AI建议删除问题ID",
        }
        missing_headers = required - set(headers)
        if missing_headers:
            raise ValueError(f"{sheet_name} missing headers: {sorted(missing_headers)}")

        current_name_col = headers["当前全部问题"]
        current_id_col = headers["当前全部问题ID"]
        keep_col = headers["AI建议保留问题ID"]
        remove_col = headers["AI建议删除问题ID"]
        sheet.cell(1, keep_col, "AI建议保留问题名称")
        sheet.cell(1, remove_col, "AI建议删除问题名称")
        sheet.column_dimensions[sheet.cell(1, keep_col).column_letter].width = 48
        sheet.column_dimensions[sheet.cell(1, remove_col).column_letter].width = 48

        for row_number in range(2, sheet.max_row + 1):
            current_ids = sheet.cell(row_number, current_id_col).value
            keep_ids = sheet.cell(row_number, keep_col).value
            remove_ids = sheet.cell(row_number, remove_col).value
            for value in (current_ids, keep_ids, remove_ids):
                mapped_ids.update(parse_issue_ids(value))
            sheet.cell(row_number, current_name_col, map_issue_ids(current_ids, issue_names))
            sheet.cell(row_number, keep_col, map_issue_ids(keep_ids, issue_names))
            sheet.cell(row_number, remove_col, map_issue_ids(remove_ids, issue_names))
            changed_rows += 1

    guide = workbook["操作说明"]
    guide.append([
        "问题名称展示",
        "当前全部问题、AI建议保留问题名称和AI建议删除问题名称均由问题ID严格映射为中文名称；"
        "当前全部问题ID继续保留用于追溯，AI原始ID保存在DeepSeek checkpoint。",
    ])
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(output)
    return {"output": str(output), "changed_rows": changed_rows, "mapped_issue_ids": len(mapped_ids)}


def load_assets(current_path=CURRENT_TAXONOMY, legacy_root=LEGACY_ASSETS):
    current = json.loads(Path(current_path).read_text(encoding="utf-8"))
    legacy = [
        json.loads(path.read_text(encoding="utf-8"))
        for path in sorted(Path(legacy_root).glob("*/review_assets.json"))
    ]
    return current, legacy


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    current, legacy = load_assets()
    issue_names = build_issue_name_map(current, legacy)
    result = localize_workbook(args.source, args.output, issue_names)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
