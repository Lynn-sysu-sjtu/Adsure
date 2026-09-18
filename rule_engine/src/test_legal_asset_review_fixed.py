# -*- coding: utf-8 -*-
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from legal_asset_review import assess_candidate_quality, export_review_workbook


class LegalAssetReviewTests(unittest.TestCase):
    def test_flags_track_specific_rule_mapped_to_gen(self):
        records = {"R1": {"track": "游戏", "rule": {"rule_uid": "R1", "title": "游戏版号核验", "legal_basis": [{"quote": "游戏上线应取得批准。"}]}}}
        issues = {"issues": [{"issue_id": "GEN.ACCESS.GAME_LICENSE", "track": "通用", "candidate_rule_uids": ["R1"]}]}
        warnings = assess_candidate_quality(issues, {"mappings": []}, {"checks": []}, records)
        self.assertTrue(any(item["code"] == "TRACK_RULE_MAPPED_TO_GEN" for item in warnings))

    def test_flags_proactive_topic_not_supported_by_rule_source(self):
        records = {"R1": {"track": "游戏", "rule": {"rule_uid": "R1", "title": "禁止虚构使用效果", "legal_basis": [{"quote": "不得虚构使用效果。"}]}}}
        checks = {"checks": [{"check_id": "GAME.FALSE.PROBABILITY", "name": "抽卡概率核查", "requirement": "核验抽卡概率", "required_materials": ["概率公示"], "basis_rule_uids": ["R1"]}]}
        warnings = assess_candidate_quality({"issues": []}, {"mappings": []}, checks, records)
        self.assertTrue(any(item["code"] == "UNSUPPORTED_PROACTIVE_TOPIC" for item in warnings))

    def test_exports_review_workbook_with_merged_issue_rule_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            issues = {"issues": [{
                "issue_id": "HF.CLAIM.TREATMENT",
                "track": "保健食品",
                "name": "疾病治疗宣称",
                "issue_path": ["HF", "CLAIM", "TREATMENT"],
                "definition": "保健食品宣称疾病治疗作用。",
                "in_scope": ["治疗高血压"],
                "out_of_scope": ["一般营养支持"],
                "candidate_rule_uids": ["RUID-HF-1"],
            }]}
            mappings = {"mappings": [{
                "rule_uid": "RUID-HF-1",
                "rule_id": "HF-001",
                "primary_issue_id": "HF.CLAIM.TREATMENT",
                "secondary_issue_ids": [],
                "elements": [],
                "evidence_policy": {},
            }]}
            records = {"RUID-HF-1": {
                "track": "保健食品",
                "source_file": "保健食品/广告法.json",
                "legal_sources": [{"id": "AL", "name": "中华人民共和国广告法"}],
                "rule": {
                    "rule_uid": "RUID-HF-1",
                    "rule_id": "HF-001",
                    "title": "保健食品不得涉及疾病治疗功能",
                    "legal_basis": [{"source_id": "AL", "article": "第十八条", "text": "保健食品广告不得涉及疾病预防、治疗功能。"}],
                },
            }}
            export_review_workbook(path, issues, mappings, {"checks": []}, [], records)
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(5, len(workbook.sheetnames))
                merged = workbook[workbook.sheetnames[2]]
                headers = [cell.value for cell in merged[1]]
                row = dict(zip(headers, [cell.value for cell in merged[2]]))
                self.assertEqual("问题目录-规则原文对照", merged.title)
                self.assertEqual("一致", row["映射一致性"])
                self.assertEqual("主问题", row["映射关系"])
                self.assertEqual("中华人民共和国广告法", row["规则来源名称"])
                self.assertEqual("保健食品广告不得涉及疾病预防、治疗功能。", row["原始规则原文"])
                self.assertEqual("保健食品/广告法.json", row["JSON源文件"])
            finally:
                workbook.close()

    def test_merged_sheet_exposes_one_sided_mapping_declarations(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            issues = {"issues": [{"issue_id": "GAME.A", "candidate_rule_uids": ["R1"]}]}
            mappings = {"mappings": [{"rule_uid": "R1", "primary_issue_id": "GAME.B", "secondary_issue_ids": []}]}
            export_review_workbook(path, issues, mappings, {"checks": []}, [], {"R1": {"rule": {"rule_uid": "R1"}}})
            workbook = load_workbook(path, read_only=True)
            try:
                merged = workbook["问题目录-规则原文对照"]
                headers = [cell.value for cell in merged[1]]
                rows = [dict(zip(headers, [cell.value for cell in row])) for row in merged.iter_rows(min_row=2)]
                self.assertEqual({"仅问题目录声明", "仅规则映射声明"}, {row["映射一致性"] for row in rows})
            finally:
                workbook.close()


if __name__ == "__main__": unittest.main()
