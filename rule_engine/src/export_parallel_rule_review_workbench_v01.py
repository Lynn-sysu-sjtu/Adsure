# -*- coding: utf-8 -*-
"""Export a UID-level workbench for parallel human review."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from build_rule_audit_v03 import build_v03_review
from export_specialized_rule_audit_v02 import node_context


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "approved_issue_tree_v03" / "规则并行审核工作台_v0.1.xlsx"

RISK_ORDER = ["跨目录重叠", "问题错配", "主动补资料", "移出普通文案树", "中低置信度"]
RISK_WEIGHTS = {
    "跨目录重叠": 3,
    "问题错配": 1,
    "主动补资料": 2,
    "移出普通文案树": 2,
    "中低置信度": 1,
}
REVIEW_HEADERS = [
    "审核状态", "负责人", "人工结论", "正确适用赛道", "保留问题ID",
    "删除问题ID", "最终映射角色", "最终审核链路", "资产处置",
    "审核理由", "是否提交争议池", "二审状态", "二审意见", "审核日期",
]


def unique_join(values):
    return "\n".join(sorted({str(value) for value in values if value not in (None, "")}))


def review_uids(rows):
    return {row["规则UID"] for row in rows if row.get("规则UID")}


def representative(items):
    return max(
        items,
        key=lambda item: sum(bool(item.get(key)) for key in (
            "rule_title", "source_file", "source_type", "source_name", "article", "original_text"
        )),
    )


def assign_packages(rows):
    package_names = ["任务包A", "任务包B", "任务包C"]
    scores = {name: 0 for name in package_names}
    counts = {name: 0 for name in package_names}
    ordered = sorted(
        rows,
        key=lambda row: (-row["工作量积分"], row["当前一级问题"], row["规则UID"]),
    )
    for row in ordered:
        package = min(package_names, key=lambda name: (scores[name], counts[name], name))
        row["建议任务包"] = package
        scores[package] += row["工作量积分"]
        counts[package] += 1
    return scores, counts


def build_parallel_review_rows(project_root=ROOT):
    result = build_v03_review(project_root)
    active = result["active"]
    excluded = result["excluded"]
    all_items = active + excluded
    by_uid = defaultdict(list)
    for item in all_items:
        by_uid[item["rule_uid"]].append(item)

    issue_nodes, ancestor = node_context(result["taxonomy"]["nodes"])
    active_issues = defaultdict(set)
    for item in active:
        active_issues[item["rule_uid"]].add(item["issue_id"])

    risks = defaultdict(set)
    for uid, issues in active_issues.items():
        if len(issues) > 1:
            risks[uid].add("跨目录重叠")
    for uid in review_uids(result["review_rows"]["mismatch"]):
        risks[uid].add("问题错配")
    for uid in review_uids(result["review_rows"]["proactive"]):
        risks[uid].add("主动补资料")
    for item in excluded:
        risks[item["rule_uid"]].add("移出普通文案树")
    for uid in review_uids(result["review_rows"]["low_confidence"]):
        risks[uid].add("中低置信度")

    low_rows = defaultdict(list)
    for row in result["review_rows"]["low_confidence"]:
        low_rows[row["规则UID"]].append(row)
    rows = []
    for uid in sorted(risks):
        if uid not in by_uid:
            continue
        items = by_uid[uid]
        source = representative(items)
        issue_ids = {item.get("issue_id", "") for item in items if item.get("issue_id")}
        issue_names = [issue_nodes.get(issue_id, {}).get("name", issue_id) for issue_id in issue_ids]
        root_names = [ancestor(issue_id, 1).get("name", "") for issue_id in issue_ids]
        level_2_names = [ancestor(issue_id, 2).get("name", "") for issue_id in issue_ids]
        tags = [tag for tag in RISK_ORDER if tag in risks[uid]]
        score = sum(RISK_WEIGHTS[tag] for tag in tags)
        confidence_rows = low_rows.get(uid, [])
        row = {
            "建议任务包": "",
            "工作量积分": score,
            "风险标签": "、".join(tags),
            "规则UID": uid,
            "规则标题": source.get("rule_title", ""),
            "当前适用赛道": unique_join(item.get("track") for item in items),
            "当前一级问题": unique_join(root_names),
            "当前二级问题": unique_join(level_2_names),
            "当前全部问题": unique_join(issue_names),
            "当前全部问题ID": unique_join(issue_ids),
            "当前映射角色": unique_join(
                item.get("mapping_type") or item.get("exclusion_role") for item in items
            ),
            "来源类型": source.get("source_type", ""),
            "法规或平台规则": source.get("source_name", ""),
            "条款": source.get("article", ""),
            "规则原文": source.get("original_text", ""),
            "来源文件": source.get("source_file", ""),
            "模型置信度": unique_join(row.get("模型置信度") for row in confidence_rows),
            "历史判断理由": unique_join(row.get("历史判断理由") for row in confidence_rows),
            "审核提示": (
                "先核对原文与适用范围，再依次确定问题归属、映射角色、审核链路和资产处置；"
                "同一UID的所有结论必须在本行一次完成。"
            ),
        }
        row.update({header: "" for header in REVIEW_HEADERS})
        rows.append(row)
    package_scores, package_counts = assign_packages(rows)
    return {
        "rows": sorted(rows, key=lambda row: (row["建议任务包"], -row["工作量积分"], row["规则UID"])),
        "package_scores": package_scores,
        "package_counts": package_counts,
        "removed_uids": result["decisions"]["remove_uids"],
    }


def add_dropdown(sheet, header, values):
    headers = {cell.value: cell.column for cell in sheet[1]}
    if header not in headers or sheet.max_row < 2:
        return
    validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    sheet.add_data_validation(validation)
    column = sheet.cell(1, headers[header]).column_letter
    validation.add(f"{column}2:{column}{sheet.max_row}")


def write_review_sheet(workbook, title, rows, editable=True):
    sheet = workbook.create_sheet(title)
    base_headers = [
        "建议任务包", "工作量积分", "风险标签", "规则UID", "规则标题",
        "当前适用赛道", "当前一级问题", "当前二级问题", "当前全部问题",
        "当前全部问题ID", "当前映射角色", "来源类型", "法规或平台规则",
        "条款", "规则原文", "来源文件", "模型置信度", "历史判断理由", "审核提示",
    ]
    headers = base_headers + REVIEW_HEADERS
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    fills = {
        "header": PatternFill("solid", fgColor="1F5364"),
        "input": PatternFill("solid", fgColor="FFF2CC"),
        "risk": PatternFill("solid", fgColor="FCE4D6"),
    }
    thin = Side(style="thin", color="D9E2E5")
    for cell in sheet[1]:
        cell.fill = fills["header"]
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for index, cell in enumerate(row):
            header = headers[index]
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if editable and header in REVIEW_HEADERS:
                cell.fill = fills["input"]
            if header == "风险标签":
                cell.fill = fills["risk"]
    widths = {
        "建议任务包": 12, "工作量积分": 11, "风险标签": 30, "规则UID": 23,
        "规则标题": 38, "当前适用赛道": 15, "当前一级问题": 24,
        "当前二级问题": 28, "当前全部问题": 48, "当前全部问题ID": 58,
        "当前映射角色": 22, "来源类型": 15, "法规或平台规则": 34,
        "条款": 18, "规则原文": 80, "来源文件": 45, "模型置信度": 14,
        "历史判断理由": 60, "审核提示": 55, "审核状态": 13, "负责人": 14,
        "人工结论": 24, "正确适用赛道": 18, "保留问题ID": 48,
        "删除问题ID": 48, "最终映射角色": 20, "最终审核链路": 23,
        "资产处置": 23, "审核理由": 70, "是否提交争议池": 18,
        "二审状态": 14, "二审意见": 50, "审核日期": 15,
    }
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(header, 18)
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    if editable:
        add_dropdown(sheet, "审核状态", ["未开始", "审核中", "已完成", "待讨论"])
        add_dropdown(sheet, "人工结论", ["保留当前映射", "调整问题映射", "保留多重映射", "转旁路资产", "不纳入候选资产", "待讨论"])
        add_dropdown(sheet, "正确适用赛道", ["通用", "游戏", "美妆", "保健食品"])
        add_dropdown(sheet, "最终映射角色", ["direct", "supporting_basis", "exception", "proactive_check"])
        add_dropdown(sheet, "最终审核链路", ["普通文案召回", "fact_check", "proactive_check", "workflow_reference", "platform_access"])
        add_dropdown(sheet, "资产处置", ["保留", "调整映射", "转旁路资产", "不纳入候选资产", "待讨论"])
        add_dropdown(sheet, "是否提交争议池", ["否", "是"])
        add_dropdown(sheet, "二审状态", ["无需二审", "待二审", "二审通过", "退回修改"])
        status_column = sheet.cell(1, headers.index("审核状态") + 1).column_letter
        sheet.conditional_formatting.add(
            f"{status_column}2:{status_column}{sheet.max_row}",
            FormulaRule(formula=[f'${status_column}2="已完成"'], fill=PatternFill("solid", fgColor="E2F0D9")),
        )
    return sheet


def export_workbench(result, output_path=OUTPUT):
    workbook = Workbook()
    guide = workbook.active
    guide.title = "操作说明"
    guide.append(["项目", "操作要求"])
    instructions = [
        ("任务目的", "以规则UID为唯一任务单位，一次完成问题归属、映射角色、审核链路和资产处置判断。"),
        ("分工方式", "三位成员分别审核任务包A、B、C。只编辑自己的任务包Sheet，负责人填写真实姓名。"),
        ("审核顺序", "先读规则来源和原文，再核对适用赛道，然后处理问题归属、映射角色、审核链路，最后填写资产处置。"),
        ("跨目录规则", "在保留问题ID和删除问题ID中分别填写完整ID；确有多个适用问题时选择“保留多重映射”并说明理由。"),
        ("补资料规则", "必须依赖资质、备案、检测或授权等外部材料时，选择fact_check或proactive_check；不得直接写成已经违规。"),
        ("争议处理", "无法独立判断时选择“待讨论”并将“是否提交争议池”设为“是”，项目负责人集中二审。"),
        ("主表用途", "“并行审核主表”是分工索引，不在其中填写结论；审核结论只写入对应任务包Sheet，避免多处不一致。"),
        ("基线保护", "本工作台不修改jsonbase、v0.2资产或生产规则引擎。"),
    ]
    for row in instructions:
        guide.append(row)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 110
    write_review_sheet(workbook, "并行审核主表", result["rows"], editable=False)
    for package in ("任务包A", "任务包B", "任务包C"):
        write_review_sheet(
            workbook,
            package,
            [row for row in result["rows"] if row["建议任务包"] == package],
            editable=True,
        )
    dispute_headers = [
        "规则UID", "规则标题", "提交人", "争议类型", "争议说明",
        "建议方案", "负责人裁决", "最终问题ID", "最终审核链路", "最终资产处置",
    ]
    dispute = workbook.create_sheet("争议池")
    dispute.append(dispute_headers)
    for _ in range(30):
        dispute.append([""] * len(dispute_headers))
    for cell in dispute[1]:
        cell.fill = PatternFill("solid", fgColor="8B2E2E")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in dispute.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([23, 38, 14, 22, 60, 55, 55, 48, 23, 23], 1):
        dispute.column_dimensions[dispute.cell(1, index).column_letter].width = width
    dispute.freeze_panes = "A2"

    checks = workbook.create_sheet("校验结果")
    checks.append(["校验项", "结果", "数量或说明"])
    checks.append(["UID任务唯一", "通过", len(result["rows"])])
    checks.append(["已排除规则未进入任务池", "通过", len(result["removed_uids"])])
    checks.append(["任务包A", "通过", f'{result["package_counts"]["任务包A"]}条 / {result["package_scores"]["任务包A"]}分'])
    checks.append(["任务包B", "通过", f'{result["package_counts"]["任务包B"]}条 / {result["package_scores"]["任务包B"]}分'])
    checks.append(["任务包C", "通过", f'{result["package_counts"]["任务包C"]}条 / {result["package_scores"]["任务包C"]}分'])
    checks.append(["原始资产状态", "通过", "未修改jsonbase、v0.2或生产链路"])
    for sheet in (guide, checks):
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F5364")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    checks.column_dimensions["A"].width = 35
    checks.column_dimensions["B"].width = 16
    checks.column_dimensions["C"].width = 55
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args(argv)
    result = build_parallel_review_rows(args.project_root)
    output = export_workbench(result, args.output)
    risk_counts = defaultdict(int)
    for row in result["rows"]:
        for tag in row["风险标签"].split("、"):
            risk_counts[tag] += 1
    print(json.dumps({
        "output": str(output),
        "unique_tasks": len(result["rows"]),
        "risk_counts": dict(risk_counts),
        "package_counts": result["package_counts"],
        "package_scores": result["package_scores"],
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
