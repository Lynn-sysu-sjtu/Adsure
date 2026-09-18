# -*- coding: utf-8 -*-

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from localize_workbench_issue_names_v03 import (
    build_issue_name_map,
    localize_workbook,
    map_issue_ids,
)


class LocalizeWorkbenchIssueNamesTests(unittest.TestCase):
    def test_current_taxonomy_has_priority_over_legacy_name(self):
        current = {"nodes": [{"issue_id": "A.B", "name": "现行名称"}]}
        legacy = {"nodes": [{"issue_id": "A.B", "name": "历史名称"}]}
        self.assertEqual("现行名称", build_issue_name_map(current, [legacy])["A.B"])

    def test_multiline_ids_are_strictly_mapped_to_names(self):
        names = {"A.B": "问题甲", "C.D": "问题乙"}
        self.assertEqual("问题甲\n问题乙", map_issue_ids("A.B\nC.D", names))
        with self.assertRaisesRegex(ValueError, "UNKNOWN.ID"):
            map_issue_ids("UNKNOWN.ID", names)

    def test_workbook_replaces_three_display_columns_and_keeps_audit_ids(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            output = Path(directory) / "output.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "并行审核主表"
            headers = [
                "规则UID", "当前全部问题", "当前全部问题ID",
                "AI建议保留问题ID", "AI建议删除问题ID",
            ]
            for sheet_name in ("并行审核主表", "任务包A", "任务包B", "任务包C"):
                target = sheet if sheet_name == "并行审核主表" else workbook.create_sheet(sheet_name)
                target.append(headers)
                target.append(["R1", "A.B", "A.B\nC.D", "C.D", "A.B"])
            workbook.create_sheet("操作说明").append(["项目", "操作要求"])
            workbook.save(source)

            localize_workbook(source, output, {"A.B": "问题甲", "C.D": "问题乙"})

            result = load_workbook(output)
            row = list(result["并行审核主表"].iter_rows(values_only=True))
            self.assertEqual(
                ("规则UID", "当前全部问题", "当前全部问题ID", "AI建议保留问题名称", "AI建议删除问题名称"),
                row[0],
            )
            self.assertEqual(("R1", "问题甲\n问题乙", "A.B\nC.D", "问题乙", "问题甲"), row[1])


if __name__ == "__main__":
    unittest.main()
