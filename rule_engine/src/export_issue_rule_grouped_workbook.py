# -*- coding: utf-8 -*-
"""Export a question-first workbook for reviewing rule-to-issue mappings."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from export_all_primary_issue_tree_html import ROOT_NAMES, ROOT_ORDER


ROOT = Path(__file__).resolve().parents[1]


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def source_text(item, field):
    return "\n\n".join(str(detail.get(field) or "") for detail in item.get("source_details") or [])


def load_mapping_rows(project_root):
    project_root = Path(project_root)
    rows = []
    node_info = {}
    expected = {
        item["rule_uid"]: item
        for item in read_json(project_root / "assets" / "rule_issue_mapping_draft_v0.3.json")["mappings"]
    }
    for root_position, root_id in enumerate(ROOT_ORDER):
        if root_id == "ENDORSEMENT_REVIEW":
            taxonomy = read_json(project_root / "assets" / "endorsement_issue_taxonomy_draft_v0.3.json")
            asset = {"nodes": taxonomy["nodes"], "mappings": read_json(project_root / "assets" / "endorsement_rule_mapping_draft_v0.3.json")["mappings"]}
            proactive = read_json(project_root / "assets" / "endorsement_proactive_checklist_draft_v0.3.json")
            source_root = "ENDORSEMENT"
        else:
            asset = read_json(project_root / "assets" / "all_primary_issue_review_draft_v0.1" / root_id / "review_assets.json")
            source_root = root_id
        nodes = {item["issue_id"]: item for item in asset["nodes"]}
        l2_order = {
            item["issue_id"]: position
            for position, item in enumerate(node for node in asset["nodes"] if node.get("level") == 2)
        }
        l3_order = defaultdict(dict)
        for item in asset["nodes"]:
            if item.get("level") != 3:
                continue
            parent_id = item.get("parent_issue_id") or ""
            l3_order[parent_id][item["issue_id"]] = len(l3_order[parent_id])
        proactive_by_uid = {
            uid: check
            for check in (proactive.get("checks") if root_id == "ENDORSEMENT_REVIEW" else [])
            for uid in check.get("source_rule_uids") or []
        }
        if root_id == "ENDORSEMENT_REVIEW" and proactive_by_uid:
            nodes.setdefault("ENDORSEMENT.PROACTIVE_CHECKLIST", {
                "issue_id": "ENDORSEMENT.PROACTIVE_CHECKLIST",
                "parent_issue_id": "ENDORSEMENT",
                "level": 2,
                "node_type": "directory",
                "name": "主动补资料核查",
            })
        for node in asset["nodes"]:
            if node.get("level") in {2, 3}:
                node_info[node["issue_id"]] = {
                    "root_id": root_id,
                    "root_name": ROOT_NAMES[root_id],
                    "node": node,
                    "source_root": source_root,
                }
        for mapping in asset["mappings"]:
            issue_id = mapping["issue_id"]
            if not issue_id:
                check = proactive_by_uid.get(mapping["rule_uid"], {})
                issue_id = "ENDORSEMENT.PROACTIVE_CHECKLIST." + str(check.get("check_id") or "CHECK")
                nodes.setdefault(issue_id, {
                    "issue_id": issue_id,
                    "parent_issue_id": "ENDORSEMENT.PROACTIVE_CHECKLIST",
                    "level": 3,
                    "node_type": "legal_issue",
                    "name": check.get("name") or "主动补资料核查项",
                    "definition": "场景触发核查义务，但不能仅凭当前文案确认是否已履行。",
                })
            node = nodes.get(issue_id)
            if not node:
                raise ValueError(f"mapping points to unknown issue node: {root_id} / {issue_id}")
            parent = node.get("parent_issue_id")
            parent_node = nodes.get(parent) if parent else None
            if node.get("level") == 3:
                l2 = parent_node or {}
                l2_name = l2.get("name") or parent or ""
                l3_name = node.get("name") or issue_id
            else:
                l2_name = node.get("name") or issue_id
                l3_name = ""
            expected_issue = expected.get(mapping["rule_uid"], {}).get("primary_issue_id") or ""
            expected_root = expected_issue.split(".", 1)[0]
            row = {
                "一级问题ID": root_id,
                "一级问题": ROOT_NAMES[root_id],
                "二级问题ID": parent if node.get("level") == 3 else issue_id,
                "二级问题": l2_name,
                "三级问题ID": issue_id if node.get("level") == 3 else "",
                "三级问题": l3_name,
                "规则UID": mapping["rule_uid"],
                "规则标题": mapping.get("rule_title") or "",
                "规则映射类型": mapping.get("mapping_type") or "",
                "当前映射性质": "主映射" if expected_root == root_id else "跨目录候选",
                "原始赛道": mapping.get("track") or "",
                "来源文件": mapping.get("source_file") or "",
                "来源类型": source_text(mapping, "source_type"),
                "法规/平台规则": source_text(mapping, "source_name"),
                "条款": source_text(mapping, "article"),
                "规则原文": source_text(mapping, "original_text"),
                "模型适用对象": ", ".join((mapping.get("fingerprint") or {}).get("object_scope") or []),
                "法律效果": (mapping.get("fingerprint") or {}).get("legal_effect") or "",
                "模型置信度": (mapping.get("fingerprint") or {}).get("confidence") or "",
                "赛道污染": (mapping.get("fingerprint") or {}).get("source_scope_conflict", False),
                "跨一级问题": (mapping.get("fingerprint") or {}).get("root_scope_conflict", False),
                "建议一级问题": (mapping.get("fingerprint") or {}).get("suggested_root_id") or "",
                "人工问题结论": "",
                "规则对应结论": "",
                "调整后问题": "",
                "确认适用范围": "",
                "备注": "",
                "_sort_key": (
                    root_position,
                    l2_order.get(parent if node.get("level") == 3 else issue_id, len(l2_order)),
                    l3_order.get(parent if node.get("level") == 3 else issue_id, {}).get(
                        issue_id if node.get("level") == 3 else "",
                        len(l3_order.get(parent if node.get("level") == 3 else issue_id, {})),
                    ),
                    issue_id,
                    mapping["rule_uid"],
                ),
            }
            rows.append(row)
    rows.sort(key=lambda item: item["_sort_key"])
    for row in rows:
        row.pop("_sort_key", None)
    return rows


def merge_contiguous_values(sheet, column, start_row, end_row):
    if end_row > start_row:
        sheet.merge_cells(start_row=start_row, start_column=column, end_row=end_row, end_column=column)
        sheet.cell(start_row, column).alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)


def export_workbook(project_root, output_path):
    rows = load_mapping_rows(project_root)
    if not rows:
        raise ValueError("no mapping rows found")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "按问题分组的规则审核"
    headers = [
        "一级问题", "二级问题", "三级问题", "规则UID", "规则标题", "规则映射类型", "当前映射性质",
        "原始赛道", "来源文件", "来源类型", "法规/平台规则", "条款", "规则原文", "模型适用对象",
        "法律效果", "模型置信度", "赛道污染", "跨一级问题", "建议一级问题", "人工问题结论",
        "规则对应结论", "调整后问题", "确认适用范围", "备注",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append([row[header] for header in headers])

    # Merge hierarchy columns within their complete parent path.
    # Rule columns remain one row per mapping occurrence.
    for column in (3, 2, 1):
        start = 2
        previous = tuple(sheet.cell(start, key_column).value for key_column in range(1, column + 1))
        for current in range(3, sheet.max_row + 2):
            value = (
                tuple(sheet.cell(current, key_column).value for key_column in range(1, column + 1))
                if current <= sheet.max_row
                else object()
            )
            if value != previous:
                merge_contiguous_values(sheet, column, start, current - 1)
                start, previous = current, value

    title_fill = PatternFill("solid", fgColor="1F5364")
    hierarchy_fill = PatternFill("solid", fgColor="EAF3F5")
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    thin_gray = Side(style="thin", color="D9E2E5")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    review_headers = {"人工问题结论", "规则对应结论", "调整后问题", "确认适用范围", "备注"}
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        for index, header in enumerate(headers, start=1):
            if header in {"一级问题", "二级问题", "三级问题"}:
                row[index - 1].fill = hierarchy_fill
            if header in review_headers:
                row[index - 1].fill = review_fill
        if row[6].value == "跨目录候选":
            row[6].fill = warning_fill
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    widths = {
        "A": 20, "B": 24, "C": 30, "D": 23, "E": 34, "F": 16, "G": 13, "H": 12,
        "I": 42, "J": 12, "K": 28, "L": 16, "M": 70, "N": 24, "O": 18, "P": 12,
        "Q": 12, "R": 14, "S": 18, "T": 20, "U": 20, "V": 24, "W": 24, "X": 30,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 72

    conclusion_validation = DataValidation(
        type="list",
        formula1='"保留为独立问题,合并到已有问题,移动到其他目录,拆分,无法确定"',
        allow_blank=True,
    )
    rule_validation = DataValidation(
        type="list",
        formula1='"当前规则准确对应,规则需要移出,问题需要补规则,规则不足以支撑问题,无法确定"',
        allow_blank=True,
    )
    sheet.add_data_validation(conclusion_validation)
    sheet.add_data_validation(rule_validation)
    conclusion_validation.add(f"T2:T{sheet.max_row}")
    rule_validation.add(f"U2:U{sheet.max_row}")
    sheet.conditional_formatting.add(f"Q2:Q{sheet.max_row}", FormulaRule(formula=["Q2=TRUE"], fill=warning_fill))
    sheet.conditional_formatting.add(f"R2:R{sheet.max_row}", FormulaRule(formula=["R2=TRUE"], fill=warning_fill))

    guide = workbook.create_sheet("使用说明")
    guide_rows = [
        ["用途", "按问题树分组检查每个问题下面挂接的规则是否准确。"],
        ["阅读方式", "A-C 列是问题层级，已按连续规则行合并；D 列以后每一行是一条具体规则。"],
        ["规则对应结论", "当前规则准确对应 / 规则需要移出 / 问题需要补规则 / 规则不足以支撑问题 / 无法确定。"],
        ["人工问题结论", "保留为独立问题 / 合并到已有问题 / 移动到其他目录 / 拆分 / 无法确定。"],
        ["当前映射性质", "主映射表示该规则的原始主目录；跨目录候选表示同一 rule_uid 还被其他问题目录收纳，需要重点确认。"],
        ["主动补资料", "规则映射类型为 proactive_check 时，不代表文案已经违规，而是需要核验外部事实、资质或证明材料。"],
        ["注意", "不要删除 rule_uid、来源、条款或原文；请把判断填写在黄色审核列，并在备注中写明依据。"],
    ]
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide.freeze_panes = "A1"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {"output": str(output_path), "mapping_rows": len(rows), "unique_rule_uids": len({row["规则UID"] for row in rows})}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=ROOT)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "reports" / "all_primary_issue_fingerprint_review_v01" / "问题-规则对应关系人工审核表_v0.1.xlsx",
    )
    args = parser.parse_args(argv)
    print(json.dumps(export_workbook(args.project_root, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
