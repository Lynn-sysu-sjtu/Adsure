# -*- coding: utf-8 -*-
import argparse
import html
import json
from pathlib import Path

from openpyxl import load_workbook


def export_html(workbook_path, output_path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["标准问题树"]
        headers = [cell.value for cell in sheet[1]]
        rows = [dict(zip(headers, [cell.value for cell in row])) for row in sheet.iter_rows(min_row=2)]
    finally:
        workbook.close()

    children = {}
    for row in rows:
        children.setdefault(row.get("父节点"), []).append(row)
    for items in children.values():
        items.sort(key=lambda item: (int(item.get("层级") or 99), str(item.get("标准问题名称") or "")))

    def esc(value):
        return html.escape(str(value or ""), quote=True)

    def badge(label, cls=""):
        return f'<span class="badge {cls}">{esc(label)}</span>'

    def render(row):
        node_id = row.get("标准问题ID")
        kids = children.get(node_id, [])
        confidence = row.get("模型置信度") or "未标注"
        tone = {"high": "high", "medium": "medium", "low": "low"}.get(confidence, "")
        detail = []
        if row.get("定义"):
            detail.append(f'<p><b>定义</b><br>{esc(row.get("定义"))}</p>')
        if row.get("别名"):
            detail.append(f'<p><b>别名</b><br>{esc(row.get("别名"))}</p>')
        if row.get("模型理由"):
            detail.append(f'<p><b>模型理由</b><br>{esc(row.get("模型理由"))}</p>')
        meta = "".join([
            badge(f"L{row.get('层级')}"),
            badge("目录节点" if row.get("节点类型") == "directory" else "法律问题", "type"),
            badge(f"旧问题 {row.get('旧问题数') or 0}"),
            badge(f"规则 {row.get('规则数') or 0}"),
            badge(f"置信度 {confidence}", tone),
        ])
        control = '<button class="toggle">+</button>' if kids else '<span class="leaf">•</span>'
        child_html = "".join(render(child) for child in kids)
        return f'''<li class="node level-{row.get("层级")} {"branch" if kids else "leaf-node"}" data-confidence="{esc(confidence)}" data-search="{esc((node_id or "") + " " + (row.get("标准问题名称") or "") + " " + (row.get("定义") or ""))}">
<div class="row">{control}<div class="main"><div class="title"><span class="name">{esc(row.get("标准问题名称"))}</span><code>{esc(node_id)}</code></div><div class="meta">{meta}</div></div></div>
<div class="detail">{"".join(detail)}</div>
{f'<ul>{child_html}</ul>' if kids else ''}</li>'''

    roots = [row for row in rows if not row.get("父节点")]
    roots.sort(key=lambda item: str(item.get("标准问题名称") or ""))
    tree_html = "".join(render(row) for row in roots)
    stats = {
        "total": len(rows),
        "directories": sum(row.get("节点类型") == "directory" for row in rows),
        "issues": sum(row.get("节点类型") == "legal_issue" for row in rows),
        "rules": sum(int(row.get("规则数") or 0) for row in rows),
        "high": sum(row.get("模型置信度") == "high" for row in rows),
        "medium": sum(row.get("模型置信度") == "medium" for row in rows),
        "low": sum(row.get("模型置信度") == "low" for row in rows),
    }
    document = f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Adsure 标准问题树 V0.2</title>
<style>
:root{{--ink:#17212b;--muted:#68757e;--line:#dce5ea;--soft:#f5f8fa;--blue:#176b87;--blue-soft:#e7f3f7;--gold:#996b12;--gold-soft:#fff5d9;--red:#a63d3d;--red-soft:#fff0f0;--green:#23734c;--green-soft:#e9f7ef}}
*{{box-sizing:border-box}}body{{margin:0;background:#edf2f5;color:var(--ink);font:14px/1.55 "Microsoft YaHei","Noto Sans CJK SC",sans-serif}}.app{{max-width:1500px;margin:auto;padding:28px 34px 50px}}header,.tree-shell{{background:#fff;border:1px solid var(--line);border-radius:10px;box-shadow:0 4px 16px #21313d0d}}header{{padding:24px 28px}}h1{{margin:0 0 5px;font-size:28px}}.subtitle{{margin:0;color:var(--muted)}}.stats{{display:flex;gap:10px;flex-wrap:wrap;margin-top:18px}}.stat{{min-width:120px;padding:9px 12px;background:var(--soft);border:1px solid var(--line);border-radius:8px}}.stat strong{{display:block;font-size:21px}}.stat span{{color:var(--muted);font-size:12px}}
.toolbar{{position:sticky;top:0;z-index:3;display:flex;gap:9px;align-items:center;flex-wrap:wrap;margin:16px 0 13px;padding:11px;background:#ffffffed;border:1px solid var(--line);border-radius:8px;backdrop-filter:blur(6px)}}input,select,button.action{{height:37px;border:1px solid #cbd5dc;border-radius:6px;background:#fff;padding:0 11px;font:inherit;color:var(--ink)}}input{{flex:1;min-width:250px}}button.action{{cursor:pointer}}button.action:hover{{border-color:var(--blue);color:var(--blue)}}.hint{{margin-left:auto;color:var(--muted);font-size:12px}}
.tree-shell{{padding:20px 25px 28px}}ul{{list-style:none;margin:0;padding:0}}.node{{position:relative;border-bottom:1px solid #edf1f3}}.node ul{{margin-left:25px;padding-left:17px;border-left:1px solid #d9e2e7}}.row{{display:flex;gap:8px;align-items:flex-start;padding:10px 6px;border-radius:7px}}.row:hover{{background:#f7fafb}}.toggle,.leaf{{flex:0 0 24px;width:24px;height:24px;text-align:center}}.toggle{{border:1px solid #b9cbd3;border-radius:5px;color:var(--blue);background:var(--blue-soft);cursor:pointer;font-weight:bold}}.leaf{{color:#9eabb3;font-size:17px}}.main{{min-width:0;flex:1}}.title{{display:flex;gap:10px;align-items:baseline;flex-wrap:wrap}}.name{{font-weight:700;font-size:15px;cursor:pointer}}.level-1>.row .name{{font-size:18px;color:var(--blue)}}.level-2>.row .name{{color:#2f5567}}code{{color:#80909a;font-size:11px;overflow-wrap:anywhere}}.meta{{display:flex;flex-wrap:wrap;gap:5px;margin-top:5px}}.badge{{display:inline-flex;border-radius:999px;padding:2px 8px;font-size:11px;color:#53636d;background:#f0f3f5;border:1px solid #e0e6e9}}.badge.type{{color:#5b4b82;background:#f0ecfa;border-color:#ddd4ef}}.badge.high{{color:var(--green);background:var(--green-soft);border-color:#c8e8d5}}.badge.medium{{color:var(--gold);background:var(--gold-soft);border-color:#f1dfaa}}.badge.low{{color:var(--red);background:var(--red-soft);border-color:#efcaca}}.detail{{display:none;margin:0 6px 10px 32px;max-width:1100px;padding:11px 14px;background:#fbfcfd;border-left:3px solid #c8d8df;color:#46535c}}.detail.open{{display:block}}.detail p{{margin:0 0 8px;white-space:pre-wrap}}.detail p:last-child{{margin-bottom:0}}.collapsed>ul{{display:none}}.hidden{{display:none}}.match>.row{{background:var(--gold-soft)}}.empty{{display:none;color:var(--muted);text-align:center;padding:42px}}
@media(max-width:720px){{.app{{padding:14px 10px 30px}}header{{padding:19px 17px}}h1{{font-size:23px}}.tree-shell{{padding:14px 10px 20px}}.node ul{{margin-left:13px;padding-left:10px}}.hint{{width:100%;margin-left:0}}input{{min-width:180px}}}}
</style></head><body><main class="app"><header><h1>标准问题树 V0.2</h1><p class="subtitle">法律问题目录浏览页。节点默认收起；点击左侧按钮展开，点击问题名称查看定义和模型理由。</p><div class="stats" id="stats"></div></header>
<div class="toolbar"><input id="search" type="search" placeholder="搜索问题名称、问题 ID 或定义…"><select id="confidence"><option value="">全部置信度</option><option value="high">High</option><option value="medium">Medium</option><option value="low">Low</option></select><button class="action" id="expand">全部展开</button><button class="action" id="collapse">全部收起</button><button class="action" id="clear">清除筛选</button><span class="hint">draft，仅供人工审核，尚未接入生产。</span></div><section class="tree-shell"><ul id="tree">{tree_html}</ul><div class="empty" id="empty">没有匹配的节点</div></section></main>
<script>const stats={json.dumps(stats,ensure_ascii=False)};document.getElementById('stats').innerHTML=[['total','节点总数'],['directories','目录节点'],['issues','标准问题'],['rules','关联规则'],['high','High'],['medium','Medium'],['low','Low']].map(([k,l])=>`<div class="stat"><strong>${{stats[k]}}</strong><span>${{l}}</span></div>`).join('');const nodes=[...document.querySelectorAll('.node')];function expand(open){{nodes.forEach(n=>{{if(n.classList.contains('branch')){{n.classList.toggle('collapsed',!open);n.querySelector(':scope>.row .toggle').textContent=open?'−':'+'}}}})}}document.querySelectorAll('.toggle').forEach(b=>b.onclick=()=>{{const n=b.closest('.node');const c=n.classList.toggle('collapsed');b.textContent=c?'+':'−'}});document.querySelectorAll('.name').forEach(n=>n.onclick=()=>{{const d=n.closest('.node').querySelector(':scope>.detail');if(d)d.classList.toggle('open')}});document.getElementById('expand').onclick=()=>expand(true);document.getElementById('collapse').onclick=()=>expand(false);function filter(){{const q=document.getElementById('search').value.toLowerCase().trim(),c=document.getElementById('confidence').value;let count=0;nodes.forEach(n=>{{const ok=(!q||n.dataset.search.toLowerCase().includes(q))&&(!c||n.dataset.confidence===c);n.classList.toggle('hidden',!ok);n.classList.toggle('match',ok&&!!q);if(ok)count++}});document.getElementById('empty').style.display=count?'none':'block'}}document.getElementById('search').oninput=filter;document.getElementById('confidence').onchange=filter;document.getElementById('clear').onclick=()=>{{document.getElementById('search').value='';document.getElementById('confidence').value='';filter()}};expand(false);</script></body></html>'''
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(document, encoding="utf-8")
    return {"output": str(output_path), "stats": stats}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(export_html(args.workbook, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
