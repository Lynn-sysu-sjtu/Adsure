# -*- coding: utf-8 -*-
"""Draft-only source-first review pipeline for non-endorsement issue roots."""

from __future__ import annotations

import hashlib
import json
import os
import re
import unicodedata
from collections import defaultdict
from pathlib import Path


MODEL = "deepseek-chat"
PROMPT_VERSION = "all_primary_issue_fingerprint_v01_20260913"
CONFIDENCE = {"high", "medium", "low"}
EVIDENCE_REQUIREMENTS = {"content", "fact", "context", "mixed"}
LEGAL_EFFECTS = {
    "prohibited", "required", "disclosure_required", "warning_required",
    "evidence_required", "qualification_required", "registration_required",
    "consent_required", "truthfulness_required", "substantiation_required",
    "actual_use_required", "relationship_determination", "joint_liability",
    "administrative_liability", "civil_liability", "workflow_required",
    "scope_limitation", "other",
}
ROOT_NAMES = {
    "SCOPE_ACCESS": "适用主体与准入",
    "PROHIBITED_CONTENT": "禁止投放内容",
    "TRUTHFULNESS": "宣传真实性",
    "CLAIM_EXPRESSION": "宣称表达方式",
    "EFFICACY_PERFORMANCE": "功效与性能宣称",
    "PRICE_PROMOTION": "价格与促销",
    "EVIDENCE_FACT": "证明材料与事实核验",
    "DISCLOSURE_WARNING": "信息披露与警示语",
    "ENDORSEMENT_REVIEW": "代言、推荐与用户评价",
    "IP_PERSONALITY": "知识产权与人格权益",
    "MINORS_PUBLIC_ORDER": "未成年人及公序良俗",
    "MATERIAL_PLATFORM": "素材形式与平台规范",
    "WORKFLOW_DUTY": "投放流程与持续义务",
}
PROCESS_ROOTS = tuple(key for key in ROOT_NAMES if key != "ENDORSEMENT_REVIEW")


def atomic_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def records_for_root(records, mappings, root_id):
    expected = {
        item["rule_uid"] for item in mappings
        if item.get("primary_issue_id") == root_id
        or str(item.get("primary_issue_id") or "").startswith(root_id + ".")
    }
    by_uid = {item["rule"]["rule_uid"]: item for item in records}
    missing = expected - set(by_uid)
    if missing:
        raise ValueError(f"{root_id} has mappings without source records: {sorted(missing)[:5]}")
    return [by_uid[uid] for uid in sorted(expected)]


def source_details(record):
    rule = record.get("rule") or {}
    sources = {str(item.get("id")): item for item in record.get("legal_sources") or []}
    details = []
    for basis in rule.get("legal_basis") or []:
        source = sources.get(str(basis.get("source_id"))) or {}
        details.append({
            "source_id": basis.get("source_id"),
            "source_name": source.get("name") or basis.get("source_name") or basis.get("source_id") or "未知来源",
            "source_type": source.get("type") or rule.get("source_type") or "",
            "legal_level": int(basis.get("legal_level") or source.get("legal_level") or 99),
            "article": basis.get("article") or "未分条",
            "original_text": basis.get("text") or "",
        })
    return details


def compact_rule(record):
    rule = record.get("rule") or {}
    return {
        "rule_uid": rule.get("rule_uid"),
        "rule_id": rule.get("rule_id"),
        "track_hint_only": record.get("track"),
        "source_file": record.get("source_file"),
        "title": rule.get("title"),
        "source_type": rule.get("source_type"),
        "applies_to": rule.get("applies_to"),
        "rule_applicability": rule.get("rule_applicability"),
        "legal_basis": source_details(record),
    }


def build_skeleton_messages(root_id, root_name, l2_nodes, old_nodes):
    payload = {
        "task": "归并一个一级广告合规问题目录下的三级问题骨架",
        "root": {"issue_id": root_id, "name": root_name},
        "fixed_l2_nodes": [
            {"issue_id": item["issue_id"], "name": item["name"], "definition": item.get("definition") or ""}
            for item in l2_nodes
        ],
        "existing_lower_nodes": [
            {"issue_id": item["issue_id"], "parent_issue_id": item.get("parent_issue_id"), "name": item["name"], "definition": item.get("definition") or ""}
            for item in old_nodes if item.get("level", 0) >= 3
        ],
        "requirements": [
            "只归并现有问题，不补充外部法律知识",
            "同一主体、行为和法律效果的问题合并，行业与平台差异作为适用范围而非新问题",
            "每个节点必须直接挂在给定二级节点下",
            "issue_id 必须以一级和父二级 issue_id 为前缀并使用 ASCII 大写下划线",
            "输出 3 至 20 个三级节点；确有必要时可以少于 3 个",
        ],
        "output": {"nodes": [{
            "issue_id": f"{root_id}.L2.CANONICAL_ISSUE",
            "parent_issue_id": f"{root_id}.L2",
            "name": "标准问题名称", "definition": "边界清晰的问题定义",
            "confidence": "high/medium/low", "reason": "归并依据",
        }]},
    }
    system = (
        "你是广告合规法律问题目录整理员。只依据输入中的既有问题名称和定义做归并，"
        "不得引入外部法律规则。固定一级和二级节点不可修改。行业、品类、平台和渠道只是适用范围，"
        "除非法律效果或行为结构不同，否则不得据此拆成不同问题。只返回合法 JSON。"
    )
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(payload, ensure_ascii=False, separators=(",", ":"))}]


def validate_skeleton_payload(payload, root_id, allowed_l2_ids):
    nodes = payload.get("nodes") or []
    if not nodes:
        raise ValueError("skeleton must contain nodes")
    ids = [item.get("issue_id") for item in nodes]
    if any(not item for item in ids) or len(ids) != len(set(ids)):
        raise ValueError("skeleton issue_id values must be present and unique")
    for item in nodes:
        if item.get("parent_issue_id") not in allowed_l2_ids:
            raise ValueError("skeleton node has invalid L2 parent")
        if not str(item["issue_id"]).startswith(str(item["parent_issue_id"]) + "."):
            raise ValueError("skeleton issue_id must extend parent issue_id")
        if not re.fullmatch(r"[A-Z0-9_.]+", str(item["issue_id"])):
            raise ValueError("skeleton issue_id must be ASCII")
        if item.get("confidence") not in CONFIDENCE:
            raise ValueError("invalid skeleton confidence")
        if not str(item.get("name") or "").strip() or not str(item.get("definition") or "").strip():
            raise ValueError("skeleton name and definition are required")
    return nodes


def build_fingerprint_messages(root_id, root_name, issue_nodes, records):
    payload = {
        "task": "依据法条或平台规则原文生成规则指纹并归入标准问题",
        "current_root": {"issue_id": root_id, "name": root_name},
        "all_root_options": ROOT_NAMES,
        "allowed_issue_nodes": [
            {"issue_id": item["issue_id"], "name": item["name"], "definition": item["definition"]}
            for item in issue_nodes
        ],
        "rules": [compact_rule(item) for item in records],
        "enums": {
            "legal_effect": sorted(LEGAL_EFFECTS),
            "evidence_requirement": sorted(EVIDENCE_REQUIREMENTS),
            "confidence": sorted(CONFIDENCE),
        },
        "output": {"fingerprints": [{
            "rule_uid": "input rule_uid", "target_issue_id": "one allowed issue_id",
            "canonical_problem_name": "中文标准问题", "regulated_action": "ascii_snake_case",
            "legal_effect": "one enum", "actor_scope": [], "object_scope": [], "platform_scope": [],
            "evidence_requirement": "content/fact/context/mixed",
            "source_scope_conflict": False, "root_scope_conflict": False,
            "suggested_root_id": root_id, "confidence": "high/medium/low", "reason": "原文依据",
        }]},
    }
    system = (
        "你是广告合规法律规则分析员。法规名称、条款和原文是最高优先级依据；title 与 applies_to 次之；"
        "track_hint_only 只是文件存放位置的弱提示，绝不能覆盖原文。若原文明示的行业或品类与 track_hint_only 冲突，"
        "source_scope_conflict=true。若规则本质不属于 current_root，root_scope_conflict=true 并给出 suggested_root_id，"
        "但 target_issue_id 仍选择当前允许节点中最接近者。不得补充模型记忆中的法律内容。只返回合法 JSON。"
    )
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(payload, ensure_ascii=False, separators=(",", ":"))}]


def validate_fingerprint_payload(payload, expected_uids, allowed_issue_ids):
    items = payload.get("fingerprints") or []
    actual = [item.get("rule_uid") for item in items]
    if set(actual) != set(expected_uids) or len(actual) != len(set(actual)):
        raise ValueError("fingerprint batch must cover each expected rule_uid exactly once")
    by_uid = {}
    required = {
        "rule_uid", "target_issue_id", "canonical_problem_name", "regulated_action", "legal_effect",
        "actor_scope", "object_scope", "platform_scope", "evidence_requirement",
        "source_scope_conflict", "root_scope_conflict", "suggested_root_id", "confidence", "reason",
    }
    for item in items:
        missing = required - set(item)
        if missing:
            raise ValueError("missing fingerprint fields: " + ", ".join(sorted(missing)))
        if item["target_issue_id"] not in allowed_issue_ids:
            raise ValueError("fingerprint target_issue_id is not allowed")
        if item["legal_effect"] not in LEGAL_EFFECTS:
            raise ValueError("invalid legal_effect")
        if item["evidence_requirement"] not in EVIDENCE_REQUIREMENTS:
            raise ValueError("invalid evidence_requirement")
        if item["confidence"] not in CONFIDENCE:
            raise ValueError("invalid confidence")
        if item["suggested_root_id"] not in ROOT_NAMES:
            raise ValueError("invalid suggested_root_id")
        for field in ("actor_scope", "object_scope", "platform_scope"):
            if not isinstance(item[field], list):
                raise ValueError(f"{field} must be an array")
        for field in ("source_scope_conflict", "root_scope_conflict"):
            if not isinstance(item[field], bool):
                raise ValueError(f"{field} must be boolean")
        by_uid[item["rule_uid"]] = item
    return [by_uid[uid] for uid in expected_uids]


def response_payload(response):
    return json.loads(response["choices"][0]["message"]["content"])


def call_json_checkpoint(client, messages, checkpoint, prompt_version, validator=None):
    checkpoint = Path(checkpoint)
    digest = hashlib.sha256(json.dumps(messages, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    if checkpoint.exists():
        saved = json.loads(checkpoint.read_text(encoding="utf-8-sig"))
        if saved.get("digest") == digest and saved.get("model") == MODEL and saved.get("prompt_version") == prompt_version:
            payload = saved["payload"]
            return validator(payload) if validator else payload
    response = client.create_chat_completion(
        messages=messages, model=MODEL, temperature=0.0, response_format={"type": "json_object"},
    )
    payload = response_payload(response)
    result = validator(payload) if validator else payload
    atomic_json(checkpoint, {"digest": digest, "model": MODEL, "prompt_version": prompt_version, "payload": payload})
    return result


def normalize(value):
    value = unicodedata.normalize("NFKC", str(value or "")).lower()
    return re.sub(r"\s+", "", value)


def normalized_norm_text(text, article):
    value = normalize(text)
    article_value = normalize(article)
    if article_value and value.startswith(article_value):
        value = value[len(article_value):]
    return value.strip("：:。；;")


def build_exact_norm_groups(records_by_uid, fingerprints_by_uid):
    buckets = defaultdict(list)
    for uid, record in records_by_uid.items():
        details = source_details(record)
        signature = []
        for item in details:
            signature.append((
                normalize(item.get("source_name") or item.get("source_id")),
                normalize(item.get("article")),
                normalized_norm_text(item.get("original_text"), item.get("article")),
            ))
        raw = json.dumps({
            "sources": sorted(signature),
            "legal_effect": fingerprints_by_uid[uid].get("legal_effect"),
        }, ensure_ascii=False, sort_keys=True)
        key = "CRK-" + hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]
        buckets[key].append(uid)
    groups = []
    uid_to_key = {}
    for key, uids in sorted(buckets.items()):
        ordered = sorted(uids, key=lambda uid: (bool(fingerprints_by_uid[uid].get("source_scope_conflict")), uid))
        canonical_uid = ordered[0]
        groups.append({
            "canonical_rule_key": key,
            "canonical_rule_uid": canonical_uid,
            "source_occurrence_uids": sorted(uids),
            "duplicate_count": len(uids),
            "canonical_title": records_by_uid[canonical_uid]["rule"].get("title"),
            "review": {"status": "pending", "notes": None},
        })
        for uid in uids:
            uid_to_key[uid] = key
    return groups, uid_to_key


def is_proactive_candidate(fingerprint):
    return (
        fingerprint.get("evidence_requirement") in {"fact", "mixed"}
        and fingerprint.get("legal_effect") in {"evidence_required", "qualification_required"}
    )


def compile_root_assets(root_id, root_name, l2_nodes, skeleton, records, fingerprints):
    records_by_uid = {item["rule"]["rule_uid"]: item for item in records}
    fp_by_uid = {item["rule_uid"]: item for item in fingerprints}
    if set(records_by_uid) != set(fp_by_uid):
        raise ValueError("root compile must preserve complete rule_uid coverage")
    groups, uid_to_key = build_exact_norm_groups(records_by_uid, fp_by_uid)
    group_by_key = {item["canonical_rule_key"]: item for item in groups}
    nodes = []
    for item in l2_nodes:
        nodes.append({
            "issue_id": item["issue_id"], "parent_issue_id": root_id, "level": 2,
            "node_type": "directory", "name": item["name"], "definition": item.get("definition") or "",
            "rule_uids": [], "source_occurrence_uids": [],
        })
    for item in skeleton:
        nodes.append({
            **item, "level": 3, "node_type": "legal_issue",
            "rule_uids": [], "source_occurrence_uids": [],
        })
    node_by_id = {item["issue_id"]: item for item in nodes}
    mappings = []
    proactive = []
    for uid in sorted(records_by_uid):
        record = records_by_uid[uid]
        fp = fp_by_uid[uid]
        if fp["target_issue_id"] not in node_by_id:
            raise ValueError(f"unknown target issue: {fp['target_issue_id']}")
        proactive_flag = is_proactive_candidate(fp)
        group = group_by_key[uid_to_key[uid]]
        mapping = {
            "rule_uid": uid, "rule_id": record["rule"].get("rule_id"), "rule_title": record["rule"].get("title"),
            "root_id": root_id, "issue_id": fp["target_issue_id"],
            "mapping_type": "proactive_check" if proactive_flag else "legal_issue",
            "canonical_rule_key": uid_to_key[uid], "canonical_rule_uid": group["canonical_rule_uid"],
            "is_canonical_occurrence": uid == group["canonical_rule_uid"],
            "track": record.get("track"), "source_file": record.get("source_file"),
            "source_details": source_details(record), "fingerprint": fp,
            "review": {"status": "pending", "notes": None},
        }
        mappings.append(mapping)
        if proactive_flag:
            proactive.append({
                "rule_uid": uid, "issue_id": fp["target_issue_id"], "name": fp["canonical_problem_name"],
                "required_evidence": fp["evidence_requirement"], "source_details": mapping["source_details"],
                "review": {"status": "pending", "notes": None},
            })
        else:
            node = node_by_id[fp["target_issue_id"]]
            node["source_occurrence_uids"].append(uid)
            if uid == group["canonical_rule_uid"]:
                node["rule_uids"].append(uid)
    used_issue_ids = {item["issue_id"] for item in mappings}
    nodes = [
        item for item in nodes
        if item.get("level") != 3 or item["issue_id"] in used_issue_ids
    ]
    return {
        "schema_version": "0.1", "asset_status": "draft", "model": MODEL,
        "prompt_version": PROMPT_VERSION, "root_id": root_id, "root_name": root_name,
        "nodes": nodes, "mappings": mappings, "canonical_groups": groups,
        "proactive_candidates": proactive,
    }


def descendants(nodes, root_id):
    children = defaultdict(list)
    for item in nodes:
        children[item.get("parent_issue_id")].append(item)
    result = []
    stack = list(children[root_id])
    while stack:
        item = stack.pop()
        result.append(item)
        stack.extend(children[item["issue_id"]])
    return result


def run_root(client, root_id, taxonomy, mappings, records, assets_dir, report_dir, batch_size=4):
    root = next(item for item in taxonomy if item["issue_id"] == root_id)
    branch = descendants(taxonomy, root_id)
    l2_nodes = sorted((item for item in branch if item.get("level") == 2), key=lambda item: item["issue_id"])
    if not l2_nodes:
        raise ValueError(f"{root_id} has no L2 nodes")
    selected = records_for_root(records, mappings, root_id)
    checkpoint_dir = Path(report_dir) / "checkpoints" / root_id
    skeleton_messages = build_skeleton_messages(root_id, root["name"], l2_nodes, branch)
    skeleton = call_json_checkpoint(
        client, skeleton_messages, checkpoint_dir / "skeleton.json", PROMPT_VERSION + "_skeleton",
        lambda payload: validate_skeleton_payload(payload, root_id, {item["issue_id"] for item in l2_nodes}),
    )
    issue_ids = {item["issue_id"] for item in skeleton}
    fingerprints = []
    for offset in range(0, len(selected), batch_size):
        batch = selected[offset:offset + batch_size]
        uids = [item["rule"]["rule_uid"] for item in batch]
        digest = hashlib.sha256("|".join(uids).encode("utf-8")).hexdigest()[:16]
        messages = build_fingerprint_messages(root_id, root["name"], skeleton, batch)
        fingerprints.extend(call_json_checkpoint(
            client, messages, checkpoint_dir / "fingerprints" / f"{digest}.json", PROMPT_VERSION + "_fingerprint",
            lambda payload, expected=uids: validate_fingerprint_payload(payload, expected, issue_ids),
        ))
    assets = compile_root_assets(root_id, root["name"], l2_nodes, skeleton, selected, fingerprints)
    output = Path(assets_dir) / root_id
    atomic_json(output / "review_assets.json", assets)
    return assets


def export_root_workbook(path, assets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    workbook = Workbook()
    tree = workbook.active
    tree.title = "问题树"
    tree.append(["层级", "节点类型", "问题ID", "父节点ID", "问题名称", "定义", "标准规范数", "来源记录数", "人工结论", "备注"])
    for item in assets["nodes"]:
        tree.append([item["level"], item["node_type"], item["issue_id"], item.get("parent_issue_id"), item["name"], item.get("definition"), len(item.get("rule_uids") or []), len(item.get("source_occurrence_uids") or []), "", ""])
    mapping = workbook.create_sheet("规则映射与原文")
    mapping.append(["rule_uid", "问题ID", "映射类型", "规则标题", "原始赛道", "来源文件", "来源类型", "法规/平台规则", "条款", "规则原文", "适用对象", "法律效果", "建议一级问题", "赛道污染", "跨一级问题", "置信度", "人工结论", "调整后问题", "确认适用范围", "备注"])
    for item in assets["mappings"]:
        fp = item["fingerprint"]
        details = item.get("source_details") or [{}]
        mapping.append([
            item["rule_uid"], item["issue_id"], item["mapping_type"], item["rule_title"], item.get("track"), item.get("source_file"),
            "\n".join(str(x.get("source_type") or "") for x in details), "\n".join(str(x.get("source_name") or "") for x in details),
            "\n".join(str(x.get("article") or "") for x in details), "\n\n".join(str(x.get("original_text") or "") for x in details),
            ", ".join(fp.get("object_scope") or []), fp.get("legal_effect"), fp.get("suggested_root_id"),
            fp.get("source_scope_conflict"), fp.get("root_scope_conflict"), fp.get("confidence"), "", "", "", "",
        ])
    duplicates = workbook.create_sheet("完全重复规范")
    duplicates.append(["canonical_rule_key", "标准rule_uid", "来源数量", "来源rule_uid", "规则标题", "人工结论", "备注"])
    for item in assets["canonical_groups"]:
        if item["duplicate_count"] > 1:
            duplicates.append([item["canonical_rule_key"], item["canonical_rule_uid"], item["duplicate_count"], "\n".join(item["source_occurrence_uids"]), item["canonical_title"], "", ""])
    focus = workbook.create_sheet("待人工重点审核")
    focus.append(["rule_uid", "审核原因", "问题ID", "规则标题", "原始赛道", "规则原文", "模型适用对象", "建议一级问题", "标准rule_uid", "其他重复来源", "适用范围结论", "去重结论", "链路结论", "调整后问题", "确认适用范围", "备注"])
    group_by_key = {item["canonical_rule_key"]: item for item in assets["canonical_groups"]}
    for item in assets["mappings"]:
        fp = item["fingerprint"]
        group = group_by_key[item["canonical_rule_key"]]
        reasons = []
        if fp.get("source_scope_conflict"): reasons.append("赛道污染")
        if fp.get("root_scope_conflict"): reasons.append("跨一级问题")
        if group["duplicate_count"] > 1: reasons.append("完全重复规范")
        if item["mapping_type"] == "proactive_check": reasons.append("主动补资料候选")
        if fp.get("confidence") != "high": reasons.append("中低置信度")
        if not reasons: continue
        focus.append([
            item["rule_uid"], "；".join(reasons), item["issue_id"], item["rule_title"], item.get("track"),
            "\n\n".join(str(x.get("original_text") or "") for x in item.get("source_details") or []),
            ", ".join(fp.get("object_scope") or []), fp.get("suggested_root_id"), group["canonical_rule_uid"],
            "\n".join(uid for uid in group["source_occurrence_uids"] if uid != item["rule_uid"]), "", "", "", "", "", "",
        ])
    validations = [
        (focus, "K", ["适用范围准确", "确认赛道污染", "模型范围错误", "无法确定"]),
        (focus, "L", ["合并为同一标准规范", "独立保留", "不涉及去重", "无法确定"]),
        (focus, "M", ["保留在法律问题树", "保留在主动补资料", "移至workflow", "无法确定"]),
    ]
    for sheet, column, choices in validations:
        validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        sheet.add_data_validation(validation)
        if sheet.max_row >= 2:
            validation.add(f"{column}2:{column}{sheet.max_row}")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="294C60")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _global_rule_signature(mapping):
    signature = []
    for item in mapping.get("source_details") or []:
        signature.append((
            normalize(item.get("source_name") or item.get("source_id")),
            normalize(item.get("article")),
            normalized_norm_text(item.get("original_text"), item.get("article")),
        ))
    raw = json.dumps({
        "sources": sorted(signature),
        "legal_effect": (mapping.get("fingerprint") or {}).get("legal_effect"),
    }, ensure_ascii=False, sort_keys=True)
    return "GCRK-" + hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]


def compile_global_audit(mappings, expected_mappings, records_by_uid):
    """Compile one deterministic review view without changing draft assets."""
    occurrences_by_uid = defaultdict(list)
    for item in mappings:
        uid = item["rule_uid"]
        root_id = item.get("root_id") or "ENDORSEMENT_REVIEW"
        if any((saved.get("root_id") or "ENDORSEMENT_REVIEW") == root_id for saved in occurrences_by_uid[uid]):
            raise ValueError(f"duplicate processed rule_uid inside one root: {uid} / {root_id}")
        occurrences_by_uid[uid].append(item)
    expected_by_uid = {}
    for item in expected_mappings:
        uid = item["rule_uid"]
        if uid in expected_by_uid:
            raise ValueError(f"duplicate expected rule_uid: {uid}")
        expected_by_uid[uid] = item

    by_uid = {}
    cross_root_overlaps = []
    for uid, occurrences in sorted(occurrences_by_uid.items()):
        expected = expected_by_uid.get(uid) or {}
        expected_root = str(expected.get("primary_issue_id") or "").split(".", 1)[0]
        preferred = [
            item for item in occurrences
            if (item.get("root_id") or "ENDORSEMENT_REVIEW") == expected_root
        ]
        selected = preferred[0] if preferred else sorted(
            occurrences, key=lambda item: (item.get("root_id") or "ENDORSEMENT_REVIEW", item.get("issue_id") or "")
        )[0]
        by_uid[uid] = selected
        root_ids = sorted({item.get("root_id") or "ENDORSEMENT_REVIEW" for item in occurrences})
        if len(root_ids) > 1:
            cross_root_overlaps.append({
                "rule_uid": uid,
                "expected_root_id": expected_root,
                "selected_root_id": selected.get("root_id") or "ENDORSEMENT_REVIEW",
                "root_ids": root_ids,
                "occurrences": sorted(occurrences, key=lambda item: item.get("root_id") or "ENDORSEMENT_REVIEW"),
            })

    buckets = defaultdict(list)
    selected_mappings = list(by_uid.values())
    for item in selected_mappings:
        buckets[_global_rule_signature(item)].append(item)
    duplicate_groups = []
    for key, items in sorted(buckets.items()):
        if len(items) < 2:
            continue
        ordered = sorted(items, key=lambda item: (
            bool((item.get("fingerprint") or {}).get("source_scope_conflict")),
            item["rule_uid"],
        ))
        duplicate_groups.append({
            "global_canonical_rule_key": key,
            "canonical_rule_uid": ordered[0]["rule_uid"],
            "rule_uids": [item["rule_uid"] for item in ordered],
            "root_ids": sorted({item.get("root_id") or "ENDORSEMENT_REVIEW" for item in ordered}),
            "tracks": sorted({str(item.get("track") or "") for item in ordered}),
            "occurrences": ordered,
        })

    missing_uids = sorted(set(expected_by_uid) - set(by_uid))
    missing = []
    for uid in missing_uids:
        expected = expected_by_uid[uid]
        record = records_by_uid.get(uid) or {}
        rule = record.get("rule") or {}
        missing.append({
            "rule_uid": uid,
            "rule_id": rule.get("rule_id"),
            "rule_title": rule.get("title"),
            "track": record.get("track"),
            "source_file": record.get("source_file"),
            "expected_primary_issue_id": expected.get("primary_issue_id"),
            "source_details": source_details(record) if record else [],
        })

    def flagged(field):
        return sorted(
            (item for item in selected_mappings if (item.get("fingerprint") or {}).get(field)),
            key=lambda item: item["rule_uid"],
        )

    return {
        "schema_version": "0.1",
        "asset_status": "draft",
        "coverage": {
            "expected_count": len(expected_by_uid),
            "processed_count": len(by_uid),
            "missing_count": len(missing_uids),
            "missing_rule_uids": missing_uids,
            "unexpected_rule_uids": sorted(set(by_uid) - set(expected_by_uid)),
        },
        "mappings": sorted(selected_mappings, key=lambda item: item["rule_uid"]),
        "cross_root_overlaps": cross_root_overlaps,
        "duplicate_groups": duplicate_groups,
        "scope_conflicts": flagged("source_scope_conflict"),
        "root_conflicts": flagged("root_scope_conflict"),
        "proactive_candidates": sorted(
            (item for item in selected_mappings if item.get("mapping_type") == "proactive_check"),
            key=lambda item: item["rule_uid"],
        ),
        "low_or_medium": sorted(
            (item for item in selected_mappings if (item.get("fingerprint") or {}).get("confidence") != "high"),
            key=lambda item: item["rule_uid"],
        ),
        "missing": missing,
    }


def export_global_audit_workbook(path, audit, root_summaries, source_snapshot):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    summary = workbook.active
    summary.title = "一级目录处理汇总"
    summary.append(["一级目录", "目录名称", "处理规则数", "三级问题数", "目录内重复组", "赛道污染", "跨一级问题", "主动补资料", "中低置信度", "状态"])
    for item in root_summaries:
        summary.append([
            item["root_id"], item["root_name"], item["rule_count"], item.get("issue_count"),
            item.get("duplicate_group_count"), item.get("scope_conflict_count"),
            item.get("root_conflict_count"), item.get("proactive_candidate_count"),
            item.get("low_or_medium_count", 0), "draft / 待人工审核",
        ])
    coverage = audit["coverage"]
    summary.append(["TOTAL", "全局", coverage["processed_count"], "", len(audit["duplicate_groups"]), len(audit["scope_conflicts"]), len(audit["root_conflicts"]), len(audit["proactive_candidates"]), len(audit["low_or_medium"]), f"应处理 {coverage['expected_count']}，缺口 {coverage['missing_count']}"])

    def texts(item, field):
        return "\n\n".join(str(detail.get(field) or "") for detail in item.get("source_details") or [])

    index = workbook.create_sheet("全部规则索引")
    index.append(["rule_uid", "一级目录", "问题ID", "映射类型", "规则标题", "原始赛道", "来源文件", "法规/平台规则", "条款", "规则原文", "模型适用对象", "法律效果", "置信度", "赛道污染", "跨一级问题", "建议一级目录"])
    for item in audit["mappings"]:
        fp = item.get("fingerprint") or {}
        index.append([item["rule_uid"], item.get("root_id") or "ENDORSEMENT_REVIEW", item.get("issue_id"), item.get("mapping_type"), item.get("rule_title"), item.get("track"), item.get("source_file"), texts(item, "source_name"), texts(item, "article"), texts(item, "original_text"), ", ".join(fp.get("object_scope") or []), fp.get("legal_effect"), fp.get("confidence"), fp.get("source_scope_conflict", False), fp.get("root_scope_conflict", False), fp.get("suggested_root_id") or fp.get("recommended_parent")])

    duplicates = workbook.create_sheet("完全重复规范")
    duplicates.append(["全局规范键", "标准rule_uid", "来源数", "涉及一级目录", "涉及赛道", "全部rule_uid", "全部规则标题", "法规/平台规则", "条款", "规则原文", "人工结论", "确认标准rule_uid", "备注"])
    for group in audit["duplicate_groups"]:
        occurrences = group["occurrences"]
        duplicates.append([group["global_canonical_rule_key"], group["canonical_rule_uid"], len(occurrences), "\n".join(group["root_ids"]), "\n".join(group["tracks"]), "\n".join(item["rule_uid"] for item in occurrences), "\n".join(str(item.get("rule_title") or "") for item in occurrences), texts(occurrences[0], "source_name"), texts(occurrences[0], "article"), texts(occurrences[0], "original_text"), "", "", ""])

    overlaps = workbook.create_sheet("同一规则跨目录重叠")
    overlaps.append(["rule_uid", "原始主目录", "当前选定主目录", "涉及目录", "各目录问题ID", "规则标题", "法规/平台规则", "条款", "规则原文", "人工结论", "确认主目录", "备注"])
    for group in audit["cross_root_overlaps"]:
        occurrences = group["occurrences"]
        overlaps.append([group["rule_uid"], group["expected_root_id"], group["selected_root_id"], "\n".join(group["root_ids"]), "\n".join(f"{item.get('root_id') or 'ENDORSEMENT_REVIEW'}: {item.get('issue_id') or ''}" for item in occurrences), occurrences[0].get("rule_title"), texts(occurrences[0], "source_name"), texts(occurrences[0], "article"), texts(occurrences[0], "original_text"), "", "", ""])

    def add_review_sheet(title, items, reason_label):
        sheet = workbook.create_sheet(title)
        sheet.append(["rule_uid", "审核原因", "一级目录", "问题ID", "规则标题", "原始赛道", "来源文件", "法规/平台规则", "条款", "规则原文", "模型适用对象", "模型理由", "建议一级目录", "人工结论", "调整后归属", "备注"])
        for item in items:
            fp = item.get("fingerprint") or {}
            sheet.append([item["rule_uid"], reason_label, item.get("root_id") or "ENDORSEMENT_REVIEW", item.get("issue_id"), item.get("rule_title"), item.get("track"), item.get("source_file"), texts(item, "source_name"), texts(item, "article"), texts(item, "original_text"), ", ".join(fp.get("object_scope") or []), fp.get("reason"), fp.get("suggested_root_id") or fp.get("recommended_parent"), "", "", ""])
        return sheet

    scope = add_review_sheet("赛道污染", audit["scope_conflicts"], "原始赛道与规则原文适用范围可能冲突")
    root = add_review_sheet("跨一级问题候选", audit["root_conflicts"], "规则可能应归入其他一级目录")
    proactive = add_review_sheet("主动补资料候选", audit["proactive_candidates"], "事实或资质核验候选")
    confidence = add_review_sheet("中低置信度", audit["low_or_medium"], "模型映射置信度非 high")

    missing = workbook.create_sheet("未处理覆盖缺口")
    missing.append(["rule_uid", "原规则ID", "规则标题", "原始赛道", "原一级问题", "来源文件", "法规/平台规则", "条款", "规则原文", "缺口原因", "人工结论", "调整后归属", "备注"])
    for item in audit["missing"]:
        missing.append([item["rule_uid"], item.get("rule_id"), item.get("rule_title"), item.get("track"), item.get("expected_primary_issue_id"), item.get("source_file"), texts(item, "source_name"), texts(item, "article"), texts(item, "original_text"), "未进入其余12目录或代言v0.3试点", "", "", ""])

    validation_sheet = workbook.create_sheet("校验结果")
    validation_sheet.append(["校验项", "结果", "说明"])
    checks = [
        ("原始映射规则数", coverage["expected_count"], "rule_issue_mapping_draft_v0.3.json"),
        ("本轮已处理规则数", coverage["processed_count"], "其余12个一级目录 + 代言专题v0.3"),
        ("覆盖缺口", coverage["missing_count"], "详见未处理覆盖缺口"),
        ("意外新增rule_uid", len(coverage["unexpected_rule_uids"]), "应为0"),
        ("rule_uid唯一", len({item["rule_uid"] for item in audit["mappings"]}) == len(audit["mappings"]), "应为TRUE"),
        ("jsonbase规则数", source_snapshot.get("rule_count"), "规则库快照"),
        ("jsonbase SHA-256", source_snapshot.get("jsonbase_sha256"), "应与生成前快照一致"),
    ]
    for row in checks:
        validation_sheet.append(list(row))

    duplicate_validation = DataValidation(type="list", formula1='"合并为同一标准规范,独立保留,无法确定"', allow_blank=True)
    duplicates.add_data_validation(duplicate_validation)
    if duplicates.max_row >= 2:
        duplicate_validation.add(f"K2:K{duplicates.max_row}")
    for sheet in (scope, root, proactive, confidence):
        general_validation = DataValidation(type="list", formula1='"确认模型判断,调整归属,独立保留,无法确定"', allow_blank=True)
        sheet.add_data_validation(general_validation)
        if sheet.max_row >= 2:
            general_validation.add(f"N2:N{sheet.max_row}")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="294C60")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
