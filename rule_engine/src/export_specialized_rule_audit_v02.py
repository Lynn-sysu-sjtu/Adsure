# -*- coding: utf-8 -*-
"""Export the specialized rule audit workbook against approved v0.2 assets."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports" / "approved_issue_tree_v02"
DEFAULT_OUTPUT = REPORT_DIR / "全部规则专项人工审核表_v0.2.xlsx"

COMMON_HEADERS = [
    "当前一级问题", "当前二级问题", "当前问题", "规则UID", "规则标题",
    "映射角色", "原始赛道", "来源文件", "来源类型", "法规或平台规则",
    "条款", "规则原文", "人工结论", "重复目标rule_uid", "建议赛道",
    "调整后问题", "资产处置", "备注",
]


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def node_context(nodes):
    by_id = {node["issue_id"]: node for node in nodes}

    def ancestor(issue_id, level):
        current = by_id.get(issue_id)
        while current and current.get("level") != level:
            current = by_id.get(current.get("parent_issue_id"))
        return current or {}

    return by_id, ancestor


def base_row(mapping, nodes, ancestor):
    issue_id = mapping["issue_id"]
    return {
        "当前一级问题": ancestor(issue_id, 1).get("name", ""),
        "当前二级问题": ancestor(issue_id, 2).get("name", ""),
        "当前问题": nodes.get(issue_id, {}).get("name", issue_id),
        "规则UID": mapping["rule_uid"],
        "规则标题": mapping.get("rule_title", ""),
        "映射角色": mapping.get("mapping_type", ""),
        "原始赛道": mapping.get("track", ""),
        "来源文件": mapping.get("source_file", ""),
        "来源类型": mapping.get("source_type", ""),
        "法规或平台规则": mapping.get("source_name", ""),
        "条款": mapping.get("article", ""),
        "规则原文": mapping.get("original_text", ""),
        "人工结论": "",
        "重复目标rule_uid": "",
        "建议赛道": "",
        "调整后问题": "",
        "资产处置": "",
        "备注": "",
    }


def build_audit_rows(project_root=ROOT):
    project_root = Path(project_root)
    current_dir = project_root / "reports" / "approved_issue_tree_v02"
    taxonomy = read_json(current_dir / "approved_issue_taxonomy_v0.2.json")
    mapping_asset = read_json(current_dir / "approved_rule_issue_mapping_v0.2.json")
    historical = read_json(
        project_root / "reports" / "all_primary_issue_fingerprint_review_v01" / "global_audit_v0.1.json"
    )
    nodes, ancestor = node_context(taxonomy["nodes"])
    active = mapping_asset["mappings"]
    excluded = mapping_asset["excluded_mappings"]
    all_current = active + excluded
    current_by_uid = defaultdict(list)
    for mapping in all_current:
        current_by_uid[mapping["rule_uid"]].append(mapping)

    historical_mapping = {}
    for mapping in historical["mappings"]:
        historical_mapping.setdefault(mapping["rule_uid"], mapping)

    all_rows = [base_row(mapping, nodes, ancestor) for mapping in active]
    exact_duplicates = []
    duplicate_noncanonical_uids = set()
    for group in historical["duplicate_groups"]:
        present = [uid for uid in group["rule_uids"] if uid in current_by_uid]
        if len(present) < 2:
            continue
        canonical = group["canonical_rule_uid"] if group["canonical_rule_uid"] in present else present[0]
        for uid in present:
            mapping = current_by_uid[uid][0]
            row = base_row(mapping, nodes, ancestor)
            row.update({
                "重复组": group["global_canonical_rule_key"],
                "组内rule_uid": "、".join(present),
                "重复目标rule_uid": canonical,
                "候选说明": "来源、条款、原文及法律效果标准化后一致",
            })
            exact_duplicates.append(row)
            if uid != canonical:
                duplicate_noncanonical_uids.add(uid)

    cross_directory = []
    for uid, occurrences in current_by_uid.items():
        active_occurrences = [item for item in occurrences if "exclusion_role" not in item]
        issue_ids = sorted({item["issue_id"] for item in active_occurrences})
        if len(issue_ids) < 2:
            continue
        for mapping in active_occurrences:
            row = base_row(mapping, nodes, ancestor)
            row.update({
                "同一UID当前问题数": len(issue_ids),
                "同一UID全部当前问题": " | ".join(nodes.get(item, {}).get("name", item) for item in issue_ids),
                "候选说明": "同一 rule_uid 当前映射到多个问题，需判断是否属于辅助依据、例外或真实跨问题适用",
            })
            cross_directory.append(row)

    track_conflicts = []
    conflict_uids = {item["rule_uid"] for item in historical["scope_conflicts"]}
    for uid in sorted(conflict_uids & set(current_by_uid)):
        signal = historical_mapping.get(uid, {}).get("fingerprint", {})
        for mapping in current_by_uid[uid]:
            row = base_row(mapping, nodes, ancestor)
            row.update({
                "模型识别适用对象": "、".join(signal.get("object_scope") or []),
                "模型识别平台": "、".join(signal.get("platform_scope") or []),
                "历史判断理由": signal.get("reason", ""),
                "候选说明": "历史模型判断规则原文适用对象与所在赛道不一致；应以规则原文为准",
            })
            track_conflicts.append(row)

    mismatch = []
    historical_root_conflicts = {item["rule_uid"] for item in historical["root_conflicts"]}
    for mapping in active:
        changed = mapping.get("original_issue_id") and mapping["original_issue_id"] != mapping["issue_id"]
        if not changed and mapping["rule_uid"] not in historical_root_conflicts:
            continue
        row = base_row(mapping, nodes, ancestor)
        row.update({
            "原问题ID": mapping.get("original_issue_id", ""),
            "当前问题ID": mapping["issue_id"],
            "候选类型": "v0.2 已调整映射复核" if changed else "历史跨一级问题候选",
            "候选说明": "核对规则原文是否完整支持当前问题；若不支持，填写调整后问题",
        })
        mismatch.append(row)

    excluded_rows = []
    for mapping in excluded:
        row = base_row(mapping, nodes, ancestor)
        row.update({
            "原问题ID": mapping.get("original_issue_id", ""),
            "保留角色": mapping.get("exclusion_role", ""),
            "候选说明": mapping.get("exclusion_reason", ""),
        })
        excluded_rows.append(row)

    removal_candidates = []
    for mapping in excluded:
        row = base_row(mapping, nodes, ancestor)
        row.update({
            "候选来源": "已移出普通文案树",
            "候选说明": "判断应保留为旁路资产，还是与广告审核无关并移出规则库",
            "保留角色": mapping.get("exclusion_role", ""),
        })
        removal_candidates.append(row)
    for uid in sorted(duplicate_noncanonical_uids):
        mapping = current_by_uid[uid][0]
        row = base_row(mapping, nodes, ancestor)
        canonical = next(
            group["canonical_rule_uid"] for group in historical["duplicate_groups"]
            if uid in group["rule_uids"]
        )
        row.update({
            "候选来源": "完全重复规范的非主记录",
            "重复目标rule_uid": canonical,
            "候选说明": "优先考虑合并到主 rule_uid；不要仅凭重复标签直接删除来源信息",
        })
        removal_candidates.append(row)

    proactive_uids = {item["rule_uid"] for item in historical["proactive_candidates"]}
    proactive = []
    for mapping in all_current:
        if mapping.get("mapping_type") != "proactive_check" and mapping["rule_uid"] not in proactive_uids:
            continue
        row = base_row(mapping, nodes, ancestor)
        row["候选说明"] = "核对其是否必须依赖第三方资料、资质或事实证明；不得作为已违规结论"
        proactive.append(row)

    confidence_rows = []
    confidence_by_uid = {
        item["rule_uid"]: (item.get("fingerprint") or {}).get("confidence", "")
        for item in historical["low_or_medium"]
    }
    for uid, confidence in confidence_by_uid.items():
        for mapping in current_by_uid.get(uid, []):
            row = base_row(mapping, nodes, ancestor)
            signal = historical_mapping.get(uid, {}).get("fingerprint", {})
            row.update({
                "模型置信度": confidence,
                "历史判断理由": signal.get("reason", ""),
                "候选说明": "优先核对模型对适用对象、行为结构和问题归属的理解",
            })
            confidence_rows.append(row)

    return {
        "all_rules": all_rows,
        "exact_duplicates": exact_duplicates,
        "cross_directory": cross_directory,
        "track_conflicts": track_conflicts,
        "mismatch": mismatch,
        "excluded_candidates": excluded_rows,
        "removal_candidates": removal_candidates,
        "proactive": proactive,
        "low_confidence": confidence_rows,
    }


def ordered_headers(rows):
    extras = []
    for row in rows:
        for key in row:
            if key not in COMMON_HEADERS and key not in extras:
                extras.append(key)
    insert_at = COMMON_HEADERS.index("人工结论")
    return COMMON_HEADERS[:insert_at] + extras + COMMON_HEADERS[insert_at:]


def add_review_validation(sheet, headers):
    if sheet.max_row < 2 or "人工结论" not in headers:
        return
    conclusion = DataValidation(
        type="list",
        formula1='"保留,重复-保留当前,重复-合并到目标UID,赛道污染-调整赛道,问题错配-迁移,移出普通文案树,保留为旁路资产,移出规则库,待讨论"',
        allow_blank=True,
    )
    disposal = DataValidation(
        type="list",
        formula1='"保留原记录,合并去重,调整赛道,迁移问题,转fact_check,转workflow_reference,转platform_access,转supporting_basis,转exception,移出规则库,待讨论"',
        allow_blank=True,
    )
    sheet.add_data_validation(conclusion)
    conclusion.add(f"{sheet.cell(1, headers.index('人工结论') + 1).column_letter}2:{sheet.cell(1, headers.index('人工结论') + 1).column_letter}{sheet.max_row}")
    if "资产处置" in headers:
        sheet.add_data_validation(disposal)
        letter = sheet.cell(1, headers.index("资产处置") + 1).column_letter
        disposal.add(f"{letter}2:{letter}{sheet.max_row}")


def write_sheet(workbook, title, rows):
    sheet = workbook.create_sheet(title)
    headers = ordered_headers(rows)
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    header_fill = PatternFill("solid", fgColor="1F5364")
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="D9E2E5")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    review_names = {"人工结论", "重复目标rule_uid", "建议赛道", "调整后问题", "资产处置", "备注"}
    for row in sheet.iter_rows(min_row=2):
        for index, cell in enumerate(row):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if headers[index] in review_names:
                cell.fill = review_fill
        if "候选说明" in headers:
            row[headers.index("候选说明")].fill = warning_fill
    for index, header in enumerate(headers, 1):
        width = 18
        if header in {"规则原文", "历史判断理由"}:
            width = 72
        elif header in {"来源文件", "原问题ID", "当前问题ID", "同一UID全部当前问题", "候选说明"}:
            width = 43
        elif header in {"规则标题", "当前问题", "调整后问题"}:
            width = 34
        elif header in {"当前一级问题", "当前二级问题", "法规或平台规则"}:
            width = 27
        elif header in {"规则UID", "重复目标rule_uid"}:
            width = 23
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    add_review_validation(sheet, headers)
    return sheet


def export_workbook(project_root=ROOT, output_path=DEFAULT_OUTPUT):
    audit = build_audit_rows(project_root)
    workbook = Workbook()
    guide = workbook.active
    guide.title = "审核说明"
    guide.append(["审核对象", "本 Sheet 要解决的问题", "人工审核重点"])
    guidance = [
        ("全部规则索引", "查看 v0.2 当前有效映射全貌", "一条规则可以因辅助依据或例外角色出现多次；不要仅按行数判断重复。"),
        ("完全重复规范", "来源、条款、原文及法律效果标准化后相同", "确认是否同一规范的多赛道副本；指定保留 UID，其他记录选择合并去重。"),
        ("同一规则跨目录重叠", "同一 rule_uid 挂到多个当前问题", "判断是真实跨问题适用，还是问题错配；辅助依据与例外允许跨问题。"),
        ("赛道污染", "历史模型认为原文适用对象与存储赛道不一致", "以原文适用范围为准；填写建议赛道，不因目录位置机械判断。"),
        ("问题错配复核", "v0.2 已迁移规则及历史跨一级问题候选", "逐条比对问题名称与规则原文是否一一对应。"),
        ("移出普通文案树", "不适合作为普通文案直接违规规则的 31 条资产", "判断应保留为 fact/workflow/access/supporting 旁路，还是移出规则库。"),
        ("移出规则库候选", "重复副本与已移出普通文案树规则的集中处置页", "移出规则库是最后手段；先判断是否应合并或保留为旁路资产。"),
        ("主动补资料候选", "依赖第三方资料或事实核验的规则", "不得写成已经违规；确认触发条件和所需资料。"),
        ("中低置信度", "历史模型理解不稳定的规则", "优先核对适用对象、受规制行为、法律效果和问题归属。"),
    ]
    for row in guidance:
        guide.append(row)
    guide.column_dimensions["A"].width = 25
    guide.column_dimensions["B"].width = 55
    guide.column_dimensions["C"].width = 85
    for cell in guide[1]:
        cell.fill = PatternFill("solid", fgColor="1F5364")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    write_sheet(workbook, "全部规则索引", audit["all_rules"])
    write_sheet(workbook, "完全重复规范", audit["exact_duplicates"])
    write_sheet(workbook, "同一规则跨目录重叠", audit["cross_directory"])
    write_sheet(workbook, "赛道污染", audit["track_conflicts"])
    write_sheet(workbook, "问题错配复核", audit["mismatch"])
    write_sheet(workbook, "移出普通文案树", audit["excluded_candidates"])
    write_sheet(workbook, "移出规则库候选", audit["removal_candidates"])
    write_sheet(workbook, "主动补资料候选", audit["proactive"])
    write_sheet(workbook, "中低置信度", audit["low_confidence"])
    checks = workbook.create_sheet("校验结果")
    checks.append(["校验项", "结果", "数量"])
    checks.append(["v0.2 活动映射均有规则原文", "通过" if all(row["规则原文"] for row in audit["all_rules"]) else "失败", len(audit["all_rules"])])
    checks.append(["移出普通文案树规则已进入专项 Sheet", "通过", len(audit["excluded_candidates"])])
    checks.append(["专项 Sheet 数据源", "approved_issue_taxonomy_v0.2 + approved_rule_issue_mapping_v0.2 + v0.1 历史候选信号", 3])
    checks.append(["生产状态", "仅供人工审核，未接入生产召回链路", 0])
    for cell in checks[1]:
        cell.fill = PatternFill("solid", fgColor="1F5364")
        cell.font = Font(color="FFFFFF", bold=True)
    checks.column_dimensions["A"].width = 40
    checks.column_dimensions["B"].width = 95
    checks.column_dimensions["C"].width = 15
    for row in checks.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        "output": str(output_path),
        "sheet_counts": {
            "全部规则索引": len(audit["all_rules"]),
            "完全重复规范": len(audit["exact_duplicates"]),
            "同一规则跨目录重叠": len(audit["cross_directory"]),
            "赛道污染": len(audit["track_conflicts"]),
            "问题错配复核": len(audit["mismatch"]),
            "移出普通文案树": len(audit["excluded_candidates"]),
            "移出规则库候选": len(audit["removal_candidates"]),
            "主动补资料候选": len(audit["proactive"]),
            "中低置信度": len(audit["low_confidence"]),
        },
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    print(json.dumps(export_workbook(args.project_root, args.output), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
