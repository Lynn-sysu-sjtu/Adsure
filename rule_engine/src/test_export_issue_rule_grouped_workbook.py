# -*- coding: utf-8 -*-

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from export_issue_rule_grouped_workbook import export_workbook, load_mapping_rows


ROOT = Path(__file__).resolve().parents[1]


class GroupedIssueRuleWorkbookTests(unittest.TestCase):
    def assert_groups_are_contiguous(self, rows, fields):
        seen = set()
        current = None
        for row in rows:
            key = tuple(row[field] for field in fields)
            if key == current:
                continue
            self.assertNotIn(key, seen, f"group is split across the workbook: {key}")
            seen.add(key)
            current = key

    def test_rows_are_grouped_by_issue_and_keep_one_rule_per_row(self):
        rows = load_mapping_rows(ROOT)
        self.assertGreater(len(rows), 800)
        self.assertEqual(len({row["规则UID"] for row in rows}), 842)
        self.assertTrue(any(row["规则映射类型"] == "proactive_check" for row in rows))
        self.assertEqual(rows[0]["二级问题ID"], "SCOPE_ACCESS.AD_REVIEW_ACCESS")
        self.assertEqual(
            rows[0]["三级问题ID"],
            "SCOPE_ACCESS.AD_REVIEW_ACCESS.PRE_REVIEW_APPROVAL_AND_CONTENT_CONSISTENCY",
        )
        self.assert_groups_are_contiguous(rows, ("一级问题ID",))
        self.assert_groups_are_contiguous(rows, ("一级问题ID", "二级问题ID"))
        self.assert_groups_are_contiguous(rows, ("一级问题ID", "二级问题ID", "三级问题ID"))
        for first, second in zip(rows, rows[1:]):
            if first["三级问题"] == second["三级问题"] and first["三级问题"]:
                self.assertNotEqual(first["规则UID"], second["规则UID"])

    def test_workbook_contains_merged_hierarchy_and_review_columns(self):
        rows = load_mapping_rows(ROOT)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "review.xlsx"
            result = export_workbook(ROOT, output)
            workbook = load_workbook(output)
        sheet = workbook["按问题分组的规则审核"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(result["mapping_rows"], sheet.max_row - 1)
        self.assertEqual(headers[:3], ["一级问题", "二级问题", "三级问题"])
        for expected in ("规则UID", "法规/平台规则", "规则原文", "人工问题结论", "规则对应结论", "调整后问题"):
            self.assertIn(expected, headers)
        merged = {str(item) for item in sheet.merged_cells.ranges}
        self.assertTrue(any(item.startswith("A") for item in merged))
        self.assertTrue(any(item.startswith("B") for item in merged))
        self.assertTrue(any(item.startswith("C") for item in merged))
        first_issue_id = rows[0]["三级问题ID"]
        first_issue_size = sum(1 for row in rows if row["三级问题ID"] == first_issue_id)
        if first_issue_size > 1:
            self.assertIn(f"C2:C{first_issue_size + 1}", merged)
        self.assertIn("使用说明", workbook.sheetnames)


if __name__ == "__main__":
    unittest.main()
